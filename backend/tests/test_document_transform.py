from __future__ import annotations

import asyncio
from collections import UserDict
from copy import deepcopy
from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO
import json
from pathlib import Path
from threading import Event
from types import SimpleNamespace
from typing import Any, NamedTuple, Never, override
import uuid

from docx import Document as WordDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from googleapiclient.errors import HttpError
from loguru import logger
from pydantic import ValidationError
import pytest

from app.document_transform import (
    executor as executor_module,
    google_gateway as google_gateway_module,
    threaded_markdown_executor as threaded_executor_module,
)
from app.document_transform.brand import SarjBrandProfile, build_copy_mutations
from app.document_transform.custom_furniture import CustomFurnitureBinding
from app.document_transform.docx_renderer import markdown_import_manifest, render_markdown_docx
from app.document_transform.errors import (
    ActiveSuggestionsNeedsReviewError,
    BrandAssetUnavailableError,
    CheckpointTooLargeError,
    CustomHeaderNeedsReviewError,
    DestinationAccessChangedError,
    DocumentGatewayError,
    IncompatiblePendingOutputError,
    InPlaceNeedsReviewError,
    InPlacePostMutationVerificationError,
    OutputConflictError,
    OutputVerificationError,
    SourceChangedError,
    SourceTooLargeError,
    UnsupportedSourceError,
    UnsupportedTabMetadataError,
)
from app.document_transform.executor import REQUEST_BUILDER_VERSION, DocumentTransformExecutor
from app.document_transform.gateway import (
    DocumentGateway,
    InPlaceInspection,
    InPlaceStageResult,
    OutputDraft,
    OutputVerification,
    SourceSnapshot,
    TabHeader,
)
from app.document_transform.google_gateway import GoogleDocumentGateway, verify_brand_application
from app.document_transform.header_provenance import HeaderIdentityEvidence
from app.document_transform.ir import (
    BlockKind,
    DocumentIR,
    SourceRange,
    TextBlock,
    parse_google_document,
    parse_markdown,
    rewrite_markdown,
    utf16_length,
)
from app.document_transform.limits import MAX_MARKDOWN_BYTES
from app.document_transform.models import PatchGenerator
from app.document_transform.native_footer_flow import NativeFooterCheckpoint, NativeFooterStage
from app.document_transform.native_table_font import TableFontPlan
from app.document_transform.planner import TextTransformPlanner
from app.document_transform.safety import validate_patches
from app.document_transform.schemas import (
    CopyCheckpoint,
    CopyCheckpointStage,
    DocumentOutput,
    DocumentTransformCheckpoint,
    DocumentTransformRequest,
    DriveFileSource,
    FindingSeverity,
    InPlaceCheckpoint,
    InPlaceCheckpointStage,
    MarkdownSource,
    OutputStrategy,
    QaCopyPermissionEvidence,
    QaDestinationPermissionBinding,
    QaDocumentTransformRequest,
    QaDocumentTransformResult,
    QaFilePermissionEvidence,
    TabNamingCheckpoint,
    TextPatch,
    TransformFinding,
    TransformMode,
    TransformPhase,
)
from app.document_transform.tab_names import TabNamingPlan
from app.document_transform.threaded_markdown_executor import ThreadedMarkdownExecutor
from app.document_transform.unslop import (
    MAX_MODEL_EDITABLE_BLOCK_CHARS,
    deterministic_unslop_replacement,
    scan_unslop_candidates,
)
from app.drive.schemas import MIME_DOC
from app.files.errors import FileNotInScopeError
from app.files.schemas import FileResponse
from app.identifiers import DocsRevisionId, DocsSegmentId, DocsTabId, FileId, TaskId
from app.tasks import service as task_service_module
from tests.support import dependency


class _HttpErrorResponse(UserDict[str, str]):
    def __init__(self, status: int) -> None:
        super().__init__()
        self.status = status
        self.reason = "test"


class _DownloadProgress(NamedTuple):
    status: object
    done: bool


def _provider_http_error(status: int) -> HttpError:
    return HttpError(dependency(_HttpErrorResponse(status)), b'{"error":"private detail"}')


def test_markdown_source_limit_is_exactly_512_kib_of_strict_utf8() -> None:
    content_schema = MarkdownSource.model_json_schema()["properties"]["content"]
    assert content_schema["maxLength"] == MAX_MARKDOWN_BYTES

    exact = "é" * (MAX_MARKDOWN_BYTES // 2)
    source = MarkdownSource(
        kind="markdown",
        content=exact,
        name="Exact",
        destination_folder_id=FileId("folder"),
    )
    assert len(source.content.encode()) == MAX_MARKDOWN_BYTES

    with pytest.raises(ValidationError, match="UTF-8 bytes"):
        MarkdownSource(
            kind="markdown",
            content=f"{exact}a",
            name="Too large",
            destination_folder_id=FileId("folder"),
        )
    with pytest.raises(ValidationError, match="unicode"):
        MarkdownSource(
            kind="markdown",
            content="bad\ud800text",
            name="Invalid",
            destination_folder_id=FileId("folder"),
        )
    with pytest.raises(ValidationError):
        MarkdownSource(
            kind="markdown",
            content="a" * (MAX_MARKDOWN_BYTES + 1),
            name="Too many characters",
            destination_folder_id=FileId("folder"),
        )


def test_bounded_drive_markdown_download_rejects_declared_oversize_after_first_chunk(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls = 0

    class FakeDownload:
        def __init__(self, destination: BytesIO, _request: object, *, chunksize: int) -> None:
            assert chunksize < MAX_MARKDOWN_BYTES
            self.destination = destination

        def next_chunk(self, *, num_retries: int) -> _DownloadProgress:
            nonlocal calls
            assert num_retries == 0
            calls += 1
            self.destination.write(b"x")
            return _DownloadProgress(SimpleNamespace(total_size=MAX_MARKDOWN_BYTES + 1), False)

    monkeypatch.setattr(google_gateway_module, "MediaIoBaseDownload", FakeDownload)

    with pytest.raises(SourceTooLargeError):
        GoogleDocumentGateway._execute_bounded_bytes(
            dependency(object()),
            max_bytes=MAX_MARKDOWN_BYTES,
        )

    assert calls == 1


def test_bounded_drive_markdown_download_stops_when_unknown_total_crosses_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls = 0

    class FakeDownload:
        def __init__(self, destination: BytesIO, _request: object, *, chunksize: int) -> None:
            self.destination = destination
            self.chunksize = chunksize

        def next_chunk(self, *, num_retries: int) -> _DownloadProgress:
            nonlocal calls
            assert num_retries == 0
            calls += 1
            self.destination.write(b"x" * self.chunksize)
            return _DownloadProgress(SimpleNamespace(total_size=None), False)

    monkeypatch.setattr(google_gateway_module, "MediaIoBaseDownload", FakeDownload)

    with pytest.raises(SourceTooLargeError):
        GoogleDocumentGateway._execute_bounded_bytes(
            dependency(object()),
            max_bytes=MAX_MARKDOWN_BYTES,
        )

    assert calls == MAX_MARKDOWN_BYTES // google_gateway_module._MARKDOWN_DOWNLOAD_CHUNK_BYTES + 1


def test_supported_brand_only_downgrades_verified_header_conflicts() -> None:
    findings = [
        TransformFinding(
            code="header_conflict:tab-a",
            detail="Existing header preserved for review",
        ),
        TransformFinding(code="separate_review", detail="Still requires review"),
    ]

    classified = executor_module._classify_supported_copy_brand_findings(
        findings,
        supported_brand_applied=True,
    )

    assert [finding.severity for finding in classified] == [
        FindingSeverity.INFORMATIONAL,
        FindingSeverity.REVIEW,
    ]


@pytest.mark.parametrize(
    "method_name",
    ["_execute_dict", "_execute_mutation_dict", "_execute_bytes"],
)
def test_google_gateway_normalizes_transport_timeout(
    monkeypatch: pytest.MonkeyPatch,
    method_name: str,
) -> None:
    boundary_name = (
        "execute_mutation" if method_name == "_execute_mutation_dict" else "execute_with_retry"
    )
    private_detail = "private provider detail"

    def timeout(_request: object) -> Never:
        raise TimeoutError(private_detail)

    monkeypatch.setattr(google_gateway_module, boundary_name, timeout)

    with pytest.raises(DocumentGatewayError) as raised:
        getattr(GoogleDocumentGateway, method_name)(dependency(object()))

    assert raised.value.retryable
    assert raised.value.http_status is None
    assert private_detail not in str(raised.value)


@pytest.mark.parametrize(
    ("status", "retryable"),
    [(400, False), (403, False), (429, True), (503, True)],
)
def test_google_gateway_preserves_provider_retryability(status: int, retryable: bool) -> None:
    error = google_gateway_module._gateway_error(_provider_http_error(status))

    assert error.http_status == status
    assert error.retryable is retryable
    assert "private detail" not in str(error)


class _Generator(PatchGenerator):
    def __init__(self, replacements: dict[TransformPhase, str] | None = None) -> None:
        self.replacements = replacements or {}
        self.calls: list[TransformPhase] = []

    @property
    @override
    def model_ids(self) -> tuple[str, ...]:
        return ("gemini-3.1-flash-lite", "gemini-3.8-flash")

    @override
    async def generate(self, document: DocumentIR, *, phase: TransformPhase) -> list[TextPatch]:
        self.calls.append(phase)
        replacement = self.replacements.get(phase)
        if replacement is None:
            return []
        block = next(block for block in document.blocks if block.editable)
        return [
            TextPatch(
                block_id=block.id,
                expected_sha256=block.sha256,
                replacement=replacement,
                rationale="focused edit",
            )
        ]


class _PatchFixture(NamedTuple):
    document: DocumentIR
    patch: TextPatch


def _unslop_patch(source: str, replacement: str) -> _PatchFixture:
    document = parse_markdown(source, title="Safety")
    block = next(block for block in document.blocks if block.editable)
    return _PatchFixture(
        document=document,
        patch=TextPatch(
            block_id=block.id,
            expected_sha256=block.sha256,
            replacement=replacement,
            rationale="adversarial test",
        ),
    )


def test_unslop_rejects_changed_names_and_factual_predicate() -> None:
    document, patch = _unslop_patch(
        "Moreover, Alice approved the acquisition by Northstar.",
        "Moreover, Bob rejected the acquisition by Northstar.",
    )

    decision = validate_patches(document, [patch], phase=TransformPhase.UNSLOP)

    assert decision.accepted == {}
    assert decision.findings[0].code in {
        "proper_noun_changed",
        "noncandidate_content_changed",
    }


def test_unslop_rejects_edit_to_sentence_without_candidate() -> None:
    document, patch = _unslop_patch(
        "Moreover, we reviewed the plan. The board approved option A.",
        "We reviewed the plan. The board rejected option A.",
    )

    decision = validate_patches(document, [patch], phase=TransformPhase.UNSLOP)

    assert decision.accepted == {}
    assert decision.findings[0].code == "sentence_locality_changed"


@pytest.mark.parametrize(
    ("source", "replacement"),
    [
        ("Moreover, we reviewed the plan.", "We reviewed the plan."),
        (
            "Here's the thing: we're reviewing the user's plan.",
            "We're reviewing the user's plan.",
        ),
    ],
    ids=["transition", "contractions"],
)
def test_unslop_accepts_candidate_only_deletion(source: str, replacement: str) -> None:
    document, patch = _unslop_patch(source, replacement)

    decision = validate_patches(document, [patch], phase=TransformPhase.UNSLOP)

    assert decision.accepted == {"b0": replacement}


@pytest.mark.parametrize(
    "source",
    [
        "Moreover. We approved the plan.",
        "It is worth noting. We approved the plan.",
        "Furthermore! We approved the plan.",
        "Moreover… We approved the plan.",
    ],
)
def test_fixed_standalone_sentence_removes_terminal_punctuation(source: str) -> None:
    document = parse_markdown(source, title="Safety")
    candidates = scan_unslop_candidates(document)

    replacement = deterministic_unslop_replacement(document.blocks[0].text, candidates)
    decision = validate_patches(
        document,
        [
            TextPatch(
                block_id="b0",
                expected_sha256=document.blocks[0].sha256,
                replacement=replacement,
            )
        ],
        phase=TransformPhase.UNSLOP,
    )

    assert replacement == "We approved the plan."
    assert decision.accepted == {"b0": replacement}


@pytest.mark.parametrize(
    ("source", "expected"),
    [
        ("Moreover, Alice approved.", "Alice approved."),
        ("Moreover, SARJ approved.", "SARJ approved."),
        ("Moreover, eBay approved.", "Moreover, eBay approved."),
    ],
)
def test_fixed_prefix_preserves_existing_name_capitalization(source: str, expected: str) -> None:
    document = parse_markdown(source, title="Safety")

    replacement = deterministic_unslop_replacement(
        document.blocks[0].text,
        scan_unslop_candidates(document),
    )

    assert replacement == expected


@pytest.mark.parametrize(
    ("source", "docx_text", "expected_url"),
    [
        (
            "Moreover, we@example.com is the alias.",
            "Moreover, we@example.com is the alias.",
            None,
        ),
        ("Moreover, 'we preserve this quote.'", "Moreover, 'we preserve this quote.'", None),
        ("Moreover, ‘we preserve this quote.’", "Moreover, ‘we preserve this quote.’", None),
        ('Moreover, "we preserve this quote."', 'Moreover, "we preserve this quote."', None),
        ("Moreover, `we_keep_ID` is exact.", "Moreover, we_keep_ID is exact.", None),
        (
            "Moreover, [we use this label](https://example.com/a_(b)?x=1&y=2).",
            "Moreover, we use this label.",
            "https://example.com/a_(b)?x=1&y=2",
        ),
        ("Moreover, @we is the handle.", "Moreover, @we is the handle.", None),
        ("Moreover, /we/path is literal.", "Moreover, /we/path is literal.", None),
        ("Moreover, #we is the tag.", "Moreover, #we is the tag.", None),
        ("Moreover, _we_var is exact.", "Moreover, _we_var is exact.", None),
        ("Moreover, we://host is exact.", "Moreover, we://host is exact.", None),
        ("Moreover, we:opaque-id is exact.", "Moreover, we:opaque-id is exact.", None),
        ("Moreover, we&amp;they is literal.", "Moreover, we&they is literal.", None),
    ],
)
async def test_unslop_planner_fails_closed_before_protected_recapitalization_in_docx(
    source: str,
    docx_text: str,
    expected_url: str | None,
) -> None:
    document = parse_markdown(source, title="Protected recapitalization")
    request = DocumentTransformRequest(
        source=MarkdownSource(
            kind="markdown",
            content=source,
            name="Protected recapitalization",
            destination_folder_id=FileId("folder"),
        ),
        mode=TransformMode.UNSLOP,
    )
    generator = _Generator()

    outcome = await TextTransformPlanner(generator=generator).transform(document, request)

    assert outcome.document == document
    assert outcome.applied_count == 0
    assert outcome.deterministic_count == 0
    assert not outcome.llm_invoked
    assert generator.calls == []
    assert {finding.code for finding in outcome.findings} == {"contextual_unslop_review_required"}
    rendered = render_markdown_docx(
        outcome.document.blocks[0].text,
        title="",
        branded=False,
    )
    imported = WordDocument(BytesIO(rendered))
    assert "\n".join(paragraph.text for paragraph in imported.paragraphs) == docx_text
    hyperlink_targets = {
        relationship.target_ref
        for relationship in imported.part.rels.values()
        if relationship.reltype == RELATIONSHIP_TYPE.HYPERLINK
    }
    assert hyperlink_targets == ({expected_url} if expected_url is not None else set())


async def test_unslop_planner_value_add_preserves_inline_topology_in_docx() -> None:
    source = "Moreover, **we** approved [SARJ](https://sarj.ai) with `API_v2`."
    expected = "**We** approved [SARJ](https://sarj.ai) with `API_v2`."
    document = parse_markdown(source, title="Value-added recapitalization")
    request = DocumentTransformRequest(
        source=MarkdownSource(
            kind="markdown",
            content=source,
            name="Value-added recapitalization",
            destination_folder_id=FileId("folder"),
        ),
        mode=TransformMode.UNSLOP,
    )
    generator = _Generator()

    outcome = await TextTransformPlanner(generator=generator).transform(document, request)

    assert outcome.document.blocks[0].text == expected
    assert outcome.applied_count == outcome.deterministic_count == 1
    assert not outcome.llm_invoked
    assert outcome.findings == ()
    manifest = markdown_import_manifest(expected, title="", branded=False)
    assert manifest.blocks[0].bold == ("We",)
    assert manifest.blocks[0].inline_code == ("API_v2",)
    assert manifest.blocks[0].links == (("SARJ", "https://sarj.ai"),)
    imported = WordDocument(BytesIO(render_markdown_docx(expected, title="", branded=False)))
    assert imported.paragraphs[0].text == "We approved SARJ with API_v2."
    assert {
        relationship.target_ref
        for relationship in imported.part.rels.values()
        if relationship.reltype == RELATIONSHIP_TYPE.HYPERLINK
    } == {"https://sarj.ai"}


@pytest.mark.parametrize(
    ("source", "unsafe_replacement", "expected_code"),
    [
        (
            "Moreover, we@example.com is the alias.",
            "We@example.com is the alias.",
            "protected_content_changed",
        ),
        (
            "Moreover, 'we preserve this quote.'",
            "'We preserve this quote.'",
            "quoted_content_changed",
        ),
        (
            "Moreover, ‘we preserve this quote.’",
            "‘We preserve this quote.’",
            "quoted_content_changed",
        ),
        (
            "Moreover, we://host is exact.",
            "Moreover, We://host is exact.",
            "protected_content_changed",
        ),
        (
            "Moreover, we&amp;they is literal.",
            "Moreover, we&lt;they is literal.",
            "protected_content_changed",
        ),
    ],
)
def test_unslop_safety_rejects_email_and_single_quote_recapitalization(
    source: str,
    unsafe_replacement: str,
    expected_code: str,
) -> None:
    document, patch = _unslop_patch(source, unsafe_replacement)

    decision = validate_patches(document, [patch], phase=TransformPhase.UNSLOP)

    assert decision.accepted == {}
    assert decision.findings[0].code == expected_code


@pytest.mark.parametrize(
    ("source", "replacement"),
    [
        ("This is a game changer.", "This is a ."),
        ("These results speak for themselves.", "These results ."),
        ("We launched Game Changer yesterday.", "We launched  yesterday."),
        ("The team uses Synergy for payroll.", "The team uses  for payroll."),
    ],
)
def test_contextual_unslop_is_review_only(source: str, replacement: str) -> None:
    document, patch = _unslop_patch(source, replacement)

    decision = validate_patches(document, [patch], phase=TransformPhase.UNSLOP)

    assert decision.accepted == {}
    assert decision.findings[0].code == "contextual_unslop_review_required"


async def test_planner_reports_native_rich_text_and_oversized_unslop_blocks() -> None:
    document = DocumentIR(
        title="Safety",
        source_format="google_doc",
        blocks=(
            TextBlock(
                id="rich", kind=BlockKind.PARAGRAPH, text="Visible rich text", editable=False
            ),
            TextBlock(
                id="large",
                kind=BlockKind.PARAGRAPH,
                text="Moreover, " + "x" * MAX_MODEL_EDITABLE_BLOCK_CHARS,
            ),
            TextBlock(id="contextual", kind=BlockKind.PARAGRAPH, text="This is synergy."),
        ),
    )
    request = DocumentTransformRequest(
        source=DriveFileSource(
            kind="drive_file",
            file_id=FileId("source"),
            expected_modified_time="2026-09-03T00:00:00Z",
        ),
        mode=TransformMode.BOTH,
    )
    generator = _Generator({TransformPhase.UNSLOP: "unsafe model rewrite"})

    result = await TextTransformPlanner(generator=generator).transform(document, request)

    assert generator.calls == []
    assert {finding.code for finding in result.findings} == {
        "contextual_unslop_review_required",
        "google_doc_text_target_unsupported",
        "unslop_block_too_large",
    }
    assert result.document == document


@pytest.mark.parametrize(
    ("source", "replacement"),
    [
        ("Moreover, we reviewed the plan.", "We we reviewed the plan."),
        ("Moreover, Alice approved the plan.", "Alice Alice approved the plan."),
        ("Moreover, revenue was $47.3M.", "Revenue was 47.3M."),
        ("Moreover, adoption reached 23%.", "Adoption reached 23."),
    ],
)
def test_unslop_rejects_insertions_duplication_and_numeric_symbol_loss(
    source: str,
    replacement: str,
) -> None:
    document, patch = _unslop_patch(source, replacement)

    decision = validate_patches(document, [patch], phase=TransformPhase.UNSLOP)

    assert decision.accepted == {}


@pytest.mark.parametrize("name", ["eBay", "iPhone", "macOS", "npm"])
def test_unslop_never_recapitalizes_brand_or_identifier(name: str) -> None:
    document, patch = _unslop_patch(
        f"Moreover, {name} approved the plan.",
        f"{name[0].upper() + name[1:]} approved the plan.",
    )

    decision = validate_patches(document, [patch], phase=TransformPhase.UNSLOP)

    assert decision.accepted == {}


def test_unslop_protects_markdown_blockquotes() -> None:
    document = parse_markdown(
        "> Additionally, this quoted example must stay exact.",
        title="Quoted source",
    )
    block = next(block for block in document.blocks if block.text)
    patch = TextPatch(
        block_id=block.id,
        expected_sha256=block.sha256,
        replacement="This quoted example must stay exact.",
    )

    decision = validate_patches(document, [patch], phase=TransformPhase.UNSLOP)

    assert decision.accepted == {}
    assert decision.findings[0].code == "protected_block"


@pytest.mark.parametrize(
    "source",
    [
        ">Additionally, preserve this quotation.",
        "~~~text\nAdditionally, preserve this sample.\n~~~",
        "    Additionally, preserve this indented code.",
    ],
)
def test_markdown_structural_regions_are_not_unslop_candidates(source: str) -> None:
    document = parse_markdown(source, title="Protected examples")

    assert scan_unslop_candidates(document) == ()


def test_unslop_scanner_masks_quoted_examples_and_protects_domain_terms() -> None:
    quoted = parse_markdown(
        'The guide says "It is important to note this bad example."',
        title="Guide",
    )
    technical = parse_markdown(
        "The robust hash verification rejected the stale block.",
        title="Design",
    )

    assert scan_unslop_candidates(quoted) == ()
    assert scan_unslop_candidates(technical) == ()


def test_unslop_scanner_masks_single_quoted_examples_but_not_contractions() -> None:
    quoted = parse_markdown(
        "The guide labels 'It is important to note that' as filler.",
        title="Guide",
    )
    contraction = parse_markdown("Here's the thing: we're ready.", title="Draft")

    assert scan_unslop_candidates(quoted) == ()
    assert {candidate.code for candidate in scan_unslop_candidates(contraction)} == {
        "throat_clearing"
    }


def test_literal_and_legal_uses_are_not_candidates() -> None:
    document = parse_markdown(
        "The team will delve into the mountain. The guard stands as the anthem plays. "
        "Additionally, the agreement remains subject to Section 4.",
        title="Literal and legal examples",
    )

    assert scan_unslop_candidates(document) == ()


@pytest.mark.parametrize(
    "source",
    [
        "Whether it is worth noting that taxes rose is disputed.",
        "She asked whether it is important to note that the deadline moved.",
        "The phrase here's the thing: is common in speech.",
    ],
)
def test_embedded_meta_phrases_are_never_unslop_candidates(source: str) -> None:
    assert scan_unslop_candidates(parse_markdown(source, title="Quoted claim")) == ()


@pytest.mark.parametrize(
    "source",
    [
        "```text\nIt is important to note this example.\n```not-a-close\n```",
        "---\ntitle: It is important to note this value\n---\nApproved prose.",
        "Label | Detail\n--- | ---\nA | It is important to note this value",
    ],
    ids=["fence", "frontmatter", "table-without-outer-pipes"],
)
def test_markdown_metadata_and_structures_are_not_unslop_candidates(source: str) -> None:
    assert scan_unslop_candidates(parse_markdown(source, title="Protected")) == ()


def test_raw_html_markdown_is_rejected_before_transform() -> None:
    with pytest.raises(UnsupportedSourceError):
        parse_markdown("<!-- It is important to note this -->", title="Raw HTML")


def test_markdown_block_count_is_bounded_before_model_or_drive_work() -> None:
    source = "\n".join(f"Paragraph {index}" for index in range(751))

    with pytest.raises(SourceTooLargeError):
        parse_markdown(source, title="Too many blocks")


def test_native_google_doc_block_count_is_bounded_before_output_creation() -> None:
    content = [
        {
            "startIndex": index * 2 + 1,
            "endIndex": index * 2 + 3,
            "paragraph": {"elements": [{"textRun": {"content": "x\n"}}]},
        }
        for index in range(751)
    ]
    document = {
        "revisionId": "revision-1",
        "tabs": [
            {
                "tabProperties": {"tabId": "tab-a"},
                "documentTab": {"body": {"content": content}},
                "childTabs": [],
            }
        ],
    }

    with pytest.raises(SourceTooLargeError):
        parse_google_document(document, title="Too many blocks")


def test_native_mixed_text_paragraph_is_preserved_as_uneditable_for_review() -> None:
    document = {
        "revisionId": "revision-1",
        "tabs": [
            {
                "tabProperties": {"tabId": "tab-a"},
                "documentTab": {
                    "body": {
                        "content": [
                            {
                                "startIndex": 1,
                                "endIndex": 14,
                                "paragraph": {
                                    "elements": [
                                        {
                                            "startIndex": 1,
                                            "endIndex": 8,
                                            "textRun": {"content": "Visible"},
                                        },
                                        {
                                            "startIndex": 8,
                                            "endIndex": 13,
                                            "inlineObjectElement": {"inlineObjectId": "image-1"},
                                        },
                                        {
                                            "startIndex": 13,
                                            "endIndex": 14,
                                            "textRun": {"content": "\n"},
                                        },
                                    ]
                                },
                            }
                        ]
                    }
                },
                "childTabs": [],
            }
        ],
    }

    parsed = parse_google_document(document, title="Mixed")

    assert parsed.blocks[0].text == "Visible"
    assert parsed.blocks[0].editable is False


@pytest.mark.parametrize("native_list", [False, True])
async def test_native_mutation_budget_is_checked_before_copy_creation(
    native_list: bool, monkeypatch: pytest.MonkeyPatch
) -> None:
    content: list[dict[str, Any]] = []
    cursor = 1
    paragraph = "Review evidence.\n" if native_list else "Moreover, we approved.\n"
    if native_list:
        # 750 unsplit paragraphs fit; the additional 750 proper subranges do not.
        monkeypatch.setattr("app.document_transform.brand.MAX_DOCS_MUTATION_REQUESTS", 2_000)
    for _ in range(750):
        end = cursor + utf16_length(paragraph)
        content.append(
            {
                "startIndex": cursor,
                "endIndex": end,
                "paragraph": {
                    "elements": [{"textRun": {"content": paragraph}}],
                    **(
                        {
                            "bullet": {"listId": "native-list", "nestingLevel": 0},
                            "paragraphStyle": {"namedStyleType": "NORMAL_TEXT"},
                        }
                        if native_list
                        else {}
                    ),
                },
            }
        )
        cursor = end
    document = {
        "revisionId": "source-revision-1",
        "tabs": [
            {
                "tabProperties": {"tabId": "tab-a"},
                "documentTab": {"body": {"content": content}},
                "childTabs": [],
            }
        ],
    }
    gateway = _Gateway(
        SourceSnapshot(
            title="Large review",
            parent_folder_id=FileId("folder"),
            modified_time="2026-09-03T00:00:00Z",
            revision_id=DocsRevisionId("source-revision-1"),
            media_type="application/vnd.google-apps.document",
            google_document=document,
        )
    )
    request = DocumentTransformRequest(
        source=DriveFileSource(
            kind="drive_file",
            file_id=FileId("source"),
            expected_modified_time="2026-09-03T00:00:00Z",
        ),
        mode=TransformMode.BOTH,
    )

    with pytest.raises(SourceTooLargeError):
        await DocumentTransformExecutor(gateway=gateway, generator=_Generator()).execute(
            operation_id=uuid.uuid7(),
            request=request,
        )

    assert gateway.brand_prepares == 0
    assert gateway.draft is None


def test_unslop_scanner_catches_high_signal_upstream_phrases() -> None:
    document = parse_markdown(
        "Here's the thing: at the end of the day, this is a game changer.",
        title="Draft",
    )

    codes = {candidate.code for candidate in scan_unslop_candidates(document)}

    assert {"throat_clearing", "filler", "jargon"} <= codes


def test_unslop_scanner_handles_native_docs_typography() -> None:
    document = parse_markdown(
        "Here’s the thing: in today’s fast‑paced market, this is a game‑changer.",
        title="Draft",
    )

    assert {candidate.code for candidate in scan_unslop_candidates(document)} == {
        "canned_intro",
        "jargon",
        "throat_clearing",
    }


class _Gateway(DocumentGateway):
    @override
    async def plan_tab_naming(
        self,
        *,
        document_id: FileId,
        title: str,
        operation_id: uuid.UUID,
        required_revision_id: DocsRevisionId,
        prefer_document_title: bool = False,
    ) -> None:
        return None

    def __init__(
        self,
        snapshot: SourceSnapshot | None = None,
        *,
        verification: OutputVerification | None = None,
    ) -> None:
        self.snapshot = snapshot
        self.verification = verification or OutputVerification(
            content_preserved=True,
            structure_preserved=True,
        )
        self.draft: OutputDraft | None = None
        self.requests: list[dict[str, Any]] = []
        self.applies = 0
        self.completed_result: dict[str, Any] | None = None
        self.brand_prepares = 0
        self.imported_docx: bytes | None = None
        self.verified_header_titles: list[str | None] = []

    @override
    async def load_drive_source(self, source: DriveFileSource) -> SourceSnapshot:
        assert self.snapshot is not None
        return self.snapshot

    @override
    async def recover_output(self, operation_id: uuid.UUID) -> OutputDraft | None:
        return self.draft

    @override
    async def source_is_unchanged(
        self,
        source: DriveFileSource,
        *,
        expected_modified_time: str,
        expected_revision_id: DocsRevisionId | None,
        expected_parent_folder_id: FileId,
    ) -> bool:
        return True

    @override
    async def create_from_drive_copy(
        self,
        *,
        operation_id: uuid.UUID,
        source: DriveFileSource,
        source_sha256: str,
        output_name: str,
        parent_folder_id: FileId,
    ) -> OutputDraft:
        self.brand_prepares += 1
        return self._create(source_sha256, output_name, parent_folder_id=parent_folder_id)

    @override
    async def create_from_drive_markdown(
        self,
        *,
        operation_id: uuid.UUID,
        source: DriveFileSource,
        source_sha256: str,
        output_name: str,
        parent_folder_id: FileId,
        imported_docx: bytes | None = None,
    ) -> OutputDraft:
        self.imported_docx = imported_docx
        return self._create(source_sha256, output_name, parent_folder_id=parent_folder_id)

    @override
    async def create_from_markdown(
        self,
        *,
        operation_id: uuid.UUID,
        source: MarkdownSource,
        source_sha256: str,
        output_name: str,
        imported_docx: bytes | None = None,
    ) -> OutputDraft:
        self.imported_docx = imported_docx
        return self._create(source_sha256, output_name)

    @override
    async def prepare_brand_shell(
        self,
        *,
        operation_id: uuid.UUID,
        draft: OutputDraft,
        managed_headers: dict[str, str] | None = None,
    ) -> OutputDraft:
        del managed_headers
        self.draft = OutputDraft(
            output=draft.output,
            document_id=draft.document_id,
            tab_id=draft.tab_id,
            revision_id=draft.revision_id,
            source_sha256=draft.source_sha256,
            headers=(
                TabHeader(
                    tab_id=DocsTabId("tab-1"),
                    header_id=DocsSegmentId("header-1"),
                ),
            ),
            logo_uri="https://assets.example.test/sarj-logo.png",
        )
        return self.draft

    def _create(
        self,
        source_sha256: str,
        output_name: str,
        *,
        parent_folder_id: FileId | None = None,
    ) -> OutputDraft:
        parent_folder_id = parent_folder_id or FileId("folder")
        self.draft = OutputDraft(
            output=DocumentOutput(
                file_id=FileId("output"),
                name=output_name,
                parent_folder_id=parent_folder_id,
            ),
            document_id=FileId("output"),
            tab_id=DocsTabId("tab-1"),
            revision_id=DocsRevisionId("output-revision-1"),
            source_sha256=source_sha256,
        )
        return self.draft

    @override
    async def apply_requests_once(
        self,
        *,
        operation_id: uuid.UUID,
        draft: OutputDraft,
        mutation_sha256: str,
        requests: list[dict[str, Any]],
        required_revision_id: DocsRevisionId,
    ) -> OutputDraft:
        assert required_revision_id == "output-revision-1"
        assert mutation_sha256
        self.applies += 1
        self.requests = requests
        self.draft = OutputDraft(
            output=draft.output,
            document_id=draft.document_id,
            tab_id=draft.tab_id,
            revision_id=DocsRevisionId("output-revision-2"),
            source_sha256=draft.source_sha256,
            headers=draft.headers,
            logo_uri=draft.logo_uri,
        )
        return self.draft

    @override
    async def verify_output(
        self,
        *,
        draft: OutputDraft,
        source: SourceSnapshot,
        expected: DocumentIR,
        verify_brand: bool,
        header_title: str | None = None,
    ) -> OutputVerification:
        del source, expected, verify_brand
        self.verified_header_titles.append(header_title)
        return OutputVerification(
            content_preserved=self.verification.content_preserved,
            structure_preserved=self.verification.structure_preserved,
            brand_applied=self.verification.brand_applied,
            revision_id=draft.revision_id,
            supported_brand_applied=self.verification.supported_brand_applied,
        )

    @override
    async def complete(
        self,
        *,
        operation_id: uuid.UUID,
        draft: OutputDraft,
        result: dict[str, Any],
    ) -> None:
        self.completed_result = result
        self.draft = OutputDraft(
            output=draft.output,
            document_id=draft.document_id,
            tab_id=draft.tab_id,
            revision_id=draft.revision_id,
            source_sha256=draft.source_sha256,
            headers=draft.headers,
            logo_uri=draft.logo_uri,
            completed=True,
            result=result,
            result_sha256=google_gateway_module._json_sha256(result),
            verified_revision_sha256=google_gateway_module._revision_sha256(draft.revision_id),
            brand_warnings=draft.brand_warnings,
        )


async def test_markdown_both_runs_unslop_then_format_and_creates_sibling() -> None:
    gateway = _Gateway()
    generator = _Generator(
        {
            TransformPhase.UNSLOP: "We reviewed the options.",
            TransformPhase.FORMAT: "We reviewed the options carefully.",
        }
    )
    request = DocumentTransformRequest(
        source=MarkdownSource(
            kind="markdown",
            content="It is important to note that we reviewed the options.",
            name="Review",
            destination_folder_id=FileId("folder"),
        ),
        mode=TransformMode.BOTH,
    )

    result = await DocumentTransformExecutor(gateway=gateway, generator=generator).execute(
        operation_id=uuid.uuid7(), request=request
    )

    assert generator.calls == []
    assert result.output.name == "Review — Sarj Polished"
    assert result.applied_patch_count == 1
    assert result.tool_versions["unslop_version"] == "2.3.0"
    assert result.validation.structure_preserved is True
    assert {finding.code for finding in result.findings} == {"logo_unavailable"}
    assert result.needs_review is True
    assert gateway.imported_docx is not None
    assert gateway.imported_docx.startswith(b"PK")
    assert gateway.requests == []


async def test_markdown_render_keeps_worker_event_loop_responsive(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = _Gateway()
    request = DocumentTransformRequest(
        source=MarkdownSource(
            kind="markdown",
            content="A short report.",
            name="Report",
            destination_folder_id=FileId("folder"),
        ),
        mode=TransformMode.FORMAT,
    )
    original_render = threaded_executor_module.render_markdown_docx
    started = Event()
    release = Event()

    def delayed_render(
        markdown: str,
        *,
        title: str,
        logo_bytes: bytes | None = None,
        branded: bool = True,
        header_title: str | None = None,
    ) -> bytes:
        started.set()
        assert release.wait(2)
        return original_render(
            markdown,
            title=title,
            logo_bytes=logo_bytes,
            branded=branded,
            header_title=header_title,
        )

    monkeypatch.setattr(threaded_executor_module, "render_markdown_docx", delayed_render)
    transform = asyncio.create_task(
        ThreadedMarkdownExecutor(gateway=gateway, generator=_Generator()).execute(
            operation_id=uuid.uuid7(), request=request
        )
    )
    try:
        assert await asyncio.to_thread(started.wait, 1)
        assert not transform.done()
    finally:
        release.set()
    assert (await transform).validation.structure_preserved is True


async def test_drive_markdown_uses_threaded_renderer_and_keeps_source_parent() -> None:
    gateway = _Gateway(
        SourceSnapshot(
            title="Drive report",
            parent_folder_id=FileId("source-parent"),
            modified_time="2026-09-23T00:00:00Z",
            revision_id=None,
            media_type="text/markdown",
            markdown="# Report\n\nBody",
        )
    )
    request = DocumentTransformRequest(
        source=DriveFileSource(
            kind="drive_file",
            file_id=FileId("source"),
            expected_modified_time="2026-09-23T00:00:00Z",
        ),
        mode=TransformMode.FORMAT,
    )

    result = await ThreadedMarkdownExecutor(gateway=gateway, generator=_Generator()).execute(
        operation_id=uuid.uuid7(), request=request
    )

    assert gateway.imported_docx is not None
    assert gateway.imported_docx.startswith(b"PK")
    assert result.output.parent_folder_id == FileId("source-parent")


async def test_markdown_docx_import_reports_unverified_brand_without_stale_fidelity_claims() -> (
    None
):
    gateway = _Gateway()
    logo_bytes = (
        Path(__file__).parents[2] / "frontend/public/brand/sarj-logo-light.png"
    ).read_bytes()
    request = DocumentTransformRequest(
        source=MarkdownSource(
            kind="markdown",
            content="**Decision**\n\n| Owner | Status |\n| --- | --- |\n| Wiki | Ready |",
            name="Review",
            destination_folder_id=FileId("folder"),
        ),
        mode=TransformMode.FORMAT,
    )

    result = await DocumentTransformExecutor(
        gateway=gateway,
        generator=_Generator(),
        logo_bytes=logo_bytes,
    ).execute(operation_id=uuid.uuid7(), request=request)

    codes = {finding.code for finding in result.findings}
    assert codes == set()
    assert "markdown_table_preserved_as_text" not in codes
    assert "markdown_inline_markup_preserved_as_text" not in codes
    assert result.needs_review is False
    assert result.validation.structure_preserved is True
    assert gateway.imported_docx is not None


async def test_unslop_clean_document_is_exact_content_noop_without_model_call() -> None:
    gateway = _Gateway()
    generator = _Generator()
    request = DocumentTransformRequest(
        source=MarkdownSource(
            kind="markdown",
            content="The team approved the proposal.",
            name="Decision",
            destination_folder_id=FileId("folder"),
        ),
        mode=TransformMode.UNSLOP,
    )

    result = await DocumentTransformExecutor(gateway=gateway, generator=generator).execute(
        operation_id=uuid.uuid7(), request=request
    )

    assert generator.calls == []
    assert result.applied_patch_count == 0
    assert gateway.requests == []
    assert gateway.imported_docx is not None
    imported = WordDocument(BytesIO(gateway.imported_docx))
    assert [paragraph.text for paragraph in imported.paragraphs] == [
        "Decision",
        "The team approved the proposal.",
    ]
    assert gateway.brand_prepares == 0
    assert "updateDocumentStyle" not in str(gateway.requests)
    assert "insertInlineImage" not in str(gateway.requests)


async def test_completed_output_recovery_never_reruns_model_or_mutations() -> None:
    gateway = _Gateway()
    generator = _Generator({TransformPhase.UNSLOP: "Changed text."})
    operation_id = uuid.uuid7()
    source = MarkdownSource(
        kind="markdown",
        content="The board approved option A.",
        name="Decision",
        destination_folder_id=FileId("folder"),
    )
    request = DocumentTransformRequest(source=source, mode=TransformMode.UNSLOP)
    first = await DocumentTransformExecutor(gateway=gateway, generator=generator).execute(
        operation_id=operation_id,
        request=request,
    )
    assert gateway.draft is not None
    gateway.draft = OutputDraft(
        output=gateway.draft.output,
        document_id=gateway.draft.document_id,
        tab_id=gateway.draft.tab_id,
        revision_id=gateway.draft.revision_id,
        source_sha256=first.source_sha256,
        completed=True,
    )
    generator.calls.clear()
    gateway.applies = 0

    with pytest.raises(OutputConflictError):
        await DocumentTransformExecutor(gateway=gateway, generator=generator).execute(
            operation_id=operation_id,
            request=request,
        )

    assert generator.calls == []
    assert gateway.applies == 0


async def test_force_language_change_is_rejected_and_requires_review() -> None:
    gateway = _Gateway()
    generator = _Generator(
        {TransformPhase.UNSLOP: "This is a game changer: teams may encrypt all records."}
    )
    request = DocumentTransformRequest(
        source=MarkdownSource(
            kind="markdown",
            content="This is a game changer: teams must encrypt all records.",
            name="Policy",
            destination_folder_id=FileId("folder"),
        ),
        mode=TransformMode.UNSLOP,
    )

    result = await DocumentTransformExecutor(gateway=gateway, generator=generator).execute(
        operation_id=uuid.uuid7(), request=request
    )

    assert result.applied_patch_count == 0
    assert result.needs_review is True
    assert generator.calls == []
    assert result.findings[0].code == "contextual_unslop_review_required"
    assert gateway.requests == []
    assert gateway.imported_docx is not None
    imported = WordDocument(BytesIO(gateway.imported_docx))
    assert "must encrypt" in "\n".join(paragraph.text for paragraph in imported.paragraphs)


async def test_completed_operation_is_replayed_without_second_mutation() -> None:
    gateway = _Gateway()
    generator = _Generator()
    executor = DocumentTransformExecutor(gateway=gateway, generator=generator)
    request = DocumentTransformRequest(
        source=MarkdownSource(
            kind="markdown",
            content="A concise source.",
            name="Source",
            destination_folder_id=FileId("folder"),
        ),
        mode=TransformMode.FORMAT,
    )
    operation_id = uuid.uuid7()

    first = await executor.execute(operation_id=operation_id, request=request)
    second = await executor.execute(operation_id=operation_id, request=request)

    assert first == second
    assert gateway.applies == 1


@pytest.mark.parametrize(
    ("mode", "verification", "expected_codes", "expected_review"),
    [
        (
            TransformMode.FORMAT,
            OutputVerification(
                content_preserved=True,
                structure_preserved=True,
                brand_applied=True,
            ),
            set(),
            False,
        ),
        (
            TransformMode.UNSLOP,
            OutputVerification(content_preserved=False, structure_preserved=False),
            {"post_write_content_mismatch", "post_write_structure_mismatch"},
            True,
        ),
        (
            TransformMode.FORMAT,
            OutputVerification(
                content_preserved=True,
                structure_preserved=True,
                brand_applied=False,
            ),
            {"post_write_brand_mismatch"},
            True,
        ),
        (
            TransformMode.FORMAT,
            OutputVerification(
                content_preserved=True,
                structure_preserved=True,
                brand_applied=False,
                supported_brand_applied=True,
            ),
            {"custom_furniture_preserved"},
            False,
        ),
    ],
    ids=["verified", "content-and-structure", "brand", "supported-brand"],
)
@pytest.mark.parametrize("rich", [False, True], ids=["plain", "preserved-rich"])
async def test_native_doc_post_write_mismatch_is_reported_for_review(
    mode: TransformMode,
    verification: OutputVerification,
    expected_codes: set[str],
    expected_review: bool,
    rich: bool,
) -> None:
    document: dict[str, Any] = {
        "revisionId": "source-revision-1",
        "tabs": [
            {
                "tabProperties": {"tabId": "tab-1"},
                "documentTab": {
                    "body": {
                        "content": [
                            {
                                "startIndex": 1,
                                "endIndex": 29,
                                "paragraph": {
                                    "elements": [
                                        {"textRun": {"content": "The team approved the plan.\n"}}
                                    ]
                                },
                            }
                        ]
                    }
                },
                "childTabs": [],
            }
        ],
    }
    if rich:
        document["tabs"][0]["documentTab"]["body"]["content"][0]["paragraph"]["elements"].append(
            {"inlineObjectElement": {"inlineObjectId": "authored-object"}}
        )
    gateway = _Gateway(
        SourceSnapshot(
            title="Decision",
            parent_folder_id=FileId("folder"),
            modified_time="2026-09-03T00:00:00Z",
            revision_id=DocsRevisionId("source-revision-1"),
            media_type="application/vnd.google-apps.document",
            google_document=document,
        ),
        verification=verification,
    )
    request = DocumentTransformRequest(
        source=DriveFileSource(
            kind="drive_file",
            file_id=FileId("source"),
            expected_modified_time="2026-09-03T00:00:00Z",
        ),
        mode=mode,
    )

    if not verification.content_preserved or not verification.structure_preserved:
        with pytest.raises(OutputVerificationError):
            await DocumentTransformExecutor(gateway=gateway, generator=_Generator()).execute(
                operation_id=uuid.uuid7(),
                request=request,
            )
        assert gateway.completed_result is None
        return

    executor = DocumentTransformExecutor(gateway=gateway, generator=_Generator())
    operation_id = uuid.uuid7()
    result = await executor.execute(operation_id=operation_id, request=request)
    applies = gateway.applies
    assert await executor.execute(operation_id=operation_id, request=request) == result
    assert gateway.applies == applies

    assert result.needs_review is expected_review
    assert result.validation.protected_content_preserved is verification.content_preserved
    assert result.validation.structure_preserved is verification.structure_preserved
    assert {finding.code for finding in result.findings} == expected_codes | (
        {"google_doc_text_target_unsupported"} if rich else set()
    )
    rich_findings = [
        finding
        for finding in result.findings
        if finding.code == "google_doc_text_target_unsupported"
    ]
    assert all(finding.severity.value == "informational" for finding in rich_findings)
    if verification.supported_brand_applied:
        supported = next(
            finding for finding in result.findings if finding.code == "custom_furniture_preserved"
        )
        assert supported.severity.value == "informational"
        assert "Complete Sarj branding was not applied" in supported.detail
    assert result.validation.unsafe_patch_count == 0


async def test_markdown_import_readback_mismatch_is_reported_for_review() -> None:
    gateway = _Gateway(
        verification=OutputVerification(content_preserved=False, structure_preserved=False)
    )
    request = DocumentTransformRequest(
        source=MarkdownSource(
            kind="markdown",
            content="# Decision\n\n1. First\n2. Second",
            name="Review",
            destination_folder_id=FileId("folder"),
        ),
        mode=TransformMode.UNSLOP,
    )

    with pytest.raises(OutputVerificationError):
        await DocumentTransformExecutor(gateway=gateway, generator=_Generator()).execute(
            operation_id=uuid.uuid7(),
            request=request,
        )
    assert gateway.completed_result is None


def test_google_doc_ir_uses_tab_ranges_and_locks_rich_text() -> None:
    document = {
        "revisionId": "revision-1",
        "tabs": [
            {
                "tabProperties": {"tabId": "tab-a"},
                "documentTab": {
                    "body": {
                        "content": [
                            {
                                "startIndex": 1,
                                "endIndex": 8,
                                "paragraph": {
                                    "paragraphStyle": {"namedStyleType": "HEADING_1"},
                                    "elements": [{"textRun": {"content": "Roadmap\n"}}],
                                },
                            },
                            {
                                "startIndex": 8,
                                "endIndex": 17,
                                "paragraph": {
                                    "elements": [
                                        {
                                            "textRun": {
                                                "content": "See ",
                                                "textStyle": {},
                                            }
                                        },
                                        {
                                            "textRun": {
                                                "content": "source\n",
                                                "textStyle": {"link": {"url": "https://sarj.ai"}},
                                            }
                                        },
                                    ]
                                },
                            },
                        ]
                    }
                },
                "childTabs": [],
            }
        ],
    }

    parsed = parse_google_document(document, title="Review")

    assert parsed.blocks[0].source_range is not None
    assert parsed.blocks[0].source_range.tab_id == "tab-a"
    assert parsed.blocks[0].level == 1
    assert parsed.blocks[1].editable is False


def test_google_doc_single_font_styled_run_remains_editable() -> None:
    document = {
        "revisionId": "revision-1",
        "tabs": [
            {
                "tabProperties": {"tabId": "tab-a"},
                "documentTab": {
                    "body": {
                        "content": [
                            {
                                "startIndex": 1,
                                "endIndex": 33,
                                "paragraph": {
                                    "elements": [
                                        {
                                            "textRun": {
                                                "content": "Additionally, review the plan.\n",
                                                "textStyle": {
                                                    "weightedFontFamily": {"fontFamily": "Nunito"}
                                                },
                                            }
                                        }
                                    ]
                                },
                            }
                        ]
                    }
                },
                "childTabs": [],
            }
        ],
    }

    parsed = parse_google_document(document, title="Review")

    assert parsed.blocks[0].editable is True


def test_scanner_reports_candidates_but_never_changes_text() -> None:
    source = "In today's fast-paced market, this robust plan serves as a testament to the team."
    document = parse_markdown(source, title="Review")

    candidates = scan_unslop_candidates(document)

    assert {candidate.code for candidate in candidates} == {
        "canned_intro",
        "inflated_phrase",
    }
    assert document.blocks[0].text == source
    assert utf16_length("A😀") == 3


def test_markdown_rewrite_tracks_unchanged_repeated_blocks() -> None:
    source = "Repeated sentence.\nRepeated sentence.\n"
    original = parse_markdown(source, title="Repeated")
    transformed = original.replace_blocks({"b1": "A focused replacement."})

    assert rewrite_markdown(source, original, transformed) == (
        "Repeated sentence.\nA focused replacement.\n"
    )


def test_copy_styles_use_post_edit_utf16_ranges_for_two_changed_paragraphs() -> None:
    original = DocumentIR(
        title="Emoji review",
        source_format="google_doc",
        blocks=(
            TextBlock(
                id="b0",
                kind=BlockKind.PARAGRAPH,
                text="😀 alpha",
                source_range=SourceRange(1, 9, DocsTabId("tab-a")),
            ),
            TextBlock(
                id="b1",
                kind=BlockKind.PARAGRAPH,
                text="second paragraph",
                source_range=SourceRange(10, 26, DocsTabId("tab-a")),
            ),
        ),
    )
    transformed = original.replace_blocks({"b0": "😀 expanded alpha", "b1": "second"})

    requests = build_copy_mutations(
        original,
        transformed,
        profile=SarjBrandProfile(),
        headers=(
            TabHeader(
                tab_id=DocsTabId("tab-a"),
                header_id=DocsSegmentId("header-a"),
            ),
        ),
        logo_uri="https://assets.example.test/sarj-logo.png",
    )
    style_ranges = [
        request["updateTextStyle"]["range"]
        for request in requests
        if "updateTextStyle" in request and "segmentId" not in request["updateTextStyle"]["range"]
    ]

    assert style_ranges == [
        {"startIndex": 1, "endIndex": 18, "tabId": "tab-a"},
        {"startIndex": 19, "endIndex": 25, "tabId": "tab-a"},
    ]
    assert all("insertPageNumber" not in request for request in requests)


def test_copy_branding_skips_locked_rich_text_blocks() -> None:
    locked = TextBlock(
        id="b0",
        kind=BlockKind.PARAGRAPH,
        text="Linked or suggested text",
        source_range=SourceRange(1, 25, DocsTabId("tab-a")),
        editable=False,
    )
    document = DocumentIR(
        title="Protected",
        blocks=(locked,),
        source_format="google_doc",
        tab_ids=(DocsTabId("tab-a"),),
    )

    requests = build_copy_mutations(
        document,
        document,
        profile=SarjBrandProfile(),
        headers=(),
        logo_uri="",
    )

    assert all("updateTextStyle" not in request for request in requests)
    assert all("updateParagraphStyle" not in request for request in requests)


@pytest.mark.parametrize("level", [4, 5, 6])
def test_copy_branding_preserves_deeper_heading_levels(level: int) -> None:
    heading = TextBlock(
        id="b0",
        kind=BlockKind.HEADING,
        text="Deep heading",
        level=level,
        source_range=SourceRange(1, 13, DocsTabId("tab-a")),
    )
    document = DocumentIR(
        title="Hierarchy",
        blocks=(heading,),
        source_format="google_doc",
        tab_ids=(DocsTabId("tab-a"),),
    )

    requests = build_copy_mutations(
        document,
        document,
        profile=SarjBrandProfile(),
        headers=(),
        logo_uri="",
    )
    paragraph_styles = [
        request["updateParagraphStyle"]["paragraphStyle"]
        for request in requests
        if "updateParagraphStyle" in request
    ]

    assert paragraph_styles[0]["namedStyleType"] == f"HEADING_{level}"
    assert paragraph_styles[0]["keepWithNext"] is True


@pytest.mark.parametrize(
    ("field", "invalid_value"),
    [
        ("spaceAbove", {"magnitude": 0, "unit": "PT"}),
        ("spaceBelow", {"magnitude": 0, "unit": "PT"}),
        ("keepWithNext", True),
    ],
)
def test_post_write_brand_verification_requires_paragraph_spacing(
    field: str,
    invalid_value: object,
) -> None:
    document = _branded_google_document()
    parsed = parse_google_document(document, title="Decision")
    assert verify_brand_application(document, parsed) is True
    changed = deepcopy(document)
    paragraph_style = changed["tabs"][0]["documentTab"]["body"]["content"][0]["paragraph"][
        "paragraphStyle"
    ]
    paragraph_style[field] = invalid_value

    assert (
        verify_brand_application(changed, parse_google_document(changed, title="Decision")) is False
    )


def test_post_write_brand_verification_requires_canonical_logo_source() -> None:
    document = _branded_google_document()
    document["tabs"][0]["documentTab"]["inlineObjects"]["logo-1"]["inlineObjectProperties"][
        "embeddedObject"
    ]["imageProperties"]["sourceUri"] = "https://example.test/wrong.png"

    assert (
        verify_brand_application(document, parse_google_document(document, title="Decision"))
        is False
    )


def test_brand_verification_diagnostics_are_bounded_and_granular() -> None:
    document = _branded_google_document()
    parsed = parse_google_document(document, title="Decision")

    checks = google_gateway_module._brand_application_diagnostics(document, parsed)

    assert checks == {
        "tabs_present": True,
        "tab_count": 1,
        "margins": True,
        "header_shape": True,
        "header_paragraph_style": True,
        "header_title": True,
        "header_text_style": True,
        "logo_count": True,
        "logo_asset": True,
        "paragraph_ranges": True,
        "paragraph_named_styles": True,
        "paragraph_layout": True,
        "paragraph_text_styles": True,
        "paragraph_count": 1,
    }
    changed = deepcopy(document)
    changed["tabs"][0]["documentTab"]["body"]["content"][0]["paragraph"]["elements"][0]["textRun"][
        "textStyle"
    ]["weightedFontFamily"]["fontFamily"] = "Arial"

    changed_checks = google_gateway_module._brand_application_diagnostics(
        changed,
        parse_google_document(changed, title="Decision"),
    )

    assert changed_checks["paragraph_text_styles"] is False
    assert set(changed_checks) == set(checks)


def test_in_place_brand_verification_reports_native_list_typography_as_preserved() -> None:
    document = _branded_google_document()
    paragraph = document["tabs"][0]["documentTab"]["body"]["content"][0]["paragraph"]
    paragraph["bullet"] = {"listId": "list", "nestingLevel": 0}
    document["tabs"][0]["documentTab"]["lists"] = {
        "list": {"listProperties": {"nestingLevels": [{}]}}
    }
    paragraph["elements"][0]["textRun"]["textStyle"]["weightedFontFamily"] = {"fontFamily": "Arial"}
    paragraph["paragraphStyle"]["spaceAbove"] = {"magnitude": 0, "unit": "PT"}
    paragraph["paragraphStyle"]["spaceBelow"] = {"magnitude": 2, "unit": "PT"}
    parsed = parse_google_document(document, title="Decision")

    assert not verify_brand_application(document, parsed)
    assert verify_brand_application(
        document,
        parsed,
        preserved_native_list_text_style_ids=frozenset({"b0"}),
    )
    checks = google_gateway_module._brand_application_diagnostics(
        document,
        parsed,
        preserved_native_list_text_style_ids=frozenset({"b0"}),
    )
    assert checks["paragraph_layout"] is True
    assert checks["paragraph_text_styles"] is True


async def test_in_place_native_list_preservation_keeps_protected_hash_strict() -> None:
    document = _branded_google_document()
    document["documentId"] = "document"
    paragraph = document["tabs"][0]["documentTab"]["body"]["content"][0]["paragraph"]
    paragraph["bullet"] = {"listId": "list", "nestingLevel": 0}
    document["tabs"][0]["documentTab"]["lists"] = {
        "list": {"listProperties": {"nestingLevels": [{}]}}
    }
    paragraph["elements"][0]["textRun"]["textStyle"]["weightedFontFamily"] = {"fontFamily": "Arial"}
    paragraph["paragraphStyle"]["spaceAbove"] = {"magnitude": 0, "unit": "PT"}
    paragraph["paragraphStyle"]["spaceBelow"] = {"magnitude": 2, "unit": "PT"}
    paragraph["elements"] = [
        {
            "startIndex": 1,
            "endIndex": 10,
            "textRun": {
                "content": "Approved.",
                "textStyle": deepcopy(paragraph["elements"][0]["textRun"]["textStyle"]),
            },
        },
        {
            "startIndex": 10,
            "endIndex": 11,
            "textRun": {
                "content": "\n",
                "textStyle": {
                    "bold": True,
                    "weightedFontFamily": {"fontFamily": "Arial", "weight": 400},
                },
            },
        },
    ]
    gateway = object.__new__(GoogleDocumentGateway)
    gateway._logo_uri = "https://assets.example.test/sarj-logo.png"
    snapshot = SourceSnapshot(
        title="Decision",
        parent_folder_id=FileId("folder"),
        modified_time="now",
        revision_id=DocsRevisionId("revision-2"),
        media_type=MIME_DOC,
        google_document=document,
        owned_by_me=True,
        can_edit=True,
        is_personal_drive=True,
    )
    inspection = await gateway.inspect_in_place(snapshot=snapshot, title="Decision")
    assert inspection.supported_brand_applied
    assert not inspection.fully_branded
    expected = parse_google_document(document, title="Decision")

    messages: list[str] = []
    sink = logger.add(messages.append, format="{message}", level="WARNING")
    try:
        result = await gateway.verify_in_place(
            document=document,
            title="Decision",
            expected=expected,
            structural_sha256=inspection.structural_sha256,
            protected_structure_sha256=inspection.protected_structure_sha256,
            verify_brand=True,
        )
    finally:
        logger.remove(sink)

    assert result.content_preserved
    assert result.structure_preserved
    assert result.supported_brand_applied
    assert not any("in-place verification failed" in message for message in messages)
    terminal_normalized = deepcopy(document)
    terminal_normalized["tabs"][0]["documentTab"]["body"]["content"][0]["paragraph"]["elements"][
        -1
    ]["textRun"]["textStyle"] = {"weightedFontFamily": {"fontFamily": "Nunito", "weight": 400}}
    assert (
        await gateway.verify_in_place(
            document=terminal_normalized,
            title="Decision",
            expected=expected,
            structural_sha256=inspection.structural_sha256,
            protected_structure_sha256=inspection.protected_structure_sha256,
            verify_brand=True,
        )
    ).supported_brand_applied
    terminal_linked = deepcopy(terminal_normalized)
    terminal_linked["tabs"][0]["documentTab"]["body"]["content"][0]["paragraph"]["elements"][-1][
        "textRun"
    ]["textStyle"]["link"] = {"url": "https://example.test/changed"}
    assert not (
        await gateway.verify_in_place(
            document=terminal_linked,
            title="Decision",
            expected=expected,
            structural_sha256=inspection.structural_sha256,
            protected_structure_sha256=inspection.protected_structure_sha256,
            verify_brand=True,
        )
    ).structure_preserved
    typography_tampered = deepcopy(document)
    typography_tampered["tabs"][0]["documentTab"]["body"]["content"][0]["paragraph"]["elements"][0][
        "textRun"
    ]["textStyle"]["weightedFontFamily"] = {"fontFamily": "Comic Sans MS"}
    typography_rejected = await gateway.verify_in_place(
        document=typography_tampered,
        title="Decision",
        expected=expected,
        structural_sha256=inspection.structural_sha256,
        protected_structure_sha256=inspection.protected_structure_sha256,
        verify_brand=True,
    )
    assert not typography_rejected.structure_preserved
    tampered = deepcopy(document)
    tampered["tabs"][0]["documentTab"]["body"]["content"][0]["paragraph"]["bullet"]["listId"] = (
        "other"
    )
    rejected = await gateway.verify_in_place(
        document=tampered,
        title="Decision",
        expected=expected,
        structural_sha256=inspection.structural_sha256,
        protected_structure_sha256=inspection.protected_structure_sha256,
        verify_brand=True,
    )
    assert not rejected.structure_preserved


async def test_in_place_non_brand_verification_preserves_brand_short_circuit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    document = _in_place_document("Decision")
    parsed = parse_google_document(document, title="Decision")

    def unexpected_brand_verification(_document: object, _parsed: object) -> bool:
        msg = "brand verification must stay short-circuited"
        raise AssertionError(msg)

    monkeypatch.setattr(
        google_gateway_module,
        "verify_brand_application",
        unexpected_brand_verification,
    )

    result = await object.__new__(GoogleDocumentGateway).verify_in_place(
        document=document,
        title="Decision",
        expected=parsed,
        structural_sha256=google_gateway_module._inventory_sha256(
            google_gateway_module._structural_inventory(document)
        ),
        protected_structure_sha256=google_gateway_module._inventory_sha256(
            google_gateway_module._protected_structure_inventory(document)
        ),
        verify_brand=False,
    )

    assert result.content_preserved is True
    assert result.structure_preserved is True
    assert result.brand_applied is True


def test_copy_branding_styles_header_deterministically() -> None:
    document = DocumentIR(
        title="Decision",
        blocks=(),
        source_format="google_doc",
        tab_ids=(DocsTabId("tab-a"),),
    )

    requests = build_copy_mutations(
        document,
        document,
        profile=SarjBrandProfile(),
        headers=(TabHeader(tab_id=DocsTabId("tab-a"), header_id=DocsSegmentId("header-a")),),
        logo_uri="https://wiki.sarj.ai/brand/sarj-logo-light.png",
    )

    header_text_style = next(
        request["updateTextStyle"]
        for request in requests
        if "updateTextStyle" in request
        and request["updateTextStyle"]["range"].get("segmentId") == "header-a"
    )
    assert header_text_style["range"] == {
        "startIndex": 0,
        "endIndex": 13,
        "segmentId": "header-a",
        "tabId": "tab-a",
    }
    assert header_text_style["textStyle"] == {
        "weightedFontFamily": {"fontFamily": "Nunito"},
        "fontSize": {"magnitude": 9, "unit": "PT"},
        "foregroundColor": {
            "color": {
                "rgbColor": {
                    "red": 115 / 255,
                    "green": 115 / 255,
                    "blue": 115 / 255,
                }
            }
        },
        "bold": True,
    }
    header_paragraph_style = next(
        request["updateParagraphStyle"]
        for request in requests
        if "updateParagraphStyle" in request
        and request["updateParagraphStyle"]["range"].get("segmentId") == "header-a"
    )
    assert header_paragraph_style["range"]["endIndex"] == 14
    assert header_paragraph_style["paragraphStyle"] == {
        "alignment": "END",
        "lineSpacing": 100,
        "spaceAbove": {"magnitude": 0, "unit": "PT"},
        "spaceBelow": {"magnitude": 0, "unit": "PT"},
        "keepWithNext": False,
        "namedStyleType": "NORMAL_TEXT",
    }


def test_copy_branding_normalizes_header_title_and_uses_utf16_indexes() -> None:
    document = DocumentIR(
        title="Decision 😀\nBrief",
        blocks=(),
        source_format="google_doc",
        tab_ids=(DocsTabId("tab-a"),),
    )

    requests = build_copy_mutations(
        document,
        document,
        profile=SarjBrandProfile(),
        headers=(TabHeader(tab_id=DocsTabId("tab-a"), header_id=DocsSegmentId("header-a")),),
        logo_uri="https://wiki.sarj.ai/brand/sarj-logo-light.png",
    )

    header_requests = [
        request
        for request in requests
        if next(iter(request.values())).get("endOfSegmentLocation", {}).get("segmentId")
        == "header-a"
        or next(iter(request.values())).get("range", {}).get("segmentId") == "header-a"
    ]
    assert [next(iter(request)) for request in header_requests] == [
        "insertText",
        "insertInlineImage",
        "updateParagraphStyle",
        "updateTextStyle",
    ]
    assert header_requests[0]["insertText"]["text"] == "Decision 😀 Brief\u00a0\u00a0\u00a0\u00a0\t"
    assert header_requests[3]["updateTextStyle"]["range"]["endIndex"] == 22
    assert header_requests[2]["updateParagraphStyle"]["range"]["endIndex"] == 23


@pytest.mark.parametrize(
    ("field", "invalid_value"),
    [
        ("alignment", "START"),
        ("lineSpacing", 115),
        ("spaceBelow", {"magnitude": 6, "unit": "PT"}),
    ],
)
def test_post_write_brand_verification_requires_header_paragraph_style(
    field: str,
    invalid_value: object,
) -> None:
    document = _branded_google_document()
    changed = deepcopy(document)
    changed["tabs"][0]["documentTab"]["headers"]["header-a"]["content"][0]["paragraph"][
        "paragraphStyle"
    ][field] = invalid_value

    assert verify_brand_application(document, parse_google_document(document, title="Decision"))
    assert not verify_brand_application(changed, parse_google_document(changed, title="Decision"))


def test_post_write_brand_verification_requires_header_text_style_and_title() -> None:
    document = _branded_google_document()
    wrong_style = deepcopy(document)
    text_run = wrong_style["tabs"][0]["documentTab"]["headers"]["header-a"]["content"][0][
        "paragraph"
    ]["elements"][0]["textRun"]
    text_run["textStyle"]["weightedFontFamily"]["fontFamily"] = "Arial"
    wrong_title = deepcopy(document)
    wrong_title["tabs"][0]["documentTab"]["headers"]["header-a"]["content"][0]["paragraph"][
        "elements"
    ][0]["textRun"]["content"] = "Other\t"

    assert not verify_brand_application(
        wrong_style, parse_google_document(wrong_style, title="Decision")
    )
    assert not verify_brand_application(
        wrong_title, parse_google_document(wrong_title, title="Decision")
    )


def test_multi_tab_brand_verification_never_uses_top_level_inline_objects() -> None:
    document = _branded_google_document()
    second = deepcopy(document["tabs"][0])
    second["tabProperties"]["tabId"] = "tab-b"
    second_tab = second["documentTab"]
    second_tab["documentStyle"]["defaultHeaderId"] = "header-b"
    second_tab["headers"] = {"header-b": second_tab["headers"].pop("header-a")}
    second_tab.pop("inlineObjects")
    document["tabs"].append(second)
    document["inlineObjects"] = deepcopy(document["tabs"][0]["documentTab"]["inlineObjects"])

    assert not verify_brand_application(document, parse_google_document(document, title="Decision"))


def test_multi_tab_brand_verification_uses_each_document_tab_inline_objects() -> None:
    document = _branded_google_document()
    second = deepcopy(document["tabs"][0])
    second["tabProperties"]["tabId"] = "tab-b"
    second_tab = second["documentTab"]
    second_tab["documentStyle"]["defaultHeaderId"] = "header-b"
    second_tab["headers"] = {"header-b": second_tab["headers"].pop("header-a")}
    document["tabs"].append(second)
    document["inlineObjects"] = {"logo-1": {"untrusted": "must be ignored for multi-tab docs"}}

    assert verify_brand_application(document, parse_google_document(document, title="Decision"))


def test_post_write_brand_verification_rejects_extra_or_wrong_size_logo() -> None:
    document = _branded_google_document()
    extra_logo = deepcopy(document)
    elements = extra_logo["tabs"][0]["documentTab"]["headers"]["header-a"]["content"][0][
        "paragraph"
    ]["elements"]
    elements.append({"inlineObjectElement": {"inlineObjectId": "logo-2"}})
    wrong_size = deepcopy(document)
    wrong_size["tabs"][0]["documentTab"]["inlineObjects"]["logo-1"]["inlineObjectProperties"][
        "embeddedObject"
    ]["size"]["width"] = {"magnitude": 40, "unit": "PT"}

    assert not verify_brand_application(
        extra_logo, parse_google_document(extra_logo, title="Decision")
    )
    assert not verify_brand_application(
        wrong_size, parse_google_document(wrong_size, title="Decision")
    )


def test_post_write_brand_verification_rejects_table_or_reversed_header() -> None:
    document = _branded_google_document()
    extra_table = deepcopy(document)
    extra_table["tabs"][0]["documentTab"]["headers"]["header-a"]["content"].append(
        {"table": {"rows": 1, "columns": 1}}
    )
    reversed_header = deepcopy(document)
    elements = reversed_header["tabs"][0]["documentTab"]["headers"]["header-a"]["content"][0][
        "paragraph"
    ]["elements"]
    elements[0], elements[1] = elements[1], elements[0]

    assert not verify_brand_application(
        extra_table, parse_google_document(extra_table, title="Decision")
    )
    assert not verify_brand_application(
        reversed_header, parse_google_document(reversed_header, title="Decision")
    )


def test_foreign_pending_mutation_fails_terminal_before_google_write() -> None:
    gateway = object.__new__(GoogleDocumentGateway)
    draft = OutputDraft(
        output=DocumentOutput(
            file_id=FileId("output"),
            name="Output",
            parent_folder_id=FileId("parent"),
        ),
        document_id=FileId("output"),
        tab_id=DocsTabId("tab-a"),
        revision_id=DocsRevisionId("revision-2"),
        source_sha256="source-hash",
        pending_mutation_sha256="v2-mutation",
    )

    with pytest.raises(IncompatiblePendingOutputError):
        gateway._apply_requests_sync(
            draft,
            "v3-mutation",
            [{"insertText": {"text": "unsafe"}}],
            DocsRevisionId("revision-2"),
        )


def _pending_output_draft(
    *,
    pending_mutation_sha256: str,
    mutation_base_revision_sha256: str,
    logo_uri: str | None = None,
) -> OutputDraft:
    return OutputDraft(
        output=DocumentOutput(
            file_id=FileId("output"),
            name="Output",
            parent_folder_id=FileId("parent"),
        ),
        document_id=FileId("output"),
        tab_id=DocsTabId("tab-a"),
        revision_id=DocsRevisionId("revision-2"),
        source_sha256="source-hash",
        logo_uri=logo_uri,
        pending_mutation_sha256=pending_mutation_sha256,
        mutation_base_revision_sha256=mutation_base_revision_sha256,
    )


def test_copy_pending_mutation_does_not_treat_ambient_revision_as_commit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    document = _branded_google_document()
    gateway = object.__new__(GoogleDocumentGateway)
    gateway._logo_uri = "https://assets.example.test/sarj-logo.png"
    writes: list[object] = []
    monkeypatch.setattr(gateway, "_get_document_sync", lambda _document_id: document)
    monkeypatch.setattr(
        gateway,
        "_update_properties_sync",
        lambda _document_id, properties: writes.append(properties),
    )

    def reject_duplicate_write(_request: object) -> dict[str, Any]:
        pytest.fail("a committed copy mutation must not be written twice")

    monkeypatch.setattr(gateway, "_execute_mutation_dict", reject_duplicate_write)
    draft = _pending_output_draft(
        logo_uri=gateway._logo_uri,
        pending_mutation_sha256="v3-mutation",
        mutation_base_revision_sha256=google_gateway_module._revision_sha256(
            DocsRevisionId("revision-1")
        ),
    )

    prepared = gateway._prepare_brand_shell_sync(uuid.uuid7(), draft)
    with pytest.raises(SourceChangedError):
        gateway._apply_requests_sync(
            prepared,
            "v3-mutation",
            [{"insertText": {"text": "not known to be committed"}}],
            DocsRevisionId("revision-1"),
        )

    assert len(prepared.headers) == 1
    assert prepared.brand_warnings == ()
    assert writes == []


def test_copy_pending_mutation_recovers_only_with_exact_atomic_marker(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    document = _branded_google_document()
    digest = "a" * 64
    marker_name = f"sarj-transform-copy-{digest}"
    marker_range = google_gateway_module._terminal_body_marker_range(document)
    assert marker_range is not None
    document["tabs"][0]["documentTab"]["namedRanges"] = {
        marker_name: {
            "namedRanges": [
                {"name": marker_name, "ranges": [marker_range]},
            ]
        }
    }
    gateway = object.__new__(GoogleDocumentGateway)
    gateway._logo_uri = "https://assets.example.test/sarj-logo.png"
    writes: list[object] = []
    monkeypatch.setattr(gateway, "_get_document_sync", lambda _document_id: document)
    monkeypatch.setattr(
        gateway,
        "_update_properties_sync",
        lambda _document_id, properties: writes.append(properties),
    )
    monkeypatch.setattr(
        gateway,
        "_execute_mutation_dict",
        lambda _request: pytest.fail("exact marker recovery must not write twice"),
    )
    draft = _pending_output_draft(
        pending_mutation_sha256=digest,
        mutation_base_revision_sha256=google_gateway_module._revision_sha256(
            DocsRevisionId("revision-1")
        ),
    )

    applied = gateway._apply_requests_sync(
        draft,
        digest,
        [{"insertText": {"text": "already committed"}}],
        DocsRevisionId("revision-2"),
    )

    assert applied.applied_mutation_sha256 == digest
    assert applied.revision_id == "revision-2"
    assert writes == [{google_gateway_module._PROP_MUTATION: digest}]


def test_failed_ambiguous_create_preserves_terminal_conversion_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = object.__new__(GoogleDocumentGateway)
    original = DocumentGatewayError(retryable=False, http_status=400)
    monkeypatch.setattr(gateway, "_recover_output_sync", lambda _operation_id: None)

    with pytest.raises(DocumentGatewayError) as raised:
        gateway._recover_ambiguous_create(
            uuid.uuid7(),
            "source-hash",
            expected_parent_folder_id=FileId("folder"),
            original_error=original,
        )

    assert raised.value is original
    assert not raised.value.retryable


def test_ambiguous_create_rejects_recovery_in_wrong_destination(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = object.__new__(GoogleDocumentGateway)
    recovered = _pending_output_draft(
        pending_mutation_sha256="mutation",
        mutation_base_revision_sha256="revision",
    )
    monkeypatch.setattr(gateway, "_recover_output_sync", lambda _operation_id: recovered)

    with pytest.raises(DocumentGatewayError):
        gateway._recover_ambiguous_create(
            uuid.uuid7(),
            "source-hash",
            expected_parent_folder_id=FileId("different-parent"),
            original_error=DocumentGatewayError(retryable=True),
        )


@pytest.mark.parametrize(
    ("mime_type", "can_add_children"),
    [
        ("application/vnd.google-apps.document", True),
        ("application/vnd.google-apps.folder", False),
    ],
)
async def test_markdown_destination_is_revalidated_before_google_create(
    monkeypatch: pytest.MonkeyPatch,
    mime_type: str,
    can_add_children: bool,
) -> None:
    class DestinationFiles:
        async def get(self, file_id: FileId) -> FileResponse:
            return FileResponse(
                id=file_id,
                name="Destination",
                mime_type=mime_type,
                parent_folder_id=None,
                drive_id=None,
                modified_time=None,
                can_add_children=can_add_children,
            )

    gateway = object.__new__(GoogleDocumentGateway)
    monkeypatch.setattr(gateway, "_files", DestinationFiles(), raising=False)

    def reject_create(*_args: object) -> OutputDraft:
        pytest.fail("invalid destination must fail before Google Drive create")

    monkeypatch.setattr(gateway, "_create_sync", reject_create)
    source = MarkdownSource(
        kind="markdown",
        content="Draft",
        name="Decision",
        destination_folder_id=FileId("folder"),
    )

    with pytest.raises(UnsupportedSourceError):
        await gateway.create_from_markdown(
            operation_id=uuid.uuid7(),
            source=source,
            source_sha256="source-hash",
            output_name="Decision — Sarj formatted",
        )


async def test_drive_markdown_parent_move_fails_source_fence(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class MovedFiles:
        async def get(self, file_id: FileId) -> FileResponse:
            return FileResponse(
                id=file_id,
                name="Decision.md",
                mime_type="text/markdown",
                parent_folder_id=FileId("new-parent"),
                drive_id=None,
                modified_time="2026-09-03T00:00:00Z",
            )

    gateway = object.__new__(GoogleDocumentGateway)
    monkeypatch.setattr(gateway, "_files", MovedFiles(), raising=False)
    unchanged = await gateway.source_is_unchanged(
        DriveFileSource(
            kind="drive_file",
            file_id=FileId("source"),
            expected_modified_time="2026-09-03T00:00:00Z",
        ),
        expected_modified_time="2026-09-03T00:00:00Z",
        expected_revision_id=None,
        expected_parent_folder_id=FileId("old-parent"),
    )

    assert unchanged is False


async def test_source_permission_loss_is_terminal_gateway_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class MissingFiles:
        async def get(self, file_id: FileId) -> FileResponse:
            del file_id
            msg = "not visible"
            raise FileNotInScopeError(msg)

    gateway = object.__new__(GoogleDocumentGateway)
    monkeypatch.setattr(gateway, "_files", MissingFiles(), raising=False)
    with pytest.raises(DocumentGatewayError) as raised:
        await gateway.source_is_unchanged(
            DriveFileSource(
                kind="drive_file",
                file_id=FileId("source"),
                expected_modified_time="2026-09-03T00:00:00Z",
            ),
            expected_modified_time="2026-09-03T00:00:00Z",
            expected_revision_id=None,
            expected_parent_folder_id=FileId("parent"),
        )

    assert raised.value.http_status == 404
    assert raised.value.retryable is False


def test_single_tab_brand_verification_supports_legacy_top_level_inline_objects() -> None:
    document = _branded_google_document()
    document["inlineObjects"] = document["tabs"][0]["documentTab"].pop("inlineObjects")

    assert verify_brand_application(document, parse_google_document(document, title="Decision"))


def _branded_google_document() -> dict[str, Any]:
    return {
        "revisionId": "revision-2",
        "tabs": [
            {
                "tabProperties": {"tabId": "tab-a"},
                "documentTab": {
                    "namedStyles": {
                        "styles": [
                            {
                                "namedStyleType": "NORMAL_TEXT",
                                "textStyle": {
                                    "fontSize": {"magnitude": 11, "unit": "PT"},
                                    "foregroundColor": {
                                        "color": {
                                            "rgbColor": {
                                                "red": 10 / 255,
                                                "green": 10 / 255,
                                                "blue": 10 / 255,
                                            }
                                        }
                                    },
                                },
                            }
                        ]
                    },
                    "inlineObjects": {
                        "logo-1": {
                            "inlineObjectProperties": {
                                "embeddedObject": {
                                    "imageProperties": {
                                        "sourceUri": (
                                            "https://wiki.sarj.ai/brand/sarj-logo-light.png"
                                        )
                                    },
                                    "size": {
                                        "width": {"magnitude": 52, "unit": "PT"},
                                        "height": {"magnitude": 28, "unit": "PT"},
                                    },
                                }
                            }
                        }
                    },
                    "documentStyle": {
                        "marginLeft": {"magnitude": 72, "unit": "PT"},
                        "marginRight": {"magnitude": 72, "unit": "PT"},
                        "marginTop": {"magnitude": 72, "unit": "PT"},
                        "marginBottom": {"magnitude": 72, "unit": "PT"},
                        "defaultHeaderId": "header-a",
                    },
                    "headers": {
                        "header-a": {
                            "content": [
                                {
                                    "paragraph": {
                                        "paragraphStyle": {
                                            "namedStyleType": "NORMAL_TEXT",
                                            "alignment": "END",
                                            "lineSpacing": 100,
                                            "spaceAbove": {"magnitude": 0, "unit": "PT"},
                                            "spaceBelow": {"magnitude": 0, "unit": "PT"},
                                            "keepWithNext": False,
                                        },
                                        "elements": [
                                            {
                                                "textRun": {
                                                    "content": "Decision\u00a0\u00a0\u00a0\u00a0\t",
                                                    "textStyle": {
                                                        "weightedFontFamily": {
                                                            "fontFamily": "Nunito"
                                                        },
                                                        "fontSize": {
                                                            "magnitude": 9,
                                                            "unit": "PT",
                                                        },
                                                        "foregroundColor": {
                                                            "color": {
                                                                "rgbColor": {
                                                                    "red": 115 / 255,
                                                                    "green": 115 / 255,
                                                                    "blue": 115 / 255,
                                                                }
                                                            }
                                                        },
                                                        "bold": True,
                                                    },
                                                }
                                            },
                                            {"inlineObjectElement": {"inlineObjectId": "logo-1"}},
                                        ],
                                    }
                                }
                            ]
                        }
                    },
                    "body": {
                        "content": [
                            {
                                "startIndex": 1,
                                "endIndex": 11,
                                "paragraph": {
                                    "paragraphStyle": {
                                        "namedStyleType": "NORMAL_TEXT",
                                        "alignment": "START",
                                        "lineSpacing": 115,
                                        "spaceAbove": {"magnitude": 2, "unit": "PT"},
                                        "spaceBelow": {"magnitude": 6, "unit": "PT"},
                                        "keepWithNext": False,
                                    },
                                    "elements": [
                                        {
                                            "startIndex": 1,
                                            "endIndex": 11,
                                            "textRun": {
                                                "content": "Approved.\n",
                                                "textStyle": {
                                                    "weightedFontFamily": {"fontFamily": "Nunito"},
                                                    "fontSize": {
                                                        "magnitude": 11,
                                                        "unit": "PT",
                                                    },
                                                    "foregroundColor": {
                                                        "color": {
                                                            "rgbColor": {
                                                                "red": 10 / 255,
                                                                "green": 10 / 255,
                                                                "blue": 10 / 255,
                                                            }
                                                        }
                                                    },
                                                },
                                            },
                                        }
                                    ],
                                },
                            }
                        ]
                    },
                },
                "childTabs": [],
            }
        ],
    }


def _in_place_document(text: str, *, revision: str = "revision-1") -> dict[str, Any]:
    content = f"{text}\n"
    return {
        "documentId": "source",
        "revisionId": revision,
        "tabs": [
            {
                "tabProperties": {"tabId": "tab-1"},
                "documentTab": {
                    "body": {
                        "content": [
                            {
                                "startIndex": 1,
                                "endIndex": 1 + utf16_length(content),
                                "paragraph": {
                                    "paragraphStyle": {},
                                    "elements": [{"textRun": {"content": content}}],
                                },
                            }
                        ]
                    }
                },
                "childTabs": [],
            }
        ],
    }


class _CheckpointStore:
    def __init__(self, events: list[str]) -> None:
        self.events = events
        self.generation = 0
        self.checkpoints: list[InPlaceCheckpoint] = []
        self.resources: list[FileId] = []

    async def save_document_transform_checkpoint(
        self,
        task_id: TaskId,
        *,
        lease_token: uuid.UUID,
        expected_generation: int,
        checkpoint: DocumentTransformCheckpoint,
    ) -> int:
        del task_id, lease_token
        assert expected_generation == self.generation
        assert isinstance(checkpoint, InPlaceCheckpoint)
        self.events.append(f"checkpoint:{checkpoint.stage}")
        self.checkpoints.append(checkpoint)
        self.generation += 1
        return self.generation

    async def record_resource(
        self,
        task_id: TaskId,
        *,
        lease_token: uuid.UUID,
        resource_id: FileId,
    ) -> None:
        del task_id, lease_token
        self.resources.append(resource_id)


class _CopyCheckpointStore:
    def __init__(self) -> None:
        self.generation = 0
        self.checkpoints: list[CopyCheckpoint] = []
        self.resources: list[FileId] = []
        self.events: list[str] = []

    async def save_document_transform_checkpoint(
        self,
        task_id: TaskId,
        *,
        lease_token: uuid.UUID,
        expected_generation: int,
        checkpoint: DocumentTransformCheckpoint,
    ) -> int:
        del task_id, lease_token
        assert expected_generation == self.generation
        assert isinstance(checkpoint, CopyCheckpoint)
        self.checkpoints.append(checkpoint)
        self.events.append(f"checkpoint:{checkpoint.stage.value}")
        self.generation += 1
        return self.generation

    async def record_resource(
        self,
        task_id: TaskId,
        *,
        lease_token: uuid.UUID,
        resource_id: FileId,
    ) -> None:
        del task_id, lease_token
        self.resources.append(resource_id)
        self.events.append(f"resource:{resource_id}")


def _copy_executor(
    *,
    gateway: _Gateway,
    store: _CopyCheckpointStore,
    task_id: TaskId,
    checkpoint: CopyCheckpoint | None = None,
    qa_destination_permission_verifier: executor_module.QaDestinationPermissionVerifier
    | None = None,
    qa_created_copy_verifier: executor_module.QaCreatedCopyVerifier | None = None,
) -> DocumentTransformExecutor:
    return DocumentTransformExecutor(
        gateway=gateway,
        generator=_Generator(),
        checkpoint_store=store,
        task_id=task_id,
        lease_token=uuid.uuid7(),
        checkpoint=checkpoint,
        checkpoint_generation=store.generation,
        qa_destination_permission_verifier=qa_destination_permission_verifier,
        qa_created_copy_verifier=qa_created_copy_verifier,
    )


def _native_copy_gateway() -> _Gateway:
    return _Gateway(
        SourceSnapshot(
            title="Decision",
            parent_folder_id=FileId("source-folder"),
            modified_time="2026-09-03T00:00:00Z",
            revision_id=DocsRevisionId("source-revision-1"),
            media_type="application/vnd.google-apps.document",
            google_document=_in_place_document(
                "The team approved the plan.", revision="source-revision-1"
            ),
        )
    )


def _native_copy_request(*, destination: FileId | None = None) -> DocumentTransformRequest:
    return DocumentTransformRequest(
        source=DriveFileSource(
            kind="drive_file",
            file_id=FileId("source"),
            expected_modified_time="2026-09-03T00:00:00Z",
        ),
        mode=TransformMode.UNSLOP,
        copy_destination_folder_id=destination,
    )


def _qa_native_copy_request(
    *,
    destination: FileId | None = None,
    binding: QaDestinationPermissionBinding | None = None,
) -> QaDocumentTransformRequest:
    return QaDocumentTransformRequest(
        source=DriveFileSource(
            kind="drive_file",
            file_id=FileId("source"),
            expected_modified_time="2026-09-03T00:00:00Z",
        ),
        mode=TransformMode.UNSLOP,
        copy_destination_folder_id=destination or FileId("qa-folder"),
        qa_destination_permission=binding
        or QaDestinationPermissionBinding(version="17", sha256="a" * 64),
    )


def _qa_permission_evidence(
    *,
    output_file_id: FileId,
    destination_folder_id: FileId,
    operation_id: uuid.UUID,
    source_sha256: str,
    output_revision_id: DocsRevisionId,
    output_version: str = "18",
) -> QaCopyPermissionEvidence:
    observed_at = datetime.now(UTC)
    return QaCopyPermissionEvidence(
        operation_id=operation_id,
        source_sha256=source_sha256,
        output_revision_id=output_revision_id,
        destination=QaFilePermissionEvidence(
            file_id=destination_folder_id,
            mime_type="application/vnd.google-apps.folder",
            parent_folder_id=None,
            version="17",
            permission_sha256="a" * 64,
            observed_at=observed_at,
        ),
        output=QaFilePermissionEvidence(
            file_id=output_file_id,
            mime_type="application/vnd.google-apps.document",
            parent_folder_id=destination_folder_id,
            version=output_version,
            permission_sha256="b" * 64,
            observed_at=observed_at,
        ),
    )


def test_public_transform_request_rejects_internal_qa_permission_binding() -> None:
    binding = QaDestinationPermissionBinding(version="17", sha256="a" * 64)

    with pytest.raises(ValidationError, match="Extra inputs are not permitted"):
        DocumentTransformRequest.model_validate(
            {
                "source": {
                    "kind": "drive_file",
                    "file_id": "source",
                    "expected_modified_time": "2026-09-03T00:00:00Z",
                },
                "mode": "unslop",
                "output_strategy": "in_place",
                "qa_destination_permission": binding.model_dump(mode="json"),
            }
        )


async def test_native_copy_uses_and_checkpoints_explicit_destination() -> None:
    gateway = _native_copy_gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())

    result = await _copy_executor(gateway=gateway, store=store, task_id=task_id).execute(
        operation_id=task_id,
        request=_native_copy_request(destination=FileId("qa-folder")),
    )

    assert result.output.parent_folder_id == "qa-folder"
    assert {checkpoint.output_parent_folder_id for checkpoint in store.checkpoints} == {"qa-folder"}


async def test_qa_copy_revalidates_bound_destination_immediately_before_create() -> None:
    gateway = _native_copy_gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    binding = QaDestinationPermissionBinding(version="17", sha256="a" * 64)
    request = _qa_native_copy_request(binding=binding)
    verifier_calls: list[tuple[FileId, QaDestinationPermissionBinding]] = []

    async def reject_changed_access(
        destination_folder_id: FileId,
        binding: QaDestinationPermissionBinding,
    ) -> None:
        verifier_calls.append((destination_folder_id, binding))
        raise DestinationAccessChangedError

    executor = DocumentTransformExecutor(
        gateway=gateway,
        generator=_Generator(),
        checkpoint_store=store,
        task_id=task_id,
        lease_token=uuid.uuid7(),
        qa_destination_permission_verifier=reject_changed_access,
    )

    with pytest.raises(DestinationAccessChangedError):
        await executor.execute(operation_id=task_id, request=request)

    assert verifier_calls == [(FileId("qa-folder"), binding)]
    assert gateway.draft is None
    assert [checkpoint.stage for checkpoint in store.checkpoints] == [CopyCheckpointStage.PLANNED]
    assert store.checkpoints[0].qa_destination_permission == binding


async def test_qa_copy_verifies_created_output_before_continuing() -> None:
    gateway = _native_copy_gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    binding = QaDestinationPermissionBinding(version="17", sha256="a" * 64)
    request = _qa_native_copy_request(binding=binding)
    postcheck_ids: list[FileId] = []

    async def accept_destination(
        destination_folder_id: FileId,
        binding: QaDestinationPermissionBinding,
    ) -> None:
        del destination_folder_id, binding

    async def reject_created_copy(
        request: executor_module.QaCreatedCopyPermissionRequest,
    ) -> QaCopyPermissionEvidence:
        postcheck_ids.append(request.output_file_id)
        raise DestinationAccessChangedError

    executor = _copy_executor(
        gateway=gateway,
        store=store,
        task_id=task_id,
        qa_destination_permission_verifier=accept_destination,
        qa_created_copy_verifier=reject_created_copy,
    )

    with pytest.raises(DestinationAccessChangedError):
        await executor.execute(operation_id=task_id, request=request)

    assert gateway.draft is not None
    assert postcheck_ids == [gateway.draft.output.file_id]
    assert [checkpoint.stage for checkpoint in store.checkpoints] == [
        CopyCheckpointStage.PLANNED,
        CopyCheckpointStage.CREATING,
        CopyCheckpointStage.CREATED,
        CopyCheckpointStage.QUARANTINED,
    ]
    quarantine = store.checkpoints[-1].qa_quarantine
    assert quarantine is not None
    assert quarantine.status == "quarantined_unverified"
    assert quarantine.operation_id == task_id
    assert quarantine.output_file_id == gateway.draft.output.file_id
    assert quarantine.destination_folder_id == "qa-folder"
    assert quarantine.manual_recovery_required


async def test_qa_copy_quarantines_when_final_permission_readback_fails() -> None:
    gateway = _native_copy_gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    binding = QaDestinationPermissionBinding(version="17", sha256="a" * 64)
    request = _qa_native_copy_request(binding=binding)
    verifier_calls: list[executor_module.QaCreatedCopyPermissionRequest] = []

    async def accept_destination(
        destination_folder_id: FileId,
        binding: QaDestinationPermissionBinding,
    ) -> None:
        del destination_folder_id, binding

    async def fail_final_readback(
        request: executor_module.QaCreatedCopyPermissionRequest,
    ) -> QaCopyPermissionEvidence:
        verifier_calls.append(request)
        if len(verifier_calls) == 2:
            raise DestinationAccessChangedError
        return _qa_permission_evidence(
            output_file_id=request.output_file_id,
            destination_folder_id=request.destination_folder_id,
            operation_id=request.operation_id,
            source_sha256=request.source_sha256,
            output_revision_id=request.output_revision_id,
        )

    with pytest.raises(DestinationAccessChangedError):
        await _copy_executor(
            gateway=gateway,
            store=store,
            task_id=task_id,
            qa_destination_permission_verifier=accept_destination,
            qa_created_copy_verifier=fail_final_readback,
        ).execute(operation_id=task_id, request=request)

    assert len(verifier_calls) == 2
    assert gateway.draft is not None
    assert verifier_calls[-1].output_revision_id == gateway.draft.revision_id
    quarantine = store.checkpoints[-1].qa_quarantine
    assert quarantine is not None
    assert quarantine.status == "quarantined_unverified"
    assert quarantine.output_file_id == gateway.draft.output.file_id
    assert quarantine.output_revision_id == gateway.draft.revision_id
    assert store.checkpoints[-1].stage is CopyCheckpointStage.QUARANTINED
    assert store.checkpoints[-1].result is None
    assert store.checkpoints[-1].qa_permission_evidence is None


async def test_qa_copy_revalidates_recovered_verified_output_before_return() -> None:
    gateway = _native_copy_gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    binding = QaDestinationPermissionBinding(version="17", sha256="a" * 64)
    request = _qa_native_copy_request(binding=binding)

    async def accept_destination(
        destination_folder_id: FileId,
        binding: QaDestinationPermissionBinding,
    ) -> None:
        del destination_folder_id, binding

    async def accept_created_copy(
        request: executor_module.QaCreatedCopyPermissionRequest,
    ) -> QaCopyPermissionEvidence:
        return _qa_permission_evidence(
            output_file_id=request.output_file_id,
            destination_folder_id=request.destination_folder_id,
            operation_id=request.operation_id,
            source_sha256=request.source_sha256,
            output_revision_id=request.output_revision_id,
        )

    result = await _copy_executor(
        gateway=gateway,
        store=store,
        task_id=task_id,
        qa_destination_permission_verifier=accept_destination,
        qa_created_copy_verifier=accept_created_copy,
    ).execute(operation_id=task_id, request=request)
    assert result.access_scope == "owner_only"
    assert isinstance(result, QaDocumentTransformResult)
    evidence = result.qa_permission_evidence
    assert evidence is not None
    assert evidence.operation_id == task_id
    assert evidence.source_sha256 == result.source_sha256
    assert evidence.output.file_id == result.output.file_id
    assert evidence.output.parent_folder_id == result.output.parent_folder_id
    assert evidence.output_revision_id == result.output_revision_id
    assert store.checkpoints[-1].qa_permission_evidence == evidence
    assert store.checkpoints[-1].result == result
    recovered_ids: list[FileId] = []

    async def reject_recovered_copy(
        request: executor_module.QaCreatedCopyPermissionRequest,
    ) -> QaCopyPermissionEvidence:
        recovered_ids.append(request.output_file_id)
        raise DestinationAccessChangedError

    with pytest.raises(DestinationAccessChangedError):
        await _copy_executor(
            gateway=gateway,
            store=store,
            task_id=task_id,
            checkpoint=store.checkpoints[-1],
            qa_destination_permission_verifier=accept_destination,
            qa_created_copy_verifier=reject_recovered_copy,
        ).execute(operation_id=task_id, request=request)

    assert gateway.draft is not None
    assert recovered_ids == [gateway.draft.output.file_id]
    assert store.checkpoints[-1].stage is CopyCheckpointStage.QUARANTINED
    assert store.checkpoints[-1].result is None
    assert store.checkpoints[-1].qa_permission_evidence is None


async def test_qa_copy_attests_state_after_complete_mutates_output_metadata() -> None:
    gateway = _native_copy_gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    request = _qa_native_copy_request()
    observed_complete_states: list[bool] = []

    async def accept_destination(
        destination_folder_id: FileId,
        binding: QaDestinationPermissionBinding,
    ) -> None:
        del destination_folder_id, binding

    async def attest_copy(
        request: executor_module.QaCreatedCopyPermissionRequest,
    ) -> QaCopyPermissionEvidence:
        completed = gateway.completed_result is not None
        observed_complete_states.append(completed)
        return _qa_permission_evidence(
            output_file_id=request.output_file_id,
            destination_folder_id=request.destination_folder_id,
            operation_id=request.operation_id,
            source_sha256=request.source_sha256,
            output_revision_id=request.output_revision_id,
            output_version="19" if completed else "18",
        )

    result = await _copy_executor(
        gateway=gateway,
        store=store,
        task_id=task_id,
        qa_destination_permission_verifier=accept_destination,
        qa_created_copy_verifier=attest_copy,
    ).execute(operation_id=task_id, request=request)

    assert observed_complete_states == [False, True]
    assert isinstance(result, QaDocumentTransformResult)
    assert result.qa_permission_evidence.output.version == "19"
    assert store.checkpoints[-1].stage is CopyCheckpointStage.ATTESTED
    assert store.checkpoints[-1].qa_permission_evidence == result.qa_permission_evidence
    assert gateway.completed_result is not None
    assert "qa_permission_evidence" not in gateway.completed_result


async def test_qa_copy_attestation_orders_docs_acl_docs_barrier(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = _native_copy_gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    events: list[str] = []
    original_verify = gateway.verify_output
    original_complete = gateway.complete

    async def verify_output(
        *,
        draft: OutputDraft,
        source: SourceSnapshot,
        expected: DocumentIR,
        verify_brand: bool,
        header_title: str | None = None,
    ) -> OutputVerification:
        events.append("docs:complete" if gateway.completed_result is not None else "docs:draft")
        return await original_verify(
            draft=draft,
            source=source,
            expected=expected,
            verify_brand=verify_brand,
            header_title=header_title,
        )

    async def complete(
        *, operation_id: uuid.UUID, draft: OutputDraft, result: dict[str, Any]
    ) -> None:
        events.append("complete")
        await original_complete(operation_id=operation_id, draft=draft, result=result)

    async def accept_destination(
        destination_folder_id: FileId,
        binding: QaDestinationPermissionBinding,
    ) -> None:
        del destination_folder_id, binding
        events.append("destination")

    async def attest_copy(
        request: executor_module.QaCreatedCopyPermissionRequest,
    ) -> QaCopyPermissionEvidence:
        events.append("acl:terminal" if gateway.completed_result is not None else "acl:created")
        return _qa_permission_evidence(
            output_file_id=request.output_file_id,
            destination_folder_id=request.destination_folder_id,
            operation_id=request.operation_id,
            source_sha256=request.source_sha256,
            output_revision_id=request.output_revision_id,
        )

    monkeypatch.setattr(gateway, "verify_output", verify_output)
    monkeypatch.setattr(gateway, "complete", complete)
    result = await _copy_executor(
        gateway=gateway,
        store=store,
        task_id=task_id,
        qa_destination_permission_verifier=accept_destination,
        qa_created_copy_verifier=attest_copy,
    ).execute(operation_id=task_id, request=_qa_native_copy_request())

    assert isinstance(result, QaDocumentTransformResult)
    assert events == [
        "destination",
        "acl:created",
        "docs:draft",
        "docs:draft",
        "complete",
        "docs:complete",
        "acl:terminal",
        "docs:complete",
    ]
    assert store.checkpoints[-1].stage is CopyCheckpointStage.ATTESTED


async def test_qa_copy_post_acl_revision_drift_is_durably_quarantined(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = _native_copy_gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    verify_calls = 0
    acl_calls = 0
    original_verify = gateway.verify_output

    async def verify_output(
        *,
        draft: OutputDraft,
        source: SourceSnapshot,
        expected: DocumentIR,
        verify_brand: bool,
        header_title: str | None = None,
    ) -> OutputVerification:
        nonlocal verify_calls
        verify_calls += 1
        verified = await original_verify(
            draft=draft,
            source=source,
            expected=expected,
            verify_brand=verify_brand,
            header_title=header_title,
        )
        if verify_calls == 4:
            return replace(verified, revision_id=DocsRevisionId("output-revision-3"))
        return verified

    async def accept_destination(
        destination_folder_id: FileId,
        binding: QaDestinationPermissionBinding,
    ) -> None:
        del destination_folder_id, binding

    async def attest_copy(
        request: executor_module.QaCreatedCopyPermissionRequest,
    ) -> QaCopyPermissionEvidence:
        nonlocal acl_calls
        acl_calls += 1
        return _qa_permission_evidence(
            output_file_id=request.output_file_id,
            destination_folder_id=request.destination_folder_id,
            operation_id=request.operation_id,
            source_sha256=request.source_sha256,
            output_revision_id=request.output_revision_id,
        )

    monkeypatch.setattr(gateway, "verify_output", verify_output)
    with pytest.raises(OutputVerificationError):
        await _copy_executor(
            gateway=gateway,
            store=store,
            task_id=task_id,
            qa_destination_permission_verifier=accept_destination,
            qa_created_copy_verifier=attest_copy,
        ).execute(operation_id=task_id, request=_qa_native_copy_request())

    assert verify_calls == 4
    assert acl_calls == 2
    quarantined = store.checkpoints[-1]
    assert quarantined.stage is CopyCheckpointStage.QUARANTINED
    assert quarantined.output_revision_id == "output-revision-3"
    assert quarantined.result is None
    assert quarantined.qa_permission_evidence is None
    assert quarantined.qa_quarantine is not None
    assert quarantined.qa_quarantine.output_revision_id == "output-revision-3"


async def test_qa_copy_post_acl_transient_docs_failure_retains_verified_retry_state(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = _native_copy_gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    verify_calls = 0
    original_verify = gateway.verify_output

    async def verify_output(
        *,
        draft: OutputDraft,
        source: SourceSnapshot,
        expected: DocumentIR,
        verify_brand: bool,
        header_title: str | None = None,
    ) -> OutputVerification:
        nonlocal verify_calls
        verify_calls += 1
        if verify_calls == 4:
            raise DocumentGatewayError(retryable=True)
        return await original_verify(
            draft=draft,
            source=source,
            expected=expected,
            verify_brand=verify_brand,
            header_title=header_title,
        )

    async def accept_destination(
        destination_folder_id: FileId,
        binding: QaDestinationPermissionBinding,
    ) -> None:
        del destination_folder_id, binding

    async def attest_copy(
        request: executor_module.QaCreatedCopyPermissionRequest,
    ) -> QaCopyPermissionEvidence:
        return _qa_permission_evidence(
            output_file_id=request.output_file_id,
            destination_folder_id=request.destination_folder_id,
            operation_id=request.operation_id,
            source_sha256=request.source_sha256,
            output_revision_id=request.output_revision_id,
        )

    monkeypatch.setattr(gateway, "verify_output", verify_output)
    with pytest.raises(DocumentGatewayError):
        await _copy_executor(
            gateway=gateway,
            store=store,
            task_id=task_id,
            qa_destination_permission_verifier=accept_destination,
            qa_created_copy_verifier=attest_copy,
        ).execute(operation_id=task_id, request=_qa_native_copy_request())

    assert verify_calls == 4
    assert store.checkpoints[-1].stage is CopyCheckpointStage.VERIFIED
    assert store.checkpoints[-1].result is not None
    assert store.checkpoints[-1].qa_permission_evidence is None


async def test_native_copy_rejects_recovered_output_in_wrong_destination() -> None:
    gateway = _native_copy_gateway()
    gateway._create("source-hash", "Recovered", parent_folder_id=FileId("wrong-folder"))

    with pytest.raises(OutputConflictError):
        await DocumentTransformExecutor(gateway=gateway, generator=_Generator()).execute(
            operation_id=uuid.uuid7(),
            request=_native_copy_request(destination=FileId("qa-folder")),
        )


async def test_qa_copy_revalidates_wrong_parent_recovery_before_rejecting_it() -> None:
    gateway = _native_copy_gateway()
    gateway._create("source-hash", "Recovered", parent_folder_id=FileId("wrong-folder"))
    binding = QaDestinationPermissionBinding(version="17", sha256="a" * 64)
    request = _qa_native_copy_request(binding=binding)
    postchecks: list[tuple[FileId, FileId]] = []

    async def reject_untrusted_recovery(
        request: executor_module.QaCreatedCopyPermissionRequest,
    ) -> QaCopyPermissionEvidence:
        postchecks.append((request.output_file_id, request.destination_folder_id))
        raise DestinationAccessChangedError

    with pytest.raises(DestinationAccessChangedError):
        await DocumentTransformExecutor(
            gateway=gateway,
            generator=_Generator(),
            qa_created_copy_verifier=reject_untrusted_recovery,
        ).execute(operation_id=uuid.uuid7(), request=request)

    assert gateway.draft is not None
    assert postchecks == [(gateway.draft.output.file_id, FileId("qa-folder"))]


async def test_legacy_copy_checkpoint_without_output_parent_uses_source_parent() -> None:
    gateway = _native_copy_gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    request = _native_copy_request()
    first = await _copy_executor(gateway=gateway, store=store, task_id=task_id).execute(
        operation_id=task_id,
        request=request,
    )
    legacy = store.checkpoints[-1].model_copy(update={"output_parent_folder_id": None})

    repeated = await _copy_executor(
        gateway=gateway,
        store=store,
        task_id=task_id,
        checkpoint=legacy,
    ).execute(operation_id=task_id, request=request)

    assert repeated == first


@pytest.mark.parametrize(
    ("source_name", "mode", "override", "expected_name"),
    [
        ("Untitled", TransformMode.FORMAT, None, "Capacity Review — Sarj Formatted"),
        ("Untitled", TransformMode.BOTH, None, "Capacity Review — Sarj Polished"),
        ("Untitled", TransformMode.UNSLOP, None, "Untitled — Unslopped"),
        ("Board  Pack", TransformMode.FORMAT, None, "Board Pack — Sarj Formatted"),
        (
            "security review 2026-09-04",
            TransformMode.BOTH,
            None,
            "Security Review — Sarj Polished",
        ),
        ("Untitled", TransformMode.FORMAT, "Approved packet", "Approved Packet"),
        ("Untitled", TransformMode.FORMAT, "approved packet — Sep 4, 2026", "Approved Packet"),
        ("x" * 255, TransformMode.FORMAT, None, "X" + "x" * 237 + " — Sarj Formatted"),
    ],
)
async def test_copy_automatic_name_is_persisted_before_creation_and_retry_stable(
    source_name: str,
    mode: TransformMode,
    override: str | None,
    expected_name: str,
) -> None:
    gateway = _Gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    request = DocumentTransformRequest(
        source=MarkdownSource(
            kind="markdown",
            content="# Capacity review\n\nThe team approved the plan.",
            name=source_name,
            destination_folder_id=FileId("folder"),
        ),
        mode=mode,
        output_name=override,
    )
    result = await _copy_executor(gateway=gateway, store=store, task_id=task_id).execute(
        operation_id=task_id, request=request
    )
    assert result.output.name == expected_name
    assert {checkpoint.output_name for checkpoint in store.checkpoints} == {expected_name}
    assert set(gateway.verified_header_titles) == {expected_name}
    assert store.events.index("checkpoint:planned") < store.events.index("resource:output")
    applies = gateway.applies
    repeated = await _copy_executor(
        gateway=gateway, store=store, task_id=task_id, checkpoint=store.checkpoints[-1]
    ).execute(operation_id=task_id, request=request)
    assert repeated.output == result.output
    assert gateway.applies == applies
    assert store.resources == [result.output.file_id]


async def test_copy_checkpoints_plan_before_create_and_resource_immediately_after() -> None:
    gateway = _Gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    result = await DocumentTransformExecutor(
        gateway=gateway,
        generator=_Generator(),
        checkpoint_store=store,
        task_id=task_id,
        lease_token=uuid.uuid7(),
    ).execute(
        operation_id=task_id,
        request=DocumentTransformRequest(
            source=MarkdownSource(
                kind="markdown",
                content="# Decision\n\nThe team approved the plan.",
                name="Decision",
                destination_folder_id=FileId("folder"),
            ),
            mode=TransformMode.UNSLOP,
        ),
    )

    assert store.events == [
        "checkpoint:planned",
        "checkpoint:creating",
        "resource:output",
        "checkpoint:created",
        "checkpoint:verified",
    ]
    verified = store.checkpoints[-1]
    assert verified.stage is CopyCheckpointStage.VERIFIED
    assert verified.result == result
    assert verified.output_revision_id == result.output_revision_id
    persisted = task_service_module._serialize_checkpoint(verified)
    assert set(persisted) == {"kind", "version", "checkpoint_ciphertext"}
    assert task_service_module._deserialize_checkpoint(persisted) == verified


@pytest.mark.parametrize(
    ("interrupt_population", "previous_builder"),
    [
        (False, None),
        (True, None),
        (True, "google-docs-requests-v7"),
        (True, "google-docs-requests-v8"),
        (True, "google-docs-requests-v9"),
        (True, "google-docs-requests-v15"),
        (True, "google-docs-requests-v19"),
    ],
    ids=[
        "fresh-current",
        "resume-current",
        "reject-pending-v7-before-provider-write",
        "reject-pending-v8-before-provider-write",
        "reject-pending-v9-before-provider-write",
        "reject-pending-v15-before-provider-write",
        "reject-pending-v19-before-provider-write",
    ],
)
async def test_copy_native_footer_persists_each_stage_and_does_not_replay_body(
    interrupt_population: bool,
    previous_builder: str | None,
) -> None:
    store = _CopyCheckpointStore()

    class FooterGateway(_Gateway):
        supports_native_footer = True
        interrupted = False
        footer_dispatches = 0

        @override
        async def plan_copy_footer(
            self, *, draft: OutputDraft, operation_id: uuid.UUID
        ) -> NativeFooterCheckpoint:
            assert self.applies == 1
            return NativeFooterCheckpoint(
                tab_id=draft.tab_id,
                url="https://wiki.sarj.ai/d/output",
                revision_id=draft.revision_id,
                create_marker=f"create-{operation_id}",
                populate_marker=f"populate-{operation_id}",
                marker_range={},
                invariant_sha256="hash",
            )

        @override
        async def apply_copy_footer_stage(
            self, *, document_id: FileId, checkpoint: NativeFooterCheckpoint
        ) -> NativeFooterCheckpoint:
            self.footer_dispatches += 1
            assert document_id == "output"
            assert store.checkpoints[-1].native_footer == checkpoint
            if (
                checkpoint.stage is NativeFooterStage.POPULATE
                and interrupt_population
                and not self.interrupted
            ):
                self.interrupted = True
                raise DocumentGatewayError
            stage = (
                NativeFooterStage.POPULATE
                if checkpoint.stage is NativeFooterStage.CREATE
                else NativeFooterStage.COMPLETE
            )
            result = checkpoint.model_copy(
                update={
                    "stage": stage,
                    "footer_id": DocsSegmentId("footer"),
                    "revision_id": DocsRevisionId(f"footer-{stage.value}"),
                }
            )
            assert self.draft is not None
            self.draft = replace(self.draft, revision_id=result.revision_id)
            return result

    gateway = FooterGateway(
        SourceSnapshot(
            title="Decision",
            parent_folder_id=FileId("folder"),
            modified_time="2026-09-03T00:00:00Z",
            revision_id=DocsRevisionId("source-revision"),
            media_type="application/vnd.google-apps.document",
            google_document=_in_place_document(
                "The team approved the plan.", revision="source-revision"
            ),
        )
    )
    task_id = TaskId(uuid.uuid7())
    request = DocumentTransformRequest(
        source=DriveFileSource(
            kind="drive_file",
            file_id=FileId("source"),
            expected_modified_time="2026-09-03T00:00:00Z",
        ),
        mode=TransformMode.FORMAT,
    )
    executor = _copy_executor(gateway=gateway, store=store, task_id=task_id)
    if interrupt_population:
        with pytest.raises(DocumentGatewayError):
            await executor.execute(operation_id=task_id, request=request)
        pending = store.checkpoints[-1]
        assert pending.request_builder_version == REQUEST_BUILDER_VERSION
        assert pending.native_footer is not None
        assert pending.native_footer.coverage is None
        assert pending.native_footer.stage is NativeFooterStage.POPULATE
        if previous_builder:
            pending = pending.model_copy(update={"request_builder_version": previous_builder})
        executor = _copy_executor(gateway=gateway, store=store, task_id=task_id, checkpoint=pending)
    if previous_builder:
        writes_before = (gateway.applies, gateway.brand_prepares, gateway.footer_dispatches)
        persisted_before = len(store.checkpoints)
        with pytest.raises(OutputConflictError):
            await executor.execute(operation_id=task_id, request=request)
        assert (gateway.applies, gateway.brand_prepares, gateway.footer_dispatches) == writes_before
        assert len(store.checkpoints) == persisted_before
        return
    result = await executor.execute(operation_id=task_id, request=request)
    assert result.output_revision_id == "footer-complete"
    assert gateway.applies == 1
    assert gateway.brand_prepares == 1
    stages = [c.native_footer.stage for c in store.checkpoints if c.native_footer is not None]
    assert stages == [
        NativeFooterStage.CREATE,
        NativeFooterStage.POPULATE,
        NativeFooterStage.COMPLETE,
        NativeFooterStage.COMPLETE,
    ]


async def test_copy_readback_mismatch_fails_with_output_resource_tracked() -> None:
    gateway = _Gateway(
        verification=OutputVerification(content_preserved=False, structure_preserved=False)
    )
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())

    with pytest.raises(OutputVerificationError):
        await DocumentTransformExecutor(
            gateway=gateway,
            generator=_Generator(),
            checkpoint_store=store,
            task_id=task_id,
            lease_token=uuid.uuid7(),
        ).execute(
            operation_id=task_id,
            request=DocumentTransformRequest(
                source=MarkdownSource(
                    kind="markdown",
                    content="Decision",
                    name="Decision",
                    destination_folder_id=FileId("folder"),
                ),
                mode=TransformMode.UNSLOP,
            ),
        )

    assert store.resources == [FileId("output")]
    assert store.checkpoints[-1].stage is CopyCheckpointStage.CREATED
    assert gateway.completed_result is None


@pytest.mark.parametrize("mode", [TransformMode.FORMAT, TransformMode.BOTH, TransformMode.UNSLOP])
async def test_ambiguous_copy_create_checkpoint_never_issues_second_create(
    mode: TransformMode,
) -> None:
    class AmbiguousGateway(_Gateway):
        def __init__(self) -> None:
            super().__init__()
            self.creates = 0

        async def create_from_markdown(
            self,
            *,
            operation_id: uuid.UUID,
            source: MarkdownSource,
            source_sha256: str,
            output_name: str,
            imported_docx: bytes | None = None,
        ) -> OutputDraft:
            del operation_id, source, source_sha256, output_name
            assert imported_docx is not None
            document = WordDocument(BytesIO(imported_docx))
            assert document.paragraphs[0].text == "Authored decision"
            assert document.paragraphs[0].style is not None
            assert document.paragraphs[0].style.name == "Heading 1"
            self.creates += 1
            raise DocumentGatewayError(retryable=True)

    gateway = AmbiguousGateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    request = DocumentTransformRequest(
        source=MarkdownSource(
            kind="markdown",
            content="# Authored decision\n\nBody",
            name="Decision",
            destination_folder_id=FileId("folder"),
        ),
        mode=mode,
    )
    with pytest.raises(DocumentGatewayError):
        await DocumentTransformExecutor(
            gateway=gateway,
            generator=_Generator(),
            checkpoint_store=store,
            task_id=task_id,
            lease_token=uuid.uuid7(),
        ).execute(operation_id=task_id, request=request)
    assert store.checkpoints[-1].stage is CopyCheckpointStage.CREATING

    with pytest.raises(DocumentGatewayError):
        await _copy_executor(
            gateway=gateway,
            store=store,
            task_id=task_id,
            checkpoint=store.checkpoints[-1],
        ).execute(operation_id=task_id, request=request)

    assert gateway.creates == 1
    assert [item.stage for item in store.checkpoints].count(CopyCheckpointStage.PLANNED) == 1


async def test_completed_copy_recovery_rejects_revision_drift() -> None:
    gateway = _Gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    request = DocumentTransformRequest(
        source=MarkdownSource(
            kind="markdown",
            content="Decision",
            name="Decision",
            destination_folder_id=FileId("folder"),
        ),
        mode=TransformMode.UNSLOP,
    )
    await DocumentTransformExecutor(
        gateway=gateway,
        generator=_Generator(),
        checkpoint_store=store,
        task_id=task_id,
        lease_token=uuid.uuid7(),
    ).execute(operation_id=task_id, request=request)
    assert gateway.draft is not None
    gateway.draft = replace(
        gateway.draft,
        revision_id=DocsRevisionId("user-edit-revision"),
    )

    with pytest.raises(OutputVerificationError):
        await _copy_executor(
            gateway=gateway,
            store=store,
            task_id=task_id,
            checkpoint=store.checkpoints[-1],
        ).execute(operation_id=task_id, request=request)


async def test_completed_copy_recovery_rejects_tampered_result_digest() -> None:
    gateway = _Gateway()
    store = _CopyCheckpointStore()
    task_id = TaskId(uuid.uuid7())
    request = DocumentTransformRequest(
        source=MarkdownSource(
            kind="markdown",
            content="Decision",
            name="Decision",
            destination_folder_id=FileId("folder"),
        ),
        mode=TransformMode.UNSLOP,
    )
    await DocumentTransformExecutor(
        gateway=gateway,
        generator=_Generator(),
        checkpoint_store=store,
        task_id=task_id,
        lease_token=uuid.uuid7(),
    ).execute(operation_id=task_id, request=request)
    assert gateway.draft is not None
    gateway.draft = replace(gateway.draft, result_sha256="0" * 64)

    with pytest.raises(OutputConflictError):
        await _copy_executor(
            gateway=gateway,
            store=store,
            task_id=task_id,
            checkpoint=store.checkpoints[-1],
        ).execute(operation_id=task_id, request=request)


class _InPlaceGateway(_Gateway):
    def __init__(self, text: str, *, missing_header: bool = False) -> None:
        document = _in_place_document(text)
        super().__init__(
            SourceSnapshot(
                title="Source",
                parent_folder_id=FileId("folder"),
                modified_time="2026-09-03T00:00:00Z",
                revision_id=DocsRevisionId("revision-1"),
                media_type="application/vnd.google-apps.document",
                google_document=document,
                owned_by_me=True,
                can_edit=True,
                is_personal_drive=True,
                web_url="https://docs.google.com/document/d/source/edit",
            )
        )
        self.document = document
        self.missing_header = missing_header
        self.header_created = False
        self.events: list[str] = []

    @override
    async def inspect_in_place(
        self, *, snapshot: SourceSnapshot, title: str, table_font_plan: TableFontPlan | None = None
    ) -> InPlaceInspection:
        del title, table_font_plan
        headers = (
            ()
            if self.missing_header and not self.header_created
            else (TabHeader(tab_id=DocsTabId("tab-1"), header_id=DocsSegmentId("header-1")),)
        )
        return InPlaceInspection(
            output=DocumentOutput(
                file_id=FileId("source"),
                name="Source",
                parent_folder_id=FileId("folder"),
            ),
            headers=headers,
            missing_header_tab_ids=(DocsTabId("tab-1"),) if not headers else (),
            unknown_header_tab_ids=(),
            marker_range={"startIndex": 20, "endIndex": 21, "tabId": "tab-1"},
            logo_uri="https://assets.example.test/logo.png",
            structural_sha256="structure",
            protected_structure_sha256="protected",
            body_sha256="body",
            header_safe_sha256="header-safe",
        )

    @override
    async def apply_in_place_stage(
        self,
        *,
        document_id: FileId,
        required_revision_id: DocsRevisionId,
        marker_name: str,
        marker_range: dict[str, Any],
        requests: list[dict[str, Any]],
    ) -> InPlaceStageResult:
        del document_id, required_revision_id, marker_name, marker_range
        self.events.append("apply")
        if any("createHeader" in request for request in requests):
            self.header_created = True
            revision = "revision-2"
        else:
            source = parse_google_document(self.document, title="Source")
            replacement = next(
                (
                    request["insertText"]["text"]
                    for request in requests
                    if "insertText" in request
                    and request["insertText"].get("location", {}).get("segmentId") is None
                ),
                source.blocks[0].text,
            )
            self.document = _in_place_document(replacement, revision="revision-3")
            revision = "revision-3"
        self.document["revisionId"] = revision
        return InPlaceStageResult(
            revision_id=DocsRevisionId(revision),
            document=deepcopy(self.document),
        )

    @override
    async def verify_in_place(
        self,
        *,
        document: dict[str, Any],
        title: str,
        expected: DocumentIR,
        structural_sha256: str,
        protected_structure_sha256: str,
        verify_brand: bool,
        custom_furniture: tuple[CustomFurnitureBinding, ...] = (),
        custom_header_titles: tuple[Any, ...] = (),
        table_font_plan: TableFontPlan | None = None,
    ) -> OutputVerification:
        del custom_furniture, custom_header_titles, table_font_plan
        del document, title, expected, structural_sha256, protected_structure_sha256, verify_brand
        return OutputVerification(
            content_preserved=True, structure_preserved=True, brand_applied=True
        )


class _CrashReplayInPlaceGateway(_InPlaceGateway):
    def __init__(self, text: str, *, crash_stage: str) -> None:
        super().__init__(text, missing_header=True)
        self.drive_title = "Source"
        self.crash_stage = crash_stage
        self.crashed = False
        self.mutations = {"headers": 0, "content": 0}
        self.stage_results: dict[str, InPlaceStageResult] = {}

    @override
    async def load_drive_source(self, source: DriveFileSource) -> SourceSnapshot:
        del source
        return SourceSnapshot(
            title=self.drive_title,
            parent_folder_id=FileId("folder"),
            modified_time="2026-09-03T00:00:00Z",
            revision_id=DocsRevisionId(self.document["revisionId"]),
            media_type="application/vnd.google-apps.document",
            google_document=deepcopy(self.document),
            owned_by_me=True,
            can_edit=True,
            is_personal_drive=True,
            web_url="https://docs.google.com/document/d/source/edit",
        )

    @override
    async def apply_in_place_stage(
        self,
        *,
        document_id: FileId,
        required_revision_id: DocsRevisionId,
        marker_name: str,
        marker_range: dict[str, Any],
        requests: list[dict[str, Any]],
    ) -> InPlaceStageResult:
        recovered = self.stage_results.get(marker_name)
        if recovered is not None:
            return InPlaceStageResult(
                revision_id=recovered.revision_id,
                document=deepcopy(recovered.document),
            )
        stage = "headers" if any("createHeader" in request for request in requests) else "content"
        result = await super().apply_in_place_stage(
            document_id=document_id,
            required_revision_id=required_revision_id,
            marker_name=marker_name,
            marker_range=marker_range,
            requests=requests,
        )
        self.mutations[stage] += 1
        self.stage_results[marker_name] = result
        if stage == self.crash_stage and not self.crashed:
            self.crashed = True
            msg = f"simulated worker crash after {stage} batch"
            raise RuntimeError(msg)
        return result


class _PreMutationRevisionDriftGateway(_InPlaceGateway):
    def __init__(self, text: str, *, drift: str = "revision") -> None:
        super().__init__(text, missing_header=True)
        self.drift = drift
        self.failed = False
        self.required_revisions: list[DocsRevisionId] = []
        self.mutations = {"headers": 0, "content": 0}

    @override
    async def load_drive_source(self, source: DriveFileSource) -> SourceSnapshot:
        del source
        return SourceSnapshot(
            title="Source",
            parent_folder_id=FileId("folder"),
            modified_time="2026-09-03T00:00:00Z",
            revision_id=DocsRevisionId(self.document["revisionId"]),
            media_type="application/vnd.google-apps.document",
            google_document=deepcopy(self.document),
            owned_by_me=True,
            can_edit=True,
            is_personal_drive=True,
            web_url="https://docs.google.com/document/d/source/edit",
        )

    @override
    async def inspect_in_place(
        self, *, snapshot: SourceSnapshot, title: str, table_font_plan: TableFontPlan | None = None
    ) -> InPlaceInspection:
        inspection = await super().inspect_in_place(
            snapshot=snapshot, title=title, table_font_plan=table_font_plan
        )
        if not self.failed:
            return inspection
        changed_fields = {
            "structure": {"structural_sha256": "changed-structure"},
            "protected": {"protected_structure_sha256": "changed-protected"},
            "body_hash": {"body_sha256": "changed-body"},
            "header_safe": {"header_safe_sha256": "changed-header-safe"},
            "marker": {"marker_range": {"startIndex": 21, "endIndex": 22, "tabId": "tab-1"}},
        }
        updates = changed_fields.get(self.drift)
        return replace(inspection, **updates) if updates is not None else inspection

    @override
    async def apply_in_place_stage(
        self,
        *,
        document_id: FileId,
        required_revision_id: DocsRevisionId,
        marker_name: str,
        marker_range: dict[str, Any],
        requests: list[dict[str, Any]],
    ) -> InPlaceStageResult:
        self.required_revisions.append(required_revision_id)
        is_header = any("createHeader" in request for request in requests)
        if is_header and not self.failed:
            self.failed = True
            if self.drift == "body":
                self.document = _in_place_document(
                    "Collaborator edit",
                    revision="revision-metadata-only",
                )
            else:
                self.document["revisionId"] = "revision-metadata-only"
                if self.drift == "header":
                    self.header_created = True
            raise DocumentGatewayError
        if required_revision_id != self.document["revisionId"]:
            raise SourceChangedError
        result = await super().apply_in_place_stage(
            document_id=document_id,
            required_revision_id=required_revision_id,
            marker_name=marker_name,
            marker_range=marker_range,
            requests=requests,
        )
        self.mutations["headers" if is_header else "content"] += 1
        return result


def _replay_executor(
    *,
    gateway: _InPlaceGateway,
    generator: _Generator,
    store: _CheckpointStore,
    task_id: TaskId,
    checkpoint: InPlaceCheckpoint | None = None,
) -> DocumentTransformExecutor:
    return DocumentTransformExecutor(
        gateway=gateway,
        generator=generator,
        allow_format_text_edits=True,
        checkpoint_store=store,
        task_id=task_id,
        lease_token=uuid.uuid7(),
        checkpoint=checkpoint,
        checkpoint_generation=store.generation,
    )


def _in_place_request(mode: TransformMode) -> DocumentTransformRequest:
    return DocumentTransformRequest(
        source=DriveFileSource(
            kind="drive_file",
            file_id=FileId("source"),
            expected_modified_time="2026-09-03T00:00:00Z",
        ),
        mode=mode,
        output_strategy=OutputStrategy.IN_PLACE,
    )


def test_in_place_active_suggestions_are_actionable_non_success() -> None:
    gateway = object.__new__(GoogleDocumentGateway)
    document = _in_place_document("Source")
    document["tabs"][0]["documentTab"]["body"]["content"][0]["paragraph"][
        "suggestedParagraphStyleChanges"
    ] = {"suggestion-1": {}}
    snapshot = _InPlaceGateway("Source").snapshot
    assert snapshot is not None
    snapshot = replace(snapshot, google_document=document)

    with pytest.raises(ActiveSuggestionsNeedsReviewError) as caught:
        gateway._inspect_in_place_sync(snapshot, "Source")

    assert caught.value.code == "needs_review"
    assert "Resolve active Google Docs suggestions" in caught.value.public_detail


async def test_in_place_custom_header_is_actionable_non_success() -> None:
    class UnknownHeaderGateway(_InPlaceGateway):
        @override
        async def inspect_in_place(
            self,
            *,
            snapshot: SourceSnapshot,
            title: str,
            table_font_plan: TableFontPlan | None = None,
        ) -> InPlaceInspection:
            inspection = await super().inspect_in_place(
                snapshot=snapshot,
                title=title,
                table_font_plan=table_font_plan,
            )
            return replace(
                inspection,
                headers=(),
                missing_header_tab_ids=(),
                unknown_header_tab_ids=(DocsTabId("tab-1"),),
            )

    gateway = UnknownHeaderGateway("Source")
    store = _CheckpointStore(gateway.events)

    with pytest.raises(CustomHeaderNeedsReviewError) as caught:
        await _replay_executor(
            gateway=gateway,
            generator=_Generator(),
            store=store,
            task_id=TaskId(uuid.uuid7()),
        ).execute(operation_id=uuid.uuid7(), request=_in_place_request(TransformMode.FORMAT))

    assert caught.value.code == "needs_review"
    assert gateway.events == []


async def test_in_place_unprovable_structure_is_durable_safe_noop(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class ReadbackGateway(_InPlaceGateway):
        def __init__(self) -> None:
            super().__init__("Source")
            self.loads = 0
            self.inspections = 0

        @override
        async def load_drive_source(self, source: DriveFileSource) -> SourceSnapshot:
            self.loads += 1
            return await super().load_drive_source(source)

        @override
        async def inspect_in_place(
            self,
            *,
            snapshot: SourceSnapshot,
            title: str,
            table_font_plan: TableFontPlan | None = None,
        ) -> InPlaceInspection:
            self.inspections += 1
            return await super().inspect_in_place(
                snapshot=snapshot, title=title, table_font_plan=table_font_plan
            )

    gateway = ReadbackGateway()
    paragraph_style = gateway.document["tabs"][0]["documentTab"]["body"]["content"][0]["paragraph"][
        "paragraphStyle"
    ]
    paragraph_style.update(
        alignment="CENTER",
        spaceAbove={"magnitude": 24, "unit": "PT"},
    )
    store = _CheckpointStore(gateway.events)
    generator = _Generator()
    task_id = TaskId(uuid.uuid7())
    operation_id = uuid.uuid7()
    request = _in_place_request(TransformMode.FORMAT)
    monkeypatch.setattr(executor_module, "preserved_native_findings", lambda **_kwargs: None)

    result = await _replay_executor(
        gateway=gateway,
        generator=generator,
        store=store,
        task_id=task_id,
    ).execute(operation_id=operation_id, request=request)

    assert result.changed is False
    assert result.needs_review is True
    assert result.applied_patch_count == 0
    assert result.source_revision_id == result.output_revision_id == "revision-1"
    assert [finding.code for finding in result.findings] == [
        "complex_native_structure_preserved",
        "native_paragraph_layout_preserved",
    ]
    assert gateway.loads == 2
    assert gateway.inspections == 2
    assert gateway.events == [f"checkpoint:{InPlaceCheckpointStage.COMPLETE}"]
    checkpoint = store.checkpoints[-1]
    assert len(checkpoint.paragraph_layout) == 1

    replay = await _replay_executor(
        gateway=gateway,
        generator=generator,
        store=store,
        task_id=task_id,
        checkpoint=checkpoint,
    ).execute(operation_id=operation_id, request=request)

    assert replay == result
    assert gateway.loads == 3
    assert gateway.inspections == 3
    assert gateway.events == [f"checkpoint:{InPlaceCheckpointStage.COMPLETE}"]


@pytest.mark.parametrize("drift", ["revision", "structure", "protected", "body_hash"])
async def test_in_place_complex_safe_noop_rejects_concurrent_readback_drift(
    monkeypatch: pytest.MonkeyPatch,
    drift: str,
) -> None:
    class DriftGateway(_InPlaceGateway):
        def __init__(self) -> None:
            super().__init__("Source")
            self.loads = 0
            self.inspections = 0

        @override
        async def load_drive_source(self, source: DriveFileSource) -> SourceSnapshot:
            self.loads += 1
            snapshot = await super().load_drive_source(source)
            if self.loads == 2 and drift == "revision":
                return replace(snapshot, revision_id=DocsRevisionId("revision-2"))
            return snapshot

        @override
        async def inspect_in_place(
            self,
            *,
            snapshot: SourceSnapshot,
            title: str,
            table_font_plan: TableFontPlan | None = None,
        ) -> InPlaceInspection:
            self.inspections += 1
            inspection = await super().inspect_in_place(
                snapshot=snapshot, title=title, table_font_plan=table_font_plan
            )
            if self.inspections != 2:
                return inspection
            updates = {
                "structure": {"structural_sha256": "changed-structure"},
                "protected": {"protected_structure_sha256": "changed-protected"},
                "body_hash": {"body_sha256": "changed-body"},
            }.get(drift, {})
            return replace(inspection, **updates)

    gateway = DriftGateway()
    store = _CheckpointStore(gateway.events)
    monkeypatch.setattr(executor_module, "preserved_native_findings", lambda **_kwargs: None)

    expected_error = SourceChangedError if drift == "revision" else InPlaceNeedsReviewError
    with pytest.raises(expected_error):
        await _replay_executor(
            gateway=gateway,
            generator=_Generator(),
            store=store,
            task_id=TaskId(uuid.uuid7()),
        ).execute(operation_id=uuid.uuid7(), request=_in_place_request(TransformMode.FORMAT))

    assert gateway.events == []


async def test_in_place_complex_safe_noop_keeps_unknown_header_hard(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class UnknownHeaderGateway(_InPlaceGateway):
        @override
        async def inspect_in_place(
            self,
            *,
            snapshot: SourceSnapshot,
            title: str,
            table_font_plan: TableFontPlan | None = None,
        ) -> InPlaceInspection:
            inspection = await super().inspect_in_place(
                snapshot=snapshot, title=title, table_font_plan=table_font_plan
            )
            return replace(
                inspection,
                headers=(),
                missing_header_tab_ids=(),
                unknown_header_tab_ids=(DocsTabId("tab-1"),),
            )

    gateway = UnknownHeaderGateway("Source")
    store = _CheckpointStore(gateway.events)
    monkeypatch.setattr(executor_module, "preserved_native_findings", lambda **_kwargs: None)

    with pytest.raises(CustomHeaderNeedsReviewError):
        await _replay_executor(
            gateway=gateway,
            generator=_Generator(),
            store=store,
            task_id=TaskId(uuid.uuid7()),
        ).execute(operation_id=uuid.uuid7(), request=_in_place_request(TransformMode.FORMAT))

    assert gateway.events == []


async def test_in_place_complex_safe_noop_keeps_font_proof_failure_hard(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = _InPlaceGateway("Source")
    store = _CheckpointStore(gateway.events)
    monkeypatch.setattr(executor_module, "preserved_native_findings", lambda **_kwargs: None)
    monkeypatch.setattr(
        executor_module,
        "capture_native_font_proof",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(ValueError),
    )

    with pytest.raises(InPlaceNeedsReviewError):
        await _replay_executor(
            gateway=gateway,
            generator=_Generator(),
            store=store,
            task_id=TaskId(uuid.uuid7()),
        ).execute(operation_id=uuid.uuid7(), request=_in_place_request(TransformMode.FORMAT))

    assert gateway.events == []


async def test_in_place_complex_safe_noop_keeps_pagination_failure_hard(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = _InPlaceGateway("Source")
    store = _CheckpointStore(gateway.events)
    document = parse_google_document(gateway.document, title="Source")
    block = document.blocks[0]
    pagination = executor_module.NativePaginationRule(
        block_id=block.id,
        text_sha256=block.sha256,
        keep_with_next=True,
    )
    monkeypatch.setattr(executor_module, "preserved_native_findings", lambda **_kwargs: None)
    monkeypatch.setattr(executor_module, "existing_native_pagination", lambda *_args: (pagination,))

    with pytest.raises(InPlaceNeedsReviewError):
        await _replay_executor(
            gateway=gateway,
            generator=_Generator(),
            store=store,
            task_id=TaskId(uuid.uuid7()),
        ).execute(operation_id=uuid.uuid7(), request=_in_place_request(TransformMode.FORMAT))

    assert gateway.events == []


async def test_in_place_complex_safe_noop_rejects_malformed_checkpoint_claim(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = _InPlaceGateway("Source")
    gateway.document["tabs"][0]["documentTab"]["body"]["content"][0]["paragraph"]["paragraphStyle"][
        "alignment"
    ] = "CENTER"
    store = _CheckpointStore(gateway.events)
    task_id = TaskId(uuid.uuid7())
    request = _in_place_request(TransformMode.FORMAT)
    monkeypatch.setattr(executor_module, "preserved_native_findings", lambda **_kwargs: None)
    await _replay_executor(
        gateway=gateway,
        generator=_Generator(),
        store=store,
        task_id=task_id,
    ).execute(operation_id=task_id, request=request)
    checkpoint = store.checkpoints[-1]
    assert len(checkpoint.paragraph_layout) == 1
    malformed = checkpoint.model_copy(
        update={
            "findings": [checkpoint.findings[0].model_copy(update={"detail": "untrusted detail"})]
        }
    )

    with pytest.raises(InPlaceNeedsReviewError):
        await _replay_executor(
            gateway=gateway,
            generator=_Generator(),
            store=store,
            task_id=task_id,
            checkpoint=malformed,
        ).execute(operation_id=task_id, request=request)

    assert gateway.events == [f"checkpoint:{InPlaceCheckpointStage.COMPLETE}"]


async def test_in_place_missing_brand_asset_is_failure_without_mutation() -> None:
    class MissingLogoGateway(_InPlaceGateway):
        @override
        async def inspect_in_place(
            self,
            *,
            snapshot: SourceSnapshot,
            title: str,
            table_font_plan: TableFontPlan | None = None,
        ) -> InPlaceInspection:
            inspection = await super().inspect_in_place(
                snapshot=snapshot,
                title=title,
                table_font_plan=table_font_plan,
            )
            return replace(inspection, logo_uri="")

    gateway = MissingLogoGateway("Source")
    store = _CheckpointStore(gateway.events)

    with pytest.raises(BrandAssetUnavailableError) as caught:
        await _replay_executor(
            gateway=gateway,
            generator=_Generator(),
            store=store,
            task_id=TaskId(uuid.uuid7()),
        ).execute(operation_id=uuid.uuid7(), request=_in_place_request(TransformMode.FORMAT))

    assert caught.value.code == "brand_asset_unavailable"
    assert gateway.events == []


async def test_in_place_unslop_checkpoints_before_mutation_and_reports_revisions(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = _InPlaceGateway("Moreover, we approved the plan.")
    store = _CheckpointStore(gateway.events)
    result = await DocumentTransformExecutor(
        gateway=gateway,
        generator=_Generator(),
        checkpoint_store=store,
        task_id=TaskId(uuid.uuid7()),
        lease_token=uuid.uuid7(),
    ).execute(operation_id=uuid.uuid7(), request=_in_place_request(TransformMode.UNSLOP))

    assert gateway.events[:2] == ["checkpoint:content", "apply"]
    assert result.output.file_id == "source"
    assert result.output_strategy is OutputStrategy.IN_PLACE
    assert result.changed is True
    assert result.source_revision_id == "revision-1"
    assert result.output_revision_id == "revision-3"
    persisted = task_service_module._serialize_checkpoint(store.checkpoints[0])
    assert "Moreover, we approved the plan." not in json.dumps(persisted)
    assert "Source" not in json.dumps(persisted)
    assert "source_title" not in persisted
    assert set(persisted) >= {"plan_ciphertext", "stage", "source_sha256"}
    monkeypatch.setattr(task_service_module, "_MAX_CHECKPOINT_PLAN_BYTES", 1)
    with pytest.raises(CheckpointTooLargeError):
        task_service_module._serialize_checkpoint(store.checkpoints[0])


async def test_in_place_clean_unslop_is_zero_batch_noop() -> None:
    gateway = _InPlaceGateway("The team approved the plan.")
    store = _CheckpointStore(gateway.events)
    result = await DocumentTransformExecutor(
        gateway=gateway,
        generator=_Generator(),
        checkpoint_store=store,
        task_id=TaskId(uuid.uuid7()),
        lease_token=uuid.uuid7(),
    ).execute(operation_id=uuid.uuid7(), request=_in_place_request(TransformMode.UNSLOP))

    assert "apply" not in gateway.events
    assert result.changed is False
    assert result.source_revision_id == result.output_revision_id == "revision-1"


async def test_in_place_format_checkpoints_header_then_content() -> None:
    gateway = _InPlaceGateway("Quarterly review", missing_header=True)
    store = _CheckpointStore(gateway.events)
    result = await DocumentTransformExecutor(
        gateway=gateway,
        generator=_Generator(),
        checkpoint_store=store,
        task_id=TaskId(uuid.uuid7()),
        lease_token=uuid.uuid7(),
    ).execute(operation_id=uuid.uuid7(), request=_in_place_request(TransformMode.FORMAT))

    assert gateway.events[:4] == [
        "checkpoint:headers",
        "apply",
        "checkpoint:content",
        "apply",
    ]
    assert result.changed is True
    assert result.output_revision_id == "revision-3"


@pytest.mark.parametrize("failure_stage", ["headers", "content"])
async def test_in_place_verification_failure_after_mutation_is_truthful(
    failure_stage: str,
) -> None:
    class VerificationFailureGateway(_InPlaceGateway):
        @override
        async def verify_in_place(self, **kwargs: Any) -> OutputVerification:
            verified = await super().verify_in_place(**kwargs)
            is_header = kwargs["verify_brand"] is False
            if (failure_stage == "headers") == is_header:
                return replace(verified, content_preserved=False)
            return verified

    gateway = VerificationFailureGateway("Quarterly review", missing_header=True)
    store = _CheckpointStore(gateway.events)

    with pytest.raises(InPlacePostMutationVerificationError) as caught:
        await DocumentTransformExecutor(
            gateway=gateway,
            generator=_Generator(),
            checkpoint_store=store,
            task_id=TaskId(uuid.uuid7()),
            lease_token=uuid.uuid7(),
        ).execute(operation_id=uuid.uuid7(), request=_in_place_request(TransformMode.FORMAT))

    assert caught.value.code == "in_place_verification_failed"
    assert "may have changed" in caught.value.public_detail
    assert gateway.events.count("apply") == (1 if failure_stage == "headers" else 2)


@pytest.mark.parametrize("tamper_list_identity", [False, True])
async def test_in_place_provider_split_native_list_reaches_verified_post_write_state(
    tamper_list_identity: bool,
) -> None:
    document = _branded_google_document()
    document["documentId"] = "source"
    paragraph = document["tabs"][0]["documentTab"]["body"]["content"][0]["paragraph"]
    visible_style = deepcopy(paragraph["elements"][0]["textRun"]["textStyle"])
    visible_style["weightedFontFamily"] = {"fontFamily": "Arial"}
    paragraph["elements"] = [
        {
            "startIndex": 1,
            "endIndex": 5,
            "textRun": {"content": "Appr", "textStyle": deepcopy(visible_style)},
        },
        {
            "startIndex": 5,
            "endIndex": 10,
            "textRun": {"content": "oved.", "textStyle": deepcopy(visible_style)},
        },
        {
            "startIndex": 10,
            "endIndex": 11,
            "textRun": {"content": "\n", "textStyle": {"bold": False}},
        },
    ]
    paragraph["bullet"] = {"listId": "list", "nestingLevel": 0}
    paragraph["paragraphStyle"].pop("spaceAbove")
    paragraph["paragraphStyle"].pop("spaceBelow")
    document["tabs"][0]["documentTab"]["lists"] = {
        "list": {"listProperties": {"nestingLevels": [{}]}}
    }

    class ProviderSplitGateway(GoogleDocumentGateway):
        def __init__(self) -> None:
            self._logo_uri = "https://wiki.sarj.ai/brand/sarj-logo-light.png"
            self._header_identity: HeaderIdentityEvidence | None = None
            self.document = deepcopy(document)
            self.calls: list[list[dict[str, Any]]] = []

        @override
        async def load_drive_source(self, source: DriveFileSource) -> SourceSnapshot:
            assert source.file_id == "source"
            return SourceSnapshot(
                title="Decision",
                parent_folder_id=FileId("folder"),
                modified_time="2026-09-03T00:00:00Z",
                revision_id=DocsRevisionId(self.document["revisionId"]),
                media_type=MIME_DOC,
                google_document=deepcopy(self.document),
                owned_by_me=True,
                can_edit=True,
                is_personal_drive=True,
            )

        @override
        async def apply_in_place_stage(self, **kwargs: Any) -> InPlaceStageResult:
            assert kwargs["required_revision_id"] == self.document["revisionId"]
            requests = deepcopy(kwargs["requests"])
            self.calls.append(requests)
            applied_paragraph = self.document["tabs"][0]["documentTab"]["body"]["content"][0][
                "paragraph"
            ]
            for request in requests:
                text_update = request.get("updateTextStyle")
                if isinstance(text_update, dict):
                    target = text_update.get("range", {})
                    assert not (
                        target.get("tabId") == "tab-a"
                        and target.get("segmentId") is None
                        and target.get("startIndex", 0) < 11
                        and target.get("endIndex", 0) > 1
                    )
                paragraph_update = request.get("updateParagraphStyle")
                if isinstance(paragraph_update, dict):
                    target = paragraph_update.get("range", {})
                    if target.get("tabId") == "tab-a" and target.get("segmentId") is None:
                        applied_paragraph["paragraphStyle"].update(
                            deepcopy(paragraph_update["paragraphStyle"])
                        )
            if tamper_list_identity:
                applied_paragraph["bullet"]["listId"] = "other"
            self.document["revisionId"] = "revision-3"
            return InPlaceStageResult(DocsRevisionId("revision-3"), deepcopy(self.document))

        @override
        async def plan_tab_naming(self, **_kwargs: Any) -> None:
            return None

    gateway = ProviderSplitGateway()
    store = _CheckpointStore([])
    task_id = TaskId(uuid.uuid7())
    request = _in_place_request(TransformMode.FORMAT)
    executor = DocumentTransformExecutor(
        gateway=gateway,
        generator=_Generator(),
        checkpoint_store=store,
        task_id=task_id,
        lease_token=uuid.uuid7(),
    )

    if tamper_list_identity:
        with pytest.raises(InPlacePostMutationVerificationError):
            await executor.execute(operation_id=task_id, request=request)
        assert store.checkpoints[-1].stage is InPlaceCheckpointStage.CONTENT
    else:
        result = await executor.execute(operation_id=task_id, request=request)
        assert result.output_revision_id == "revision-3"
        completed = store.checkpoints[-1]
        replay_generator = _Generator({TransformPhase.FORMAT: "must not run"})
        replay = await DocumentTransformExecutor(
            gateway=gateway,
            generator=replay_generator,
            checkpoint_store=store,
            task_id=task_id,
            lease_token=uuid.uuid7(),
            checkpoint=completed,
            checkpoint_generation=store.generation,
        ).execute(operation_id=task_id, request=request)
        assert replay.output_revision_id == "revision-3"
        assert replay_generator.calls == []
    assert len(gateway.calls) == 1
    assert any("updateParagraphStyle" in request for request in gateway.calls[0])


@pytest.mark.parametrize("crash_stage", ["headers", "content"])
async def test_in_place_persisted_checkpoint_replay_skips_model_and_duplicate_batches(
    crash_stage: str,
) -> None:
    operation_id = uuid.uuid7()
    task_id = TaskId(uuid.uuid7())
    gateway = _CrashReplayInPlaceGateway("Moreover, we approved the plan.", crash_stage=crash_stage)
    store = _CheckpointStore(gateway.events)
    first_generator = _Generator()

    with pytest.raises(RuntimeError, match="simulated worker crash"):
        await _replay_executor(
            gateway=gateway,
            generator=first_generator,
            store=store,
            task_id=task_id,
        ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))

    assert first_generator.calls == [TransformPhase.FORMAT]
    persisted = store.checkpoints[-1]
    retry_generator = _Generator({TransformPhase.FORMAT: "must never be generated"})
    result = await _replay_executor(
        gateway=gateway,
        generator=retry_generator,
        store=store,
        task_id=task_id,
        checkpoint=persisted,
    ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))

    assert retry_generator.calls == []
    assert gateway.mutations == {"headers": 1, "content": 1}
    assert result.output_strategy is OutputStrategy.IN_PLACE
    assert result.applied_patch_count == 1
    assert (
        parse_google_document(gateway.document, title="Source").blocks[0].text
        == "We approved the plan."
    )
    assert result.changed is True
    assert result.output_revision_id == "revision-3"


async def test_in_place_tab_checkpoint_resume_never_replays_content(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = _CrashReplayInPlaceGateway("Moreover, we approved the plan.", crash_stage="never")
    store = _CheckpointStore(gateway.events)
    operation_id = uuid.uuid7()
    task_id = TaskId(uuid.uuid7())
    calls: list[str] = []

    async def plan_tabs(**_kwargs: object) -> TabNamingCheckpoint:
        return TabNamingCheckpoint(
            plan=TabNamingPlan((), ()),
            revision_id=DocsRevisionId("revision-3"),
            marker_name="tabs",
            marker_range={"startIndex": 1, "endIndex": 2, "tabId": "tab-1"},
            invariant_sha256="test",
        )

    async def apply_tabs(**_kwargs: object) -> DocsRevisionId:
        calls.append("tabs")
        gateway.document["revisionId"] = "revision-4"
        if len(calls) == 1:
            raise DocumentGatewayError
        return DocsRevisionId("revision-4")

    async def reject_content_replay(**_kwargs: object) -> InPlaceStageResult:
        pytest.fail("content must never be replayed after the metadata checkpoint")

    monkeypatch.setattr(gateway, "plan_tab_naming", plan_tabs)
    monkeypatch.setattr(gateway, "apply_tab_naming", apply_tabs)
    with pytest.raises(DocumentGatewayError):
        await _replay_executor(
            gateway=gateway,
            generator=_Generator(),
            store=store,
            task_id=task_id,
        ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))
    assert store.checkpoints[-1].tab_naming is not None
    monkeypatch.setattr(gateway, "apply_in_place_stage", reject_content_replay)
    result = await _replay_executor(
        gateway=gateway,
        generator=_Generator(),
        store=store,
        task_id=task_id,
        checkpoint=store.checkpoints[-1],
    ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))
    assert result.output_revision_id == "revision-4"
    assert len(calls) == 2


async def test_in_place_unsupported_tab_metadata_is_durable_decline_after_verified_body(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = _CrashReplayInPlaceGateway("Moreover, we approved the plan.", crash_stage="never")
    store = _CheckpointStore(gateway.events)
    operation_id = uuid.uuid7()
    task_id = TaskId(uuid.uuid7())
    naming_attempts = 0

    async def decline_tabs(**_kwargs: object) -> None:
        nonlocal naming_attempts
        naming_attempts += 1
        raise UnsupportedTabMetadataError

    monkeypatch.setattr(gateway, "plan_tab_naming", decline_tabs)
    result = await _replay_executor(
        gateway=gateway,
        generator=_Generator(),
        store=store,
        task_id=task_id,
    ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))

    assert gateway.mutations == {"headers": 1, "content": 1}
    assert store.checkpoints[-1].stage is InPlaceCheckpointStage.COMPLETE
    assert store.checkpoints[-1].tab_naming_declined_reason == "unsupported_metadata"
    assert [(finding.code, finding.severity) for finding in result.findings] == [
        ("tab_naming_preserved", FindingSeverity.INFORMATIONAL)
    ]
    assert not result.needs_review

    replayed = await _replay_executor(
        gateway=gateway,
        generator=_Generator(),
        store=store,
        task_id=task_id,
        checkpoint=store.checkpoints[-1],
    ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))
    assert replayed.findings == result.findings
    assert replayed.output_revision_id == result.output_revision_id
    assert replayed.validation == result.validation
    assert naming_attempts == 1
    assert gateway.mutations == {"headers": 1, "content": 1}


async def test_in_place_decline_crash_replay_never_reenters_verified_body(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = _CrashReplayInPlaceGateway("Moreover, we approved the plan.", crash_stage="never")
    store = _CheckpointStore(gateway.events)
    operation_id = uuid.uuid7()
    task_id = TaskId(uuid.uuid7())
    original_save = store.save_document_transform_checkpoint
    crashed = False
    crash_message = "simulated crash after durable tab decline"

    async def decline_tabs(**_kwargs: object) -> None:
        raise UnsupportedTabMetadataError

    async def crash_after_decline(
        saved_task_id: TaskId,
        *,
        lease_token: uuid.UUID,
        expected_generation: int,
        checkpoint: DocumentTransformCheckpoint,
    ) -> int:
        nonlocal crashed
        generation = await original_save(
            saved_task_id,
            lease_token=lease_token,
            expected_generation=expected_generation,
            checkpoint=checkpoint,
        )
        if (
            isinstance(checkpoint, InPlaceCheckpoint)
            and checkpoint.tab_naming_declined_reason is not None
            and not crashed
        ):
            crashed = True
            raise RuntimeError(crash_message)
        return generation

    monkeypatch.setattr(gateway, "plan_tab_naming", decline_tabs)
    monkeypatch.setattr(store, "save_document_transform_checkpoint", crash_after_decline)
    with pytest.raises(RuntimeError, match="durable tab decline"):
        await _replay_executor(
            gateway=gateway,
            generator=_Generator(),
            store=store,
            task_id=task_id,
        ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))

    persisted = store.checkpoints[-1]
    assert persisted.stage is InPlaceCheckpointStage.CONTENT
    assert persisted.tab_naming_declined_reason == "unsupported_metadata"
    assert gateway.mutations == {"headers": 1, "content": 1}

    async def reject_body_reentry(**_kwargs: object) -> Never:
        pytest.fail("verified body stage must not be re-entered after a durable tab decline")

    monkeypatch.setattr(gateway, "apply_in_place_stage", reject_body_reentry)
    result = await _replay_executor(
        gateway=gateway,
        generator=_Generator(),
        store=store,
        task_id=task_id,
        checkpoint=persisted,
    ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))
    assert [finding.code for finding in result.findings] == ["tab_naming_preserved"]
    assert store.checkpoints[-1].stage is InPlaceCheckpointStage.COMPLETE
    assert gateway.mutations == {"headers": 1, "content": 1}


async def test_copy_unsupported_tab_metadata_is_durable_decline_after_verified_body(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    gateway = _native_copy_gateway()
    gateway.verification = OutputVerification(
        content_preserved=True,
        structure_preserved=True,
        brand_applied=True,
    )
    store = _CopyCheckpointStore()
    operation_id = uuid.uuid7()
    task_id = TaskId(uuid.uuid7())
    naming_attempts = 0

    async def decline_tabs(**_kwargs: object) -> None:
        nonlocal naming_attempts
        naming_attempts += 1
        raise UnsupportedTabMetadataError

    monkeypatch.setattr(gateway, "plan_tab_naming", decline_tabs)
    request = _native_copy_request().model_copy(update={"mode": TransformMode.FORMAT})
    result = await _copy_executor(gateway=gateway, store=store, task_id=task_id).execute(
        operation_id=operation_id,
        request=request,
    )

    assert gateway.applies == 1
    assert store.checkpoints[-1].stage is CopyCheckpointStage.VERIFIED
    assert store.checkpoints[-1].tab_naming_declined_reason == "unsupported_metadata"
    assert [(finding.code, finding.severity) for finding in result.findings] == [
        ("tab_naming_preserved", FindingSeverity.INFORMATIONAL)
    ]
    assert not result.needs_review

    replayed = await _copy_executor(
        gateway=gateway,
        store=store,
        task_id=task_id,
        checkpoint=store.checkpoints[-1],
    ).execute(operation_id=operation_id, request=request)
    assert replayed == result
    assert naming_attempts == 1
    assert gateway.applies == 1


async def test_in_place_header_retry_rebases_metadata_only_revision_drift() -> None:
    operation_id = uuid.uuid7()
    task_id = TaskId(uuid.uuid7())
    gateway = _PreMutationRevisionDriftGateway("Moreover, we approved the plan.")
    store = _CheckpointStore(gateway.events)
    first_generator = _Generator()

    with pytest.raises(DocumentGatewayError):
        await _replay_executor(
            gateway=gateway,
            generator=first_generator,
            store=store,
            task_id=task_id,
        ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))

    persisted = store.checkpoints[-1]
    retry_generator = _Generator({TransformPhase.FORMAT: "must never be generated"})
    result = await _replay_executor(
        gateway=gateway,
        generator=retry_generator,
        store=store,
        task_id=task_id,
        checkpoint=persisted,
    ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))

    assert first_generator.calls == [TransformPhase.FORMAT]
    assert retry_generator.calls == []
    assert gateway.required_revisions[:2] == ["revision-1", "revision-metadata-only"]
    assert gateway.mutations == {"headers": 1, "content": 1}
    assert result.applied_patch_count == 1
    assert (
        parse_google_document(gateway.document, title="Source").blocks[0].text
        == "We approved the plan."
    )
    assert result.output_revision_id == "revision-3"


@pytest.mark.parametrize(
    "drift",
    ["body", "header", "structure", "protected", "body_hash", "header_safe", "marker"],
)
async def test_in_place_header_retry_does_not_rebase_user_visible_drift(drift: str) -> None:
    operation_id = uuid.uuid7()
    task_id = TaskId(uuid.uuid7())
    gateway = _PreMutationRevisionDriftGateway("Quarterly review", drift=drift)
    store = _CheckpointStore(gateway.events)

    with pytest.raises(DocumentGatewayError):
        await _replay_executor(
            gateway=gateway,
            generator=_Generator(),
            store=store,
            task_id=task_id,
        ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.FORMAT))

    persisted = store.checkpoints[-1]
    with pytest.raises(SourceChangedError):
        await _replay_executor(
            gateway=gateway,
            generator=_Generator({TransformPhase.FORMAT: "must never be generated"}),
            store=store,
            task_id=task_id,
            checkpoint=persisted,
        ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.FORMAT))

    assert gateway.required_revisions == ["revision-1", "revision-1"]
    assert gateway.mutations == {"headers": 0, "content": 0}


async def test_in_place_content_replay_rejects_post_commit_collaborator_edit() -> None:
    operation_id = uuid.uuid7()
    task_id = TaskId(uuid.uuid7())
    gateway = _CrashReplayInPlaceGateway("Moreover, we approved the plan.", crash_stage="content")
    store = _CheckpointStore(gateway.events)

    with pytest.raises(RuntimeError, match="simulated worker crash"):
        await _replay_executor(
            gateway=gateway,
            generator=_Generator(),
            store=store,
            task_id=task_id,
        ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))

    body = gateway.document["tabs"][0]["documentTab"]["body"]["content"]
    collaborator_text = "Collaborator edit\n"
    body.append(
        {
            "startIndex": body[-1]["endIndex"],
            "endIndex": body[-1]["endIndex"] + utf16_length(collaborator_text),
            "paragraph": {"elements": [{"textRun": {"content": collaborator_text}}]},
        }
    )
    gateway.document["revisionId"] = "collaborator-revision"
    persisted = task_service_module._deserialize_checkpoint(
        task_service_module._serialize_checkpoint(store.checkpoints[-1])
    )
    assert isinstance(persisted, InPlaceCheckpoint)
    retry_generator = _Generator({TransformPhase.FORMAT: "must never be generated"})

    with pytest.raises(InPlaceNeedsReviewError):
        await _replay_executor(
            gateway=gateway,
            generator=retry_generator,
            store=store,
            task_id=task_id,
            checkpoint=persisted,
        ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))

    assert retry_generator.calls == []
    assert gateway.mutations == {"headers": 1, "content": 1}


async def test_in_place_complete_checkpoint_rejects_later_revision_without_mutation() -> None:
    operation_id = uuid.uuid7()
    task_id = TaskId(uuid.uuid7())
    gateway = _CrashReplayInPlaceGateway("Quarterly review", crash_stage="never")
    store = _CheckpointStore(gateway.events)
    request = _in_place_request(TransformMode.FORMAT)

    await _replay_executor(
        gateway=gateway,
        generator=_Generator(),
        store=store,
        task_id=task_id,
    ).execute(operation_id=operation_id, request=request)

    completed = store.checkpoints[-1]
    assert completed.stage is InPlaceCheckpointStage.COMPLETE
    assert gateway.mutations == {"headers": 1, "content": 1}
    gateway.document = _in_place_document(
        "Later collaborator edit",
        revision="revision-after-complete",
    )
    retry_generator = _Generator({TransformPhase.FORMAT: "must never be generated"})

    with pytest.raises(InPlaceNeedsReviewError):
        await _replay_executor(
            gateway=gateway,
            generator=retry_generator,
            store=store,
            task_id=task_id,
            checkpoint=completed,
        ).execute(operation_id=operation_id, request=request)

    assert retry_generator.calls == []
    assert gateway.mutations == {"headers": 1, "content": 1}


async def test_in_place_checkpoint_replay_rejects_source_title_drift() -> None:
    operation_id = uuid.uuid7()
    task_id = TaskId(uuid.uuid7())
    gateway = _CrashReplayInPlaceGateway("Moreover, we approved the plan.", crash_stage="headers")
    store = _CheckpointStore(gateway.events)

    with pytest.raises(RuntimeError, match="simulated worker crash"):
        await _replay_executor(
            gateway=gateway,
            generator=_Generator(),
            store=store,
            task_id=task_id,
        ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))

    stale_title_checkpoint = store.checkpoints[-1].model_copy(
        update={"source_title": "Former source title"}
    )
    retry_generator = _Generator({TransformPhase.FORMAT: "must never be generated"})
    with pytest.raises(OutputConflictError):
        await _replay_executor(
            gateway=gateway,
            generator=retry_generator,
            store=store,
            task_id=task_id,
            checkpoint=stale_title_checkpoint,
        ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))

    assert retry_generator.calls == []
    assert gateway.mutations == {"headers": 1, "content": 0}


async def test_in_place_replay_rejects_concurrent_drive_rename_without_more_mutation() -> None:
    operation_id = uuid.uuid7()
    task_id = TaskId(uuid.uuid7())
    gateway = _CrashReplayInPlaceGateway("Moreover, we approved the plan.", crash_stage="headers")
    store = _CheckpointStore(gateway.events)

    with pytest.raises(RuntimeError, match="simulated worker crash"):
        await _replay_executor(
            gateway=gateway,
            generator=_Generator(),
            store=store,
            task_id=task_id,
        ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))

    gateway.drive_title = "Collaborator Renamed This Document"
    retry_generator = _Generator({TransformPhase.FORMAT: "must never be generated"})

    with pytest.raises(OutputConflictError):
        await _replay_executor(
            gateway=gateway,
            generator=retry_generator,
            store=store,
            task_id=task_id,
            checkpoint=store.checkpoints[-1],
        ).execute(operation_id=operation_id, request=_in_place_request(TransformMode.BOTH))

    assert retry_generator.calls == []
    assert gateway.drive_title == "Collaborator Renamed This Document"
    assert gateway.mutations == {"headers": 1, "content": 0}


def test_in_place_request_rejects_markdown_and_output_name() -> None:
    with pytest.raises(ValueError, match="native Google Doc"):
        DocumentTransformRequest(
            source=MarkdownSource(
                kind="markdown",
                content="Draft",
                name="Draft",
                destination_folder_id=FileId("folder"),
            ),
            mode=TransformMode.UNSLOP,
            output_strategy=OutputStrategy.IN_PLACE,
        )
    with pytest.raises(ValueError, match="output_name"):
        DocumentTransformRequest(
            source=DriveFileSource(
                kind="drive_file",
                file_id=FileId("source"),
                expected_modified_time="time",
            ),
            mode=TransformMode.UNSLOP,
            output_name="Fork",
            output_strategy=OutputStrategy.IN_PLACE,
        )


def test_copy_destination_request_is_drive_copy_only() -> None:
    with pytest.raises(ValueError, match="Drive file copies"):
        DocumentTransformRequest(
            source=MarkdownSource(
                kind="markdown",
                content="Draft",
                name="Draft",
                destination_folder_id=FileId("folder"),
            ),
            mode=TransformMode.FORMAT,
            copy_destination_folder_id=FileId("destination"),
        )
    with pytest.raises(ValueError, match="Drive file copies"):
        DocumentTransformRequest(
            source=DriveFileSource(
                kind="drive_file",
                file_id=FileId("source"),
                expected_modified_time="time",
            ),
            mode=TransformMode.FORMAT,
            output_strategy=OutputStrategy.IN_PLACE,
            copy_destination_folder_id=FileId("destination"),
        )


def test_copy_request_keeps_legacy_canonical_payload_hash_shape() -> None:
    request = DocumentTransformRequest(
        source=DriveFileSource(
            kind="drive_file",
            file_id=FileId("source"),
            expected_modified_time="time",
        ),
        mode=TransformMode.UNSLOP,
    )

    serialized = task_service_module._serialize(request)

    assert "output_strategy" not in serialized
    assert "copy_destination_folder_id" not in serialized
    assert serialized["output_name"] is None


def test_copy_destination_is_part_of_canonical_payload_hash_shape() -> None:
    request = _native_copy_request(destination=FileId("qa-folder"))

    serialized = task_service_module._serialize(request)

    assert serialized["copy_destination_folder_id"] == "qa-folder"


def test_in_place_request_keeps_legacy_canonical_payload_hash_shape() -> None:
    request = DocumentTransformRequest(
        source=DriveFileSource(
            kind="drive_file",
            file_id=FileId("source"),
            expected_modified_time="time",
        ),
        mode=TransformMode.UNSLOP,
        output_strategy=OutputStrategy.IN_PLACE,
    )

    serialized = task_service_module._serialize(request)

    assert serialized["output_strategy"] == "in_place"
    assert "copy_destination_folder_id" not in serialized


def test_in_place_drive_id_is_bounded_and_google_safe() -> None:
    opaque = DriveFileSource(
        kind="drive_file",
        file_id=FileId("opaque/id:still-accepted-for-copy"),
        expected_modified_time="time",
    )
    assert opaque.file_id == "opaque/id:still-accepted-for-copy"
    with pytest.raises(ValueError, match="between 1 and 128"):
        DocumentTransformRequest(
            source=DriveFileSource(
                kind="drive_file",
                file_id=FileId("a" * 129),
                expected_modified_time="time",
            ),
            mode=TransformMode.UNSLOP,
            output_strategy=OutputStrategy.IN_PLACE,
        )


def test_in_place_marker_reconciliation_is_exact_and_tab_scoped() -> None:
    expected = {"startIndex": 10, "endIndex": 11, "tabId": "tab-1"}
    marker = "sarj-transform-operation-content-hash"
    document = _in_place_document("Source")
    provider_range = {"startIndex": 10, "endIndex": 11, "segmentId": ""}
    document["tabs"][0]["documentTab"]["namedRanges"] = {
        marker: {
            "namedRanges": [{"name": marker, "namedRangeId": "range-1", "ranges": [provider_range]}]
        }
    }

    assert google_gateway_module._has_exact_named_range(document, marker, expected)
    wrong_tab = {**expected, "tabId": "tab-2"}
    assert not google_gateway_module._has_exact_named_range(document, marker, wrong_tab)
    provider_range["segmentId"] = "header-1"
    assert not google_gateway_module._has_exact_named_range(document, marker, expected)
    provider_range["segmentId"] = ""
    document["tabs"][0]["documentTab"]["namedRanges"][marker]["namedRanges"].append(
        {"name": marker, "namedRangeId": "range-2", "ranges": [expected]}
    )
    assert not google_gateway_module._has_exact_named_range(document, marker, expected)


def test_in_place_rejects_active_suggestion_metadata() -> None:
    assert google_gateway_module._has_active_suggestions(
        {"paragraph": {"suggestedParagraphStyleChanges": {"suggestion-1": {}}}}
    )


def test_google_in_place_lost_response_reconciles_atomic_tab_marker(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    marker = "sarj-transform-operation-headers-hash"
    marker_range = {"startIndex": 10, "endIndex": 11, "tabId": "tab-1"}
    provider_range = {"startIndex": 10, "endIndex": 11, "segmentId": ""}
    document = _in_place_document("Source", revision="revision-2")
    document["tabs"][0]["documentTab"]["namedRanges"] = {
        marker: {
            "namedRanges": [{"name": marker, "namedRangeId": "range-1", "ranges": [provider_range]}]
        }
    }
    captured: dict[str, Any] = {}

    def batch_update(**kwargs: Any) -> object:
        captured.update(kwargs)
        return object()

    provider = SimpleNamespace(
        documents=lambda: SimpleNamespace(batchUpdate=batch_update),
    )
    gateway = object.__new__(GoogleDocumentGateway)

    def get_provider() -> Any:
        return provider

    def request_boundary(value: Any, **kwargs: Any) -> Any:
        return value(**kwargs)

    monkeypatch.setattr(gateway, "_docs", get_provider)
    monkeypatch.setattr(google_gateway_module, "_google_http_request", request_boundary)

    def lost_response(_request: object) -> dict[str, Any]:
        raise google_gateway_module.DocumentGatewayError

    def get_document(_document_id: FileId) -> dict[str, Any]:
        return document

    monkeypatch.setattr(gateway, "_execute_mutation_dict", lost_response)
    monkeypatch.setattr(gateway, "_get_document_sync", get_document)

    result = gateway._apply_in_place_stage_sync(
        FileId("source"),
        DocsRevisionId("revision-1"),
        marker,
        marker_range,
        [{"createHeader": {"type": "DEFAULT"}}],
    )

    body = captured["body"]
    assert body["writeControl"] == {"requiredRevisionId": "revision-1"}
    assert body["requests"][0] == {"createNamedRange": {"name": marker, "range": marker_range}}
    assert body["requests"][1] == {"createHeader": {"type": "DEFAULT"}}
    assert result.revision_id == "revision-2"


def test_google_in_place_lost_response_reconciles_utf16_shifted_marker(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    marker = "sarj-transform-operation-content-hash"
    marker_range = {"startIndex": 10, "endIndex": 11, "tabId": "tab-1"}
    shifted_range = {"startIndex": 13, "endIndex": 14, "tabId": "tab-1"}
    request = {
        "insertText": {
            "location": {"index": 2, "tabId": "tab-1"},
            "text": "A😀",
        }
    }
    document = _in_place_document("A😀Source", revision="revision-2")
    document["tabs"][0]["documentTab"]["namedRanges"] = {
        marker: {
            "namedRanges": [{"name": marker, "namedRangeId": "range-1", "ranges": [shifted_range]}]
        }
    }
    captured: dict[str, Any] = {}

    gateway = object.__new__(GoogleDocumentGateway)

    def batch_update(**kwargs: Any) -> object:
        captured.update(kwargs)
        return object()

    def get_provider() -> Any:
        return SimpleNamespace(documents=lambda: SimpleNamespace(batchUpdate=batch_update))

    def lost_response(_request: object) -> dict[str, Any]:
        raise DocumentGatewayError

    monkeypatch.setattr(gateway, "_docs", get_provider)
    monkeypatch.setattr(
        google_gateway_module,
        "_google_http_request",
        lambda value, **kwargs: value(**kwargs),
    )
    monkeypatch.setattr(gateway, "_execute_mutation_dict", lost_response)
    monkeypatch.setattr(gateway, "_get_document_sync", lambda _document_id: document)

    result = gateway._apply_in_place_stage_sync(
        FileId("source"),
        DocsRevisionId("revision-1"),
        marker,
        marker_range,
        [request],
    )

    body = captured["body"]
    assert body["writeControl"] == {"requiredRevisionId": "revision-1"}
    assert body["requests"] == [
        {"createNamedRange": {"name": marker, "range": marker_range}},
        request,
    ]
    assert result.revision_id == "revision-2"


@pytest.mark.parametrize(
    ("observed_revision", "expected_error"),
    [
        ("revision-1", DocumentGatewayError),
        ("collaborator-revision", SourceChangedError),
    ],
)
def test_google_in_place_lost_response_without_marker_fails_closed(
    monkeypatch: pytest.MonkeyPatch,
    observed_revision: str,
    expected_error: type[Exception],
) -> None:
    marker_range = {"startIndex": 10, "endIndex": 11, "tabId": "tab-1"}
    document = _in_place_document("Source", revision=observed_revision)

    gateway = object.__new__(GoogleDocumentGateway)

    def get_provider() -> Any:
        return SimpleNamespace(
            documents=lambda: SimpleNamespace(batchUpdate=lambda **_kwargs: object())
        )

    def lost_response(_request: object) -> dict[str, Any]:
        raise DocumentGatewayError

    monkeypatch.setattr(gateway, "_docs", get_provider)
    monkeypatch.setattr(
        google_gateway_module,
        "_google_http_request",
        lambda value, **kwargs: value(**kwargs),
    )
    monkeypatch.setattr(gateway, "_execute_mutation_dict", lost_response)
    monkeypatch.setattr(gateway, "_get_document_sync", lambda _document_id: document)

    with pytest.raises(expected_error):
        gateway._apply_in_place_stage_sync(
            FileId("source"),
            DocsRevisionId("revision-1"),
            "sarj-transform-operation-content-hash",
            marker_range,
            [{"createHeader": {"type": "DEFAULT"}}],
        )
