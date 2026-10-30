from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring
from docx import Document as WordDocument
from markdown_it import MarkdownIt
import pytest

from app.document_transform import docx_renderer
from app.document_transform.brand_manifest import CURRENT_BRAND_MANIFEST
from app.document_transform.docx_renderer import (
    markdown_import_manifest,
    render_markdown_docx,
    validate_markdown_import,
)
from app.document_transform.errors import (
    EmptyMarkdownBlockError,
    SourceTooLargeError,
    UnsupportedSourceError,
)


if TYPE_CHECKING:
    from xml.etree.ElementTree import Element


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
LOGO_BYTES = (Path(__file__).parents[2] / "frontend/public/brand/sarj-logo-light.png").read_bytes()


def _xml_paragraph_text(paragraph: Element[str]) -> str:
    return "".join(node.text or "" for node in paragraph.findall(f".//{W}t"))


SOURCE = """# Executive summary

This is **important**, *careful*, and uses `inline_code` plus [Sarj](https://sarj.ai).

1. First numbered item
2. Second numbered item
   - Nested bullet
- [ ] Pending review
- [x] Approved

> Preserve this quoted decision.

| Owner | Status | Detail |
| --- | --- | --- |
| Platform | Ready | Native structures |
| Docs | Review | Long values wrap rather than clip |

```python
def hello(name: str) -> str:
    return f"Hello, {name}"
```
"""


@pytest.fixture
def rendered() -> bytes:
    return render_markdown_docx(SOURCE, title="Technology review", logo_bytes=LOGO_BYTES)


def test_render_is_valid_loadable_and_byte_deterministic(rendered: bytes) -> None:
    assert rendered == render_markdown_docx(
        SOURCE,
        title="Technology review",
        logo_bytes=LOGO_BYTES,
    )
    document = WordDocument(BytesIO(rendered))
    assert document.core_properties.title == "Technology review"
    assert document.core_properties.author == "Sarj Wiki"
    with ZipFile(BytesIO(rendered)) as archive:
        assert archive.testzip() is None
        assert archive.namelist() == sorted(archive.namelist())
        assert all(item.date_time == (1980, 1, 1, 0, 0, 0) for item in archive.infolist())


def test_fenced_code_preserves_blank_lines_without_inserting_spaces() -> None:
    source = "```text\nfirst\n\nlast\n\n```"
    rendered = render_markdown_docx(source, title="Code", branded=False)
    document = WordDocument(BytesIO(rendered))
    expected = markdown_import_manifest(source, title="Code", branded=False)

    assert document.paragraphs[-1].text == "first\n\nlast\n"
    assert expected.blocks[-1].text == document.paragraphs[-1].text


@pytest.mark.parametrize(
    "source",
    [
        "#",
        "##\nBody",
        "Title\n-----\n\n##   ",
        "```\n```",
        "~~~python\n~~~",
        "```text\n\n```",
    ],
)
def test_empty_heading_and_code_blocks_are_rejected_before_render(source: str) -> None:
    with pytest.raises(EmptyMarkdownBlockError, match="add text or remove"):
        validate_markdown_import(source, title="Import")
    with pytest.raises(EmptyMarkdownBlockError):
        markdown_import_manifest(source, title="Import", branded=False)
    with pytest.raises(EmptyMarkdownBlockError):
        render_markdown_docx(source, title="Import", branded=False)


def test_render_preserves_native_structure_and_inline_formatting(rendered: bytes) -> None:
    with ZipFile(BytesIO(rendered)) as archive:
        document = fromstring(archive.read("word/document.xml"))
        relationships = fromstring(archive.read("word/_rels/document.xml.rels"))
        numbering = fromstring(archive.read("word/numbering.xml"))

    assert document.find(f".//{W}pStyle[@{W}val='Heading1']") is not None
    assert len(document.findall(f".//{W}numPr")) == 5
    assert document.find(f".//{W}tbl") is not None
    table_rows = document.findall(f".//{W}tbl/{W}tr")
    assert table_rows
    assert all(row.find(f"./{W}trPr/{W}cantSplit") is not None for row in table_rows)
    raw_grid_widths = [cell.get(f"{W}w") for cell in document.findall(f".//{W}gridCol")]
    assert all(width is not None for width in raw_grid_widths)
    grid_widths = [int(width) for width in raw_grid_widths if width is not None]
    assert len(grid_widths) == 3
    assert sum(grid_widths) == 9_360
    assert min(grid_widths) >= 1_300
    assert all(cell.get(f"{W}type") == "dxa" for cell in document.findall(f".//{W}tcW"))
    assert document.find(f".//{W}tblHeader") is not None
    assert document.find(f".//{W}b") is not None
    assert document.find(f".//{W}i") is not None
    assert document.find(f".//{W}hyperlink") is not None
    subtle = CURRENT_BRAND_MANIFEST.colors.subtle.removeprefix("#")
    assert document.find(f".//{W}shd[@{W}fill='{subtle}']") is not None
    targets = [item.attrib.get("Target") for item in relationships]
    assert "https://sarj.ai" in targets
    checklist_markers = [node.get(f"{W}val") for node in numbering.findall(f".//{W}lvlText")]
    assert "☐" in checklist_markers
    assert "☒" in checklist_markers


def test_sarj_paragraph_styles_follow_manifest(rendered: bytes) -> None:
    with ZipFile(BytesIO(rendered)) as archive:
        styles = fromstring(archive.read("word/styles.xml"))
    document_tokens = CURRENT_BRAND_MANIFEST.document
    primary = CURRENT_BRAND_MANIFEST.colors.primary.removeprefix("#")
    ink = CURRENT_BRAND_MANIFEST.colors.ink.removeprefix("#")
    expected_styles = (
        ("Normal", document_tokens.body_size_pt, ink, False),
        ("Heading1", document_tokens.h1_size_pt, primary, True),
        ("Heading2", document_tokens.h2_size_pt, primary, True),
        ("Heading3", document_tokens.h3_size_pt, primary, True),
        ("SarjTitle", document_tokens.h1_size_pt, primary, True),
    )
    for style_id, size_pt, color, heading in expected_styles:
        style = styles.find(f".//{W}style[@{W}styleId='{style_id}']")
        assert style is not None
        fonts = style.find(f"./{W}rPr/{W}rFonts")
        assert fonts is not None
        assert fonts.get(f"{W}ascii") == document_tokens.font_family
        assert fonts.get(f"{W}hAnsi") == document_tokens.font_family
        assert fonts.get(f"{W}eastAsia") == document_tokens.font_family
        assert fonts.get(f"{W}cs") == document_tokens.complex_script_font_family
        assert fonts.get(f"{W}cstheme") is None
        size = style.find(f"./{W}rPr/{W}sz")
        assert size is not None
        assert size.get(f"{W}val") == str(size_pt * 2)
        foreground = style.find(f"./{W}rPr/{W}color")
        assert foreground is not None
        assert foreground.get(f"{W}val") == color
        spacing = style.find(f"./{W}pPr/{W}spacing")
        assert spacing is not None
        assert spacing.get(f"{W}before") == str(document_tokens.paragraph_space_above_pt * 20)
        assert spacing.get(f"{W}after") == str(document_tokens.paragraph_space_below_pt * 20)
        assert spacing.get(f"{W}line") == str(
            round(document_tokens.line_spacing_percent / 100 * 240)
        )
        keep_next = style.find(f"./{W}pPr/{W}keepNext")
        if heading:
            assert keep_next is not None
            assert keep_next.get(f"{W}val") != "0"
        else:
            assert keep_next is None or keep_next.get(f"{W}val") == "0"


def test_plain_runs_inherit_bold_and_italic_from_paragraph_styles() -> None:
    rendered = render_markdown_docx(
        "# Heading\n\n> Quoted text\n\nPlain **bold** and *italic*.",
        title="Styled title",
    )
    document = WordDocument(BytesIO(rendered))
    heading, quote, body = document.paragraphs
    title = WordDocument(BytesIO(render_markdown_docx("Body", title="Styled title"))).paragraphs[0]

    assert title.style is not None
    assert title.style.font.bold is True
    assert title.runs[0].bold is None
    assert heading.style is not None
    assert heading.style.font.bold is True
    assert heading.runs[0].bold is None
    assert quote.style is not None
    assert quote.style.font.italic is True
    assert quote.runs[0].italic is None
    assert next(run for run in body.runs if run.text == "bold").bold is True
    assert next(run for run in body.runs if run.text == "italic").italic is True

    with ZipFile(BytesIO(rendered)) as archive:
        document_xml = fromstring(archive.read("word/document.xml"))
    styled_paragraphs = document_xml.findall(f".//{W}body/{W}p")[:2]
    assert all(paragraph.find(f".//{W}b") is None for paragraph in styled_paragraphs)
    assert all(paragraph.find(f".//{W}i") is None for paragraph in styled_paragraphs)


def test_all_branded_paragraph_styles_pin_complex_script_font() -> None:
    rendered = render_markdown_docx("Body", title="Styles")
    with ZipFile(BytesIO(rendered)) as archive:
        styles = fromstring(archive.read("word/styles.xml"))

    style_ids = (
        "Normal",
        "Heading1",
        "Heading2",
        "Heading3",
        "Heading4",
        "Heading5",
        "Heading6",
        "SarjTitle",
        "SarjList",
        "SarjQuote",
        "SarjCode",
        "SarjTableText",
    )
    for style_id in style_ids:
        style = styles.find(f".//{W}style[@{W}styleId='{style_id}']")
        assert style is not None
        fonts = style.find(f"./{W}rPr/{W}rFonts")
        assert fonts is not None
        assert fonts.get(f"{W}cs") == "Arial"
        assert fonts.get(f"{W}cstheme") is None


def test_sarj_header_logo_footer_and_margins_follow_manifest(rendered: bytes) -> None:
    with ZipFile(BytesIO(rendered)) as archive:
        styles = fromstring(archive.read("word/styles.xml"))
        document = fromstring(archive.read("word/document.xml"))
        header = fromstring(archive.read("word/header1.xml"))
        footer = fromstring(archive.read("word/footer1.xml"))

    document_tokens = CURRENT_BRAND_MANIFEST.document
    title_style = styles.find(f".//{W}style[@{W}styleId='SarjTitle']")
    assert title_style is not None
    assert title_style.find(f".//{W}pBdr") is None
    assert header.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip") is not None
    header_alignment = header.find(f".//{W}jc")
    assert header_alignment is not None
    assert header_alignment.get(f"{W}val") == "right"
    header_tokens = document_tokens.header
    assert header_tokens is not None
    header_spacing = header.find(f".//{W}spacing")
    assert header_spacing is not None
    assert header_spacing.get(f"{W}before") == str(header_tokens.space_above_pt * 20)
    assert header_spacing.get(f"{W}after") == str(header_tokens.space_below_pt * 20)
    extent = header.find(f".//{WP}extent")
    assert extent is not None
    asset_width = CURRENT_BRAND_MANIFEST.assets.logo_pixel_width
    asset_height = CURRENT_BRAND_MANIFEST.assets.logo_pixel_height
    assert asset_width is not None
    assert asset_height is not None
    scale = min(
        header_tokens.logo_width_pt / asset_width, header_tokens.logo_height_pt / asset_height
    )
    assert int(extent.get("cx", "0")) == pytest.approx(asset_width * scale * 12_700, abs=1)
    assert int(extent.get("cy", "0")) == pytest.approx(asset_height * scale * 12_700, abs=1)
    page_field = footer.find(f".//{W}instrText")
    assert page_field is not None
    assert page_field.text is not None
    assert page_field.text.strip() == "PAGE"
    margins = document.find(f".//{W}pgMar")
    assert margins is not None
    assert (
        margins.get(f"{W}left")
        == margins.get(f"{W}right")
        == str(document_tokens.horizontal_margin_pt * 20)
    )
    assert (
        margins.get(f"{W}top")
        == margins.get(f"{W}bottom")
        == str(document_tokens.vertical_margin_pt * 20)
    )


def test_logo_authors_explicit_alt_title_and_description(rendered: bytes) -> None:
    with ZipFile(BytesIO(rendered)) as archive:
        header = fromstring(archive.read("word/header1.xml"))
    logo_properties = header.find(f".//{WP}docPr")
    assert logo_properties is not None
    assert logo_properties.get("title") == "Sarj logo"
    assert logo_properties.get("name") == "Sarj logo"
    assert logo_properties.get("descr") == docx_renderer.trusted_document_logo_description()


def test_custom_logo_is_a_native_header_picture() -> None:
    rendered = render_markdown_docx("Body", title="Custom", logo_bytes=LOGO_BYTES)
    with ZipFile(BytesIO(rendered)) as archive:
        assert archive.read("word/media/image1.png") == LOGO_BYTES


def test_missing_logo_does_not_fabricate_brand_artwork() -> None:
    rendered = render_markdown_docx("Body", title="No logo")
    with ZipFile(BytesIO(rendered)) as archive:
        assert not any(name.startswith("word/media/") for name in archive.namelist())


@pytest.mark.parametrize(
    "markdown",
    [
        "![architecture](https://example.com/diagram.png)",
        "<aside>Decision</aside>",
        "A claim with a footnote[^1].",
        "~~obsolete~~",
        ">> Nested quotation",
        "---\ntitle: Draft\n---\nBody",
    ],
)
def test_rejects_markdown_constructs_that_cannot_be_preserved(markdown: str) -> None:
    with pytest.raises(UnsupportedSourceError):
        render_markdown_docx(markdown, title="Unsupported")


def test_rejects_xml_control_characters_and_unsafe_links() -> None:
    with pytest.raises(UnsupportedSourceError):
        render_markdown_docx("bad\x00text", title="Bad")
    with pytest.raises(UnsupportedSourceError):
        render_markdown_docx("[click](javascript:alert)", title="Unsafe link")


@pytest.mark.parametrize(
    "markdown",
    [
        "[**bold link**](https://example.com)",
        "**bold [link](https://example.com)**",
        "[*italic link*](https://example.com)",
        "*[italic link](https://example.com)*",
    ],
)
def test_nested_link_and_emphasis_are_preserved_exactly(markdown: str) -> None:
    manifest = markdown_import_manifest(markdown, title="Nested inline", branded=True)
    rendered = render_markdown_docx(markdown, title="Nested inline")
    with ZipFile(BytesIO(rendered)) as archive:
        document = fromstring(archive.read("word/document.xml"))
        relationships = fromstring(archive.read("word/_rels/document.xml.rels"))

    paragraph = document.findall(f".//{W}body/{W}p")[1]
    assert _xml_paragraph_text(paragraph) == manifest.blocks[1].text
    assert paragraph.find(f".//{W}hyperlink") is not None
    assert paragraph.find(f".//{W}b") is not None or paragraph.find(f".//{W}i") is not None
    assert manifest.blocks[1].links[0][1] == "https://example.com"
    assert "https://example.com" in [item.attrib.get("Target") for item in relationships]


def test_escaped_emphasis_inside_link_remains_literal() -> None:
    rendered = render_markdown_docx(
        r"[\*literal\* and snake_case](https://example.com)",
        title="Escaped link label",
    )
    with ZipFile(BytesIO(rendered)) as archive:
        document = fromstring(archive.read("word/document.xml"))
        relationships = fromstring(archive.read("word/_rels/document.xml.rels"))

    body_paragraphs = document.findall(f".//{W}body/{W}p")
    assert _xml_paragraph_text(body_paragraphs[1]) == "*literal* and snake_case"
    assert body_paragraphs[1].find(f".//{W}hyperlink") is not None
    italic = body_paragraphs[1].find(f".//{W}i")
    assert italic is None or italic.get(f"{W}val") in {"0", "false"}
    assert "https://example.com" in [item.attrib.get("Target") for item in relationships]


def test_balanced_parentheses_link_is_preserved_from_commonmark_tokens() -> None:
    markdown = "[Wiki](https://en.wikipedia.org/wiki/Function_(mathematics))"
    manifest = markdown_import_manifest(markdown, title="Balanced destination", branded=True)
    rendered = render_markdown_docx(markdown, title="Balanced destination")
    with ZipFile(BytesIO(rendered)) as archive:
        relationships = fromstring(archive.read("word/_rels/document.xml.rels"))
    target = "https://en.wikipedia.org/wiki/Function_(mathematics)"
    assert manifest.blocks[1].links == (("Wiki", target),)
    assert target in [item.attrib.get("Target") for item in relationships]


def test_supported_link_manifest_matches_independent_ooxml_readback() -> None:
    markdown = "Before [Wiki](https://example.com/path?q=1&x=2) after."
    manifest = markdown_import_manifest(markdown, title="Link", branded=True)
    rendered = render_markdown_docx(markdown, title="Link")
    with ZipFile(BytesIO(rendered)) as archive:
        document = fromstring(archive.read("word/document.xml"))
        relationships = fromstring(archive.read("word/_rels/document.xml.rels"))

    body_paragraphs = document.findall(f".//{W}body/{W}p")
    assert manifest.blocks[1].text == "Before Wiki after."
    assert manifest.blocks[1].links == (("Wiki", "https://example.com/path?q=1&x=2"),)
    assert _xml_paragraph_text(body_paragraphs[1]) == manifest.blocks[1].text
    assert "https://example.com/path?q=1&x=2" in [
        item.attrib.get("Target") for item in relationships
    ]


@pytest.mark.parametrize("indent", ["      ", "\t  "])
def test_rejects_list_nesting_deeper_than_supported(indent: str) -> None:
    markdown = f"- Level zero\n  - Level one\n    - Level two\n{indent}- Level three"
    with pytest.raises(UnsupportedSourceError):
        render_markdown_docx(markdown, title="Deep list")


def test_first_strong_rtl_sets_native_direction_without_changing_mixed_text() -> None:
    title = "مراجعة Sarj"
    markdown = (
        "# عنوان عربي\n\n"
        "👩🏽‍💻 العربية with Latin.\n\n"
        "English before \u2067نص عربي ٤٥٦\u2069 after."
    )
    rendered = render_markdown_docx(markdown, title=title)
    assert rendered == render_markdown_docx(markdown, title=title)
    with ZipFile(BytesIO(rendered)) as archive:
        document = fromstring(archive.read("word/document.xml"))

    paragraphs = document.findall(f".//{W}body/{W}p")
    observed = {_xml_paragraph_text(paragraph): paragraph for paragraph in paragraphs}
    assert set(observed) == {
        "عنوان عربي",
        "👩🏽‍💻 العربية with Latin.",
        "English before \u2067نص عربي ٤٥٦\u2069 after.",
    }
    for text in ("عنوان عربي", "👩🏽‍💻 العربية with Latin."):
        paragraph = observed[text]
        bidi = paragraph.find(f"./{W}pPr/{W}bidi")
        alignment = paragraph.find(f"./{W}pPr/{W}jc")
        run_rtl = paragraph.find(f".//{W}rPr/{W}rtl")
        run_fonts = paragraph.find(f".//{W}rPr/{W}rFonts")
        assert bidi is not None
        assert alignment is not None
        assert run_rtl is not None
        assert run_fonts is not None
        assert bidi.get(f"{W}val") == "true"
        assert alignment.get(f"{W}val") == "right"
        assert run_rtl.get(f"{W}val") == "true"
        assert run_fonts.get(f"{W}cs") == "Arial"
        assert run_fonts.get(f"{W}cstheme") is None

    ltr_paragraph = observed["English before \u2067نص عربي ٤٥٦\u2069 after."]
    assert ltr_paragraph.find(f"./{W}pPr/{W}bidi") is None
    assert ltr_paragraph.find(f"./{W}pPr/{W}jc") is None
    assert ltr_paragraph.find(f".//{W}rPr/{W}rtl") is None


@pytest.mark.parametrize(
    ("limit_name", "limit", "markdown"),
    [
        ("MAX_DOCX_TABLE_COLUMNS", 2, "| A | B | C |\n| --- | --- | --- |"),
        ("MAX_DOCX_TABLE_ROWS", 1, "| A | B |\n| --- | --- |\n| 1 | 2 |"),
        ("MAX_DOCX_TABLE_CELLS", 3, "| A | B |\n| --- | --- |\n| 1 | 2 |"),
        ("MAX_DOCX_INLINE_RUNS", 1, "plain **bold**"),
        ("MAX_DOCX_STYLED_SPANS", 0, "**bold**"),
        ("MAX_DOCX_HYPERLINKS", 0, "[link](https://sarj.ai)"),
        ("MAX_DOCX_ESTIMATED_XML_BYTES", 1, "plain"),
    ],
)
def test_complexity_limits_reject_before_word_document_allocation(
    monkeypatch: pytest.MonkeyPatch,
    limit_name: str,
    limit: int,
    markdown: str,
) -> None:
    monkeypatch.setattr(docx_renderer, limit_name, limit)

    def fail_allocation() -> None:
        pytest.fail("WordDocument must not be allocated before complexity admission")

    monkeypatch.setattr(docx_renderer, "WordDocument", fail_allocation)

    with pytest.raises(SourceTooLargeError):
        render_markdown_docx(markdown, title="Bounded")


def test_unbranded_render_preserves_structure_without_sarj_furniture() -> None:
    rendered = render_markdown_docx(
        "# Decision\n\n1. First\n2. Second\n\n| A | B |\n| --- | --- |\n| 1 | 2 |",
        title="Plain import",
        logo_bytes=LOGO_BYTES,
        branded=False,
    )
    document = WordDocument(BytesIO(rendered))

    assert document.paragraphs[0].style is not None
    assert document.paragraphs[0].style.name == "Heading 1"
    with ZipFile(BytesIO(rendered)) as archive:
        assert not any(name.startswith("word/header") for name in archive.namelist())
        assert not any(name.startswith("word/footer") for name in archive.namelist())
        assert not any(name.startswith("word/media/") for name in archive.namelist())
        styles_bytes = archive.read("word/styles.xml")
    assert b"Nunito" not in styles_bytes
    assert b"Sarj Title" not in styles_bytes
    styles = fromstring(styles_bytes)
    for style_id in ("Normal", "Title", *(f"Heading{level}" for level in range(1, 7))):
        style = styles.find(f".//{W}style[@{W}styleId='{style_id}']")
        assert style is not None
        fonts = style.find(f"./{W}rPr/{W}rFonts")
        assert fonts is not None
        assert {
            fonts.get(f"{W}ascii"),
            fonts.get(f"{W}hAnsi"),
            fonts.get(f"{W}eastAsia"),
            fonts.get(f"{W}cs"),
        } == {"Arial"}
        assert not any(attribute.endswith("Theme") for attribute in fonts.attrib)
    title_style = styles.find(f".//{W}style[@{W}styleId='Title']")
    assert title_style is not None
    assert title_style.find(f"./{W}pPr/{W}pBdr") is None
    title_fonts = title_style.find(f"./{W}rPr/{W}rFonts")
    assert title_fonts is not None
    assert title_fonts.get(f"{W}ascii") == "Arial"
    assert title_fonts.get(f"{W}hAnsi") == "Arial"
    assert title_fonts.get(f"{W}eastAsia") == "Arial"
    assert title_fonts.get(f"{W}cs") == "Arial"
    assert not any(attribute.endswith("Theme") for attribute in title_fonts.attrib)
    title_color = title_style.find(f"./{W}rPr/{W}color")
    assert title_color is not None
    assert title_color.get(f"{W}val") == "000000"
    assert not any("theme" in attribute.casefold() for attribute in title_color.attrib)


def test_generated_property_children_follow_wordprocessingml_schema_order() -> None:
    rendered = render_markdown_docx(
        "> Quoted text\n\n---\n\n| A | B |\n| --- | --- |\n| 1 | 2 |",
        title="OOXML order",
    )
    with ZipFile(BytesIO(rendered)) as archive:
        document = fromstring(archive.read("word/document.xml"))

    paragraphs = document.findall(f".//{W}body/{W}p")
    quote_properties = next(
        properties
        for paragraph in paragraphs
        if (properties := paragraph.find(f"./{W}pPr")) is not None
        and properties.find(f"./{W}pStyle[@{W}val='SarjQuote']") is not None
    )
    assert [child.tag.removeprefix(W) for child in quote_properties] == ["pStyle", "pBdr", "shd"]

    rule_properties = next(
        properties
        for paragraph in paragraphs
        if (properties := paragraph.find(f"./{W}pPr")) is not None
        and properties.find(f"./{W}pBdr/{W}bottom") is not None
    )
    assert [child.tag.removeprefix(W) for child in rule_properties] == ["pBdr", "spacing"]

    table_properties = document.find(f".//{W}tbl/{W}tblPr")
    assert table_properties is not None
    assert [child.tag.removeprefix(W) for child in table_properties] == [
        "tblW",
        "jc",
        "tblInd",
        "tblBorders",
        "tblLayout",
        "tblCellMar",
        "tblLook",
    ]
    header_cell_properties = document.find(f".//{W}tbl/{W}tr/{W}tc/{W}tcPr")
    assert header_cell_properties is not None
    assert [child.tag.removeprefix(W) for child in header_cell_properties] == [
        "tcW",
        "shd",
        "vAlign",
    ]


def test_matching_leading_h1_does_not_duplicate_visible_title() -> None:
    rendered = render_markdown_docx(
        "# Technology review\n\nBody",
        title="Technology review",
    )
    document = WordDocument(BytesIO(rendered))
    assert [paragraph.text for paragraph in document.paragraphs].count("Technology review") == 1


@pytest.mark.parametrize("title", ["Authored Ω", "Polish Authored Ω", "Different name", ""])
@pytest.mark.parametrize("branded", [True, False])
def test_authored_title_preserves_inline_semantics_and_independent_metadata(
    title: str, branded: bool
) -> None:
    source = "# *Authored* [Ω](https://example.com)\n\nBody"
    rendered = render_markdown_docx(source, title=title, branded=branded)
    assert render_markdown_docx(source, title=title, branded=branded) == rendered
    document = WordDocument(BytesIO(rendered))
    assert document.core_properties.title == title
    manifest = markdown_import_manifest(source, title=title, branded=branded)
    assert [block.kind for block in manifest.blocks] == ["heading", "paragraph"]
    assert manifest.blocks[0].text == "Authored Ω"
    assert manifest.blocks[0].level == 1
    assert manifest.blocks[0].links == (("Ω", "https://example.com"),)
    with ZipFile(BytesIO(rendered)) as archive:
        body = fromstring(archive.read("word/document.xml"))
    paragraphs = body.findall(f".//{W}body/{W}p")
    assert [_xml_paragraph_text(paragraph) for paragraph in paragraphs] == ["Authored Ω", "Body"]
    assert paragraphs[0].find(f".//{W}i") is not None
    assert paragraphs[0].find(f".//{W}hyperlink") is not None


@pytest.mark.parametrize("source", ["Body\n\n# Later", "Body", ""])
def test_nonleading_h1_keeps_generated_display_title(source: str) -> None:
    manifest = markdown_import_manifest(source, title="Display", branded=False)
    assert manifest.blocks[0].kind == "title"
    document = WordDocument(BytesIO(render_markdown_docx(source, title="Display", branded=False)))
    assert document.paragraphs[0].text == "Display"


def test_separate_ordered_lists_restart_at_one() -> None:
    rendered = render_markdown_docx(
        "1. First\n2. Second\n\nParagraph\n\n1. Restarted",
        title="List restart",
    )
    with ZipFile(BytesIO(rendered)) as archive:
        document = fromstring(archive.read("word/document.xml"))
        numbering = fromstring(archive.read("word/numbering.xml"))

    list_paragraphs = [
        paragraph
        for paragraph in document.findall(f".//{W}body/{W}p")
        if paragraph.find(f".//{W}numId") is not None
    ]
    num_ids: list[str | None] = []
    for paragraph in list_paragraphs:
        num_id = paragraph.find(f".//{W}numId")
        assert num_id is not None
        num_ids.append(num_id.get(f"{W}val"))

    assert num_ids[0] == num_ids[1]
    assert num_ids[2] != num_ids[1]
    restarted = numbering.find(f".//{W}num[@{W}numId='{num_ids[2]}']")
    assert restarted is not None
    start_override = restarted.find(f".//{W}startOverride")
    assert start_override is not None
    assert start_override.get(f"{W}val") == "1"


def test_ordered_list_preserves_explicit_start() -> None:
    rendered = render_markdown_docx("5. Fifth\n6. Sixth", title="List start")
    with ZipFile(BytesIO(rendered)) as archive:
        document = fromstring(archive.read("word/document.xml"))
        numbering = fromstring(archive.read("word/numbering.xml"))

    first_num_id = document.find(f".//{W}body/{W}p/{W}pPr/{W}numPr/{W}numId")
    assert first_num_id is not None
    concrete = numbering.find(f".//{W}num[@{W}numId='{first_num_id.get(f'{W}val')}']")
    assert concrete is not None
    start = concrete.find(f".//{W}startOverride")
    assert start is not None
    assert start.get(f"{W}val") == "5"


@pytest.mark.parametrize(
    "markdown",
    ["Line one  \nLine two", "Line one\\\nLine two"],
    ids=["two-spaces", "backslash"],
)
def test_markdown_hard_breaks_are_native_line_breaks(markdown: str) -> None:
    rendered = render_markdown_docx(markdown, title="Hard break")
    document = WordDocument(BytesIO(rendered))

    assert document.paragraphs[1].text == "Line one\nLine two"


def test_backslash_escaped_emphasis_stays_literal() -> None:
    rendered = render_markdown_docx(
        r"Use \*literal asterisks\* exactly.",
        title="Escaped punctuation",
    )
    document = WordDocument(BytesIO(rendered))
    paragraph = document.paragraphs[1]

    assert paragraph.text == "Use *literal asterisks* exactly."
    assert not any(run.italic for run in paragraph.runs)


def test_setext_headings_remain_headings_without_extra_rules() -> None:
    rendered = render_markdown_docx(
        "Primary heading\n===\n\nSecondary heading\n---",
        title="Setext",
    )
    document = WordDocument(BytesIO(rendered))

    observed: list[tuple[str, str]] = []
    for paragraph in document.paragraphs:
        style = paragraph.style
        assert style is not None
        assert style.name is not None
        observed.append((paragraph.text, style.name))
    assert observed == [
        ("Primary heading", "Heading 1"),
        ("Secondary heading", "Heading 2"),
    ]


@pytest.mark.parametrize(
    ("markdown", "expected_text", "expected_style"),
    [
        ("~~~python\nprint('safe')\n~~~", "print('safe')", "Sarj Code"),
        ("    indented code", "indented code", "Sarj Code"),
        ("[Sarj][site]\n\n[site]: https://sarj.ai", "Sarj", "Normal"),
        ("This is ***bold and italic*** text.", "This is bold and italic text.", "Normal"),
    ],
)
def test_commonmark_safe_constructs_are_preserved(
    markdown: str,
    expected_text: str,
    expected_style: str,
) -> None:
    manifest = markdown_import_manifest(markdown, title="Supported", branded=True)
    rendered = render_markdown_docx(markdown, title="Supported")
    document = WordDocument(BytesIO(rendered))

    paragraph = document.paragraphs[1]
    assert paragraph.text == expected_text
    assert paragraph.style is not None
    assert paragraph.style.name == expected_style
    assert manifest.blocks[1].text == expected_text


@pytest.mark.parametrize("markdown", ["## Starts too deep", "# Root\n\n### Skips a level"])
def test_invalid_heading_hierarchy_rejects_before_allocation(
    monkeypatch: pytest.MonkeyPatch,
    markdown: str,
) -> None:
    def fail_allocation() -> None:
        pytest.fail("WordDocument must not be allocated for an invalid heading hierarchy")

    monkeypatch.setattr(docx_renderer, "WordDocument", fail_allocation)
    with pytest.raises(UnsupportedSourceError):
        render_markdown_docx(markdown, title="Visual title is not a semantic heading")
    with pytest.raises(UnsupportedSourceError):
        markdown_import_manifest(
            markdown, title="Visual title is not a semantic heading", branded=True
        )


def test_valid_heading_hierarchy_matches_manifest_and_ooxml() -> None:
    markdown = "# Root\n\n## Child one\n\n## Child two\n\n# New root"
    manifest = markdown_import_manifest(markdown, title="Hierarchy", branded=True)
    rendered = render_markdown_docx(markdown, title="Hierarchy")
    with ZipFile(BytesIO(rendered)) as archive:
        document = fromstring(archive.read("word/document.xml"))

    headings = [
        (_xml_paragraph_text(paragraph), paragraph_style.get(f"{W}val"))
        for paragraph in document.findall(f".//{W}body/{W}p")
        if (paragraph_style := paragraph.find(f"./{W}pPr/{W}pStyle")) is not None
        and paragraph_style.get(f"{W}val", "").startswith("Heading")
    ]
    assert headings == [
        ("Root", "Heading1"),
        ("Child one", "Heading2"),
        ("Child two", "Heading2"),
        ("New root", "Heading1"),
    ]
    assert [(block.text, block.level) for block in manifest.blocks] == [
        ("Root", 1),
        ("Child one", 2),
        ("Child two", 2),
        ("New root", 1),
    ]


def _commonmark_visible_text(markdown: str) -> str:
    tokens = MarkdownIt("commonmark").parse(markdown)
    children = next(token.children for token in tokens if token.type == "inline")
    assert children is not None
    return "".join(
        " " if token.type == "softbreak" else "\n" if token.type == "hardbreak" else token.content
        for token in children
        if token.type in {"text", "code_inline", "softbreak", "hardbreak"}
    )


def test_seeded_safe_inline_corpus_matches_commonmark_and_ooxml() -> None:
    # Values are checked independently through markdown-it and WordprocessingML
    # rather than through the import manifest.
    corpus = [
        "foo_bar_baz and a__b__c",
        "A &amp; B &#169; Sarj",
        r"Escaped \*stars\* and \# marker",
        "Use `` `padded` `` and `one`.",
        'A [link](https://example.com/a_(b)?x=1&amp;y=2 "Title") here.',
        "Combining cafe\u0301 and emoji 👩🏽‍💻",
        "first line\ncontinuation",
        "first line  \ncontinuation",
    ]
    atoms = (
        "plain",
        "snake_case",
        "**bold**",
        "*italic*",
        "`code`",
        "&amp;",
        "[link](https://example.com/a_(b))",
        "العربية",
        "👩🏽‍💻",
    )
    seed = 260904
    for _ in range(24):
        selected: list[str] = []
        for _ in range(4):
            seed = (1_103_515_245 * seed + 12_345) % (2**31)
            selected.append(atoms[seed % len(atoms)])
        corpus.append(" ".join(selected))
    for index, markdown in enumerate(corpus):
        expected = _commonmark_visible_text(markdown)
        rendered = render_markdown_docx(markdown, title=f"Seed 260904-{index}")
        with ZipFile(BytesIO(rendered)) as archive:
            document = fromstring(archive.read("word/document.xml"))
        body = document.findall(f".//{W}body/{W}p")[1]
        assert _xml_paragraph_text(body) == expected.replace("\n", "")
        assert len(body.findall(f".//{W}br")) == expected.count("\n")


def test_list_continuation_is_one_native_list_paragraph() -> None:
    markdown = "- first line\n  continuation line"
    manifest = markdown_import_manifest(markdown, title="Continuation", branded=True)
    rendered = render_markdown_docx(markdown, title="Continuation")
    with ZipFile(BytesIO(rendered)) as archive:
        document = fromstring(archive.read("word/document.xml"))

    list_paragraphs = [
        paragraph
        for paragraph in document.findall(f".//{W}body/{W}p")
        if paragraph.find(f"./{W}pPr/{W}numPr") is not None
    ]
    assert len(list_paragraphs) == 1
    assert _xml_paragraph_text(list_paragraphs[0]) == "first line continuation line"
    assert manifest.blocks[1].text == "first line continuation line"


@pytest.mark.parametrize("markdown", ["* * *", "- - -"])
def test_spaced_thematic_rules_are_native_rules(markdown: str) -> None:
    manifest = markdown_import_manifest(markdown, title="Rule", branded=True)
    rendered = render_markdown_docx(markdown, title="Rule")
    with ZipFile(BytesIO(rendered)) as archive:
        document = fromstring(archive.read("word/document.xml"))
    rule = document.findall(f".//{W}body/{W}p")[1]
    assert rule.find(f"./{W}pPr/{W}pBdr/{W}bottom") is not None
    assert manifest.blocks[1].kind == "rule"


def test_link_title_and_entity_decoded_destination_are_exact() -> None:
    markdown = '[Wiki &amp; Docs](https://example.com/a_(b)?x=1&amp;y=2 "Reference")'
    target = "https://example.com/a_(b)?x=1&y=2"
    manifest = markdown_import_manifest(markdown, title="Link", branded=True)
    rendered = render_markdown_docx(markdown, title="Link")
    with ZipFile(BytesIO(rendered)) as archive:
        relationships = fromstring(archive.read("word/_rels/document.xml.rels"))
        document = fromstring(archive.read("word/document.xml"))
    assert manifest.blocks[1].text == "Wiki & Docs"
    assert manifest.blocks[1].links == (("Wiki & Docs", target),)
    assert _xml_paragraph_text(document.findall(f".//{W}body/{W}p")[1]) == "Wiki & Docs"
    assert target in [item.attrib.get("Target") for item in relationships]


def test_escaped_and_unspaced_atx_markers_follow_commonmark() -> None:
    markdown = "# Heading \\#\n\n#Heading\n\n# Heading###"
    manifest = markdown_import_manifest(markdown, title="ATX", branded=True)
    rendered = render_markdown_docx(markdown, title="ATX")
    document = WordDocument(BytesIO(rendered))
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Heading #",
        "#Heading",
        "Heading###",
    ]
    assert [block.kind for block in manifest.blocks] == ["heading", "paragraph", "heading"]
