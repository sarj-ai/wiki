from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from io import BytesIO
import re
from typing import Final, Literal, override
from urllib.parse import urlparse
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document as WordDocument
from docx.document import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Inches, Pt, RGBColor
from docx.styles.style import BaseStyle, ParagraphStyle
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from markdown_it import MarkdownIt
from markdown_it.token import Token

from app.document_transform.brand_manifest import (
    CURRENT_BRAND_MANIFEST,
    HEADER_TITLE_SEPARATOR,
    DocumentHeaderBrandTokens,
    normalize_header_title,
    trusted_document_logo_description,
)
from app.document_transform.errors import (
    EmptyMarkdownBlockError,
    SourceTooLargeError,
    UnsupportedSourceError,
)
from app.document_transform.import_manifest import (
    MarkdownImportBlock,
    MarkdownImportInline,
    MarkdownImportManifest,
)
from app.document_transform.limits import (
    MAX_DOCX_ESTIMATED_XML_BYTES,
    MAX_DOCX_HYPERLINKS,
    MAX_DOCX_INLINE_RUNS,
    MAX_DOCX_STYLED_SPANS,
    MAX_DOCX_TABLE_CELLS,
    MAX_DOCX_TABLE_COLUMNS,
    MAX_DOCX_TABLE_ROWS,
)
from app.document_transform.text_direction import first_strong_is_rtl


_DOCUMENT = CURRENT_BRAND_MANIFEST.document
_PAGE_WIDTH_PT: Final = 612
_PAGE_HEIGHT_PT: Final = 792
_CONTENT_WIDTH_DXA: Final = int((_PAGE_WIDTH_PT - 2 * _DOCUMENT.horizontal_margin_pt) * 20)
_MAX_LIST_LEVEL: Final = 2
_SARJ_PURPLE: Final = CURRENT_BRAND_MANIFEST.colors.primary.removeprefix("#")
_SARJ_PURPLE_LIGHT: Final = CURRENT_BRAND_MANIFEST.colors.primary_tint.removeprefix("#")
_INK: Final = CURRENT_BRAND_MANIFEST.colors.ink.removeprefix("#")
_MUTED: Final = CURRENT_BRAND_MANIFEST.colors.muted.removeprefix("#")
_BORDER: Final = CURRENT_BRAND_MANIFEST.colors.border.removeprefix("#")
_CODE_FILL: Final = CURRENT_BRAND_MANIFEST.colors.subtle.removeprefix("#")
_FIXED_TIMESTAMP: Final = datetime(2000, 1, 1, tzinfo=UTC)
_INVALID_XML_CHARACTER = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


class _CommonMarkParser(MarkdownIt):
    @override
    def validateLink(self, url: str) -> bool:
        # Parse every destination, then apply our stricter scheme allowlist while
        # building the IR. Otherwise markdown-it renders an unsafe link as plain
        # text and the admission layer cannot distinguish it from literal text.
        del url
        return True


_MARKDOWN = _CommonMarkParser("commonmark").enable("table").enable("strikethrough")


def _current_header_tokens() -> DocumentHeaderBrandTokens:
    tokens = _DOCUMENT.header
    if tokens is None:  # pragma: no cover - immutable manifest invariant
        msg = "Current Sarj brand profile must define document header tokens"
        raise RuntimeError(msg)
    return tokens


def _current_complex_script_font() -> str:
    font = _DOCUMENT.complex_script_font_family
    if font is None:  # pragma: no cover - immutable manifest invariant
        msg = "Current Sarj brand profile must define a complex-script font"
        raise RuntimeError(msg)
    return font


_HEADER = _current_header_tokens()
_COMPLEX_SCRIPT_FONT = _current_complex_script_font()


@dataclass(frozen=True, slots=True)
class _Block:
    kind: Literal["heading", "paragraph", "list", "check", "quote", "code", "table", "rule"]
    text: str = ""
    level: int = 0
    ordered: bool = False
    start: int = 1
    checked: bool = False
    rows: tuple[tuple[str, ...], ...] = ()
    inlines: tuple[MarkdownImportInline, ...] = ()
    cell_inlines: tuple[tuple[tuple[MarkdownImportInline, ...], ...], ...] = ()
    has_header: bool = False


@dataclass(frozen=True, slots=True)
class _NumberingIds:
    bullet: int
    ordered: int
    unchecked: int
    checked: int
    ordered_abstract: int
    next_num: int


@dataclass(frozen=True, slots=True)
class _LogoDimensions:
    width: float
    height: float


@dataclass(frozen=True, slots=True)
class _CheckMarker:
    checked: bool | None
    inlines: tuple[MarkdownImportInline, ...]


def render_markdown_docx(
    markdown: str,
    *,
    title: str,
    logo_bytes: bytes | None = None,
    branded: bool = True,
    header_title: str | None = None,
) -> bytes:
    blocks = _validated_blocks(markdown, title=title)

    document = WordDocument()
    _set_metadata(document, title)
    _configure_page(document, branded=branded)
    _configure_styles(document, branded=branded)
    numbering = _configure_numbering(document, branded=branded)
    if branded and logo_bytes is not None:
        _configure_header(
            document, logo_bytes, title=title if header_title is None else header_title
        )
    if branded:
        _configure_footer(document)

    if _include_display_title(blocks, title):
        paragraph = document.add_paragraph(style="Sarj Title" if branded else "Title")
        _append_inlines(paragraph, _plain_inlines(title.strip()), branded=branded)
    _append_blocks(document, blocks, numbering, branded=branded)

    output = BytesIO()
    document.save(output)
    return _canonicalize_archive(output.getvalue())


def validate_markdown_import(markdown: str, *, title: str) -> None:
    _validated_blocks(markdown, title=title)


def _validated_blocks(markdown: str, *, title: str) -> tuple[_Block, ...]:
    _validate_text(markdown, field="Markdown")
    _validate_text(title, field="Title")
    normalized_lines = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if _has_frontmatter(normalized_lines) or _has_unsupported_footnote(markdown):
        raise UnsupportedSourceError
    blocks = _parse_blocks(markdown)
    if any(block.kind in {"heading", "code"} and not block.text for block in blocks):
        raise EmptyMarkdownBlockError(EmptyMarkdownBlockError.public_detail)
    _validate_heading_hierarchy(blocks)
    _validate_render_complexity(blocks, title=title, markdown=markdown)
    return blocks


def _validate_render_complexity(
    blocks: tuple[_Block, ...],
    *,
    title: str,
    markdown: str,
) -> None:
    table_cells, inline_runs, styled_spans, hyperlinks = _render_complexity_counts(blocks, title)
    if (
        table_cells > MAX_DOCX_TABLE_CELLS
        or inline_runs > MAX_DOCX_INLINE_RUNS
        or styled_spans > MAX_DOCX_STYLED_SPANS
        or hyperlinks > MAX_DOCX_HYPERLINKS
    ):
        raise SourceTooLargeError
    estimated_xml_bytes = (
        len(markdown.encode("utf-8")) * 4
        + len(title.encode("utf-8")) * 4
        + len(blocks) * 1_024
        + table_cells * 512
        + inline_runs * 256
        + hyperlinks * 512
    )
    if estimated_xml_bytes > MAX_DOCX_ESTIMATED_XML_BYTES:
        raise SourceTooLargeError


def _render_complexity_counts(blocks: tuple[_Block, ...], title: str) -> tuple[int, int, int, int]:
    table_cells = 0
    inline_runs = 0
    styled_spans = 0
    hyperlinks = 0

    def count_inlines(inlines: tuple[MarkdownImportInline, ...]) -> None:
        nonlocal inline_runs, styled_spans, hyperlinks
        inline_runs += len(inlines) + sum(inline.text.count("\n") for inline in inlines)
        styled_spans += sum(
            inline.bold or inline.italic or inline.code or inline.url is not None
            for inline in inlines
        )
        hyperlinks += sum(inline.url is not None for inline in inlines)

    if title.strip():
        count_inlines(_plain_inlines(title.strip()))
    for block in blocks:
        match block.kind:
            case "table":
                rows = len(block.rows)
                columns = max((len(row) for row in block.rows), default=0)
                cells = rows * columns
                if rows > MAX_DOCX_TABLE_ROWS or columns > MAX_DOCX_TABLE_COLUMNS:
                    raise SourceTooLargeError
                table_cells += cells
                for row in block.cell_inlines:
                    for cell in row:
                        count_inlines(cell)
            case "code":
                # Each hard line break is materialized as its own run.
                inline_runs += max(1, 2 * len(block.text.split("\n")) - 1)
            case _:
                count_inlines(block.inlines)
    return table_cells, inline_runs, styled_spans, hyperlinks


def markdown_import_manifest(
    markdown: str,
    *,
    title: str,
    branded: bool,
    header_title: str | None = None,
) -> MarkdownImportManifest:
    parsed = _validated_blocks(markdown, title=title)
    blocks: list[MarkdownImportBlock] = []
    if _include_display_title(parsed, title):
        blocks.append(
            _manifest_block(
                _Block(
                    "heading",
                    text=title.strip(),
                    level=1,
                    inlines=_plain_inlines(title.strip()),
                ),
                kind="title",
            )
        )
    active_ordered: dict[int, int] = {}
    next_list_group = 1
    for block in parsed:
        start, list_group, next_list_group = _manifest_list_position(
            block, active_ordered, next_list_group
        )
        blocks.append(_manifest_block(block, start=start, list_group=list_group))
    return MarkdownImportManifest(
        blocks=tuple(blocks),
        branded=branded,
        header_title=normalize_header_title(
            title if header_title is None else header_title, max_chars=_HEADER.max_title_chars
        )
        if branded
        else None,
    )


def _manifest_list_position(
    block: _Block, active_ordered: dict[int, int], next_list_group: int
) -> tuple[int, int, int]:
    if block.kind == "list" and block.ordered:
        for level in tuple(active_ordered):
            if level > block.level:
                del active_ordered[level]
        start = block.start if block.level not in active_ordered else 0
        if block.level not in active_ordered:
            active_ordered[block.level] = next_list_group
            next_list_group += 1
        return start, active_ordered[block.level], next_list_group
    if block.kind in {"list", "check"}:
        for level in tuple(active_ordered):
            if level >= block.level:
                del active_ordered[level]
    else:
        active_ordered.clear()
    return block.start, 0, next_list_group


def _manifest_block(
    block: _Block,
    *,
    kind: str | None = None,
    start: int | None = None,
    list_group: int = 0,
) -> MarkdownImportBlock:
    inlines = _manifest_flat_inlines(block)
    signature = inlines if block.kind != "table" else ()
    if block.kind == "code":
        signature = (MarkdownImportInline(block.text, code=True),)
    return MarkdownImportBlock(
        kind=kind or block.kind,
        text="".join(inline.text for inline in inlines) if inlines else block.text,
        level=block.level,
        ordered=block.ordered,
        start=block.start if start is None else start,
        list_group=list_group,
        rows=block.rows,
        links=_manifest_links(inlines),
        bold=tuple(inline.text for inline in inlines if inline.bold),
        italic=tuple(inline.text for inline in inlines if inline.italic),
        inline_code=tuple(inline.text for inline in inlines if inline.code),
        inline_signature=signature,
        cell_signatures=block.cell_inlines,
        checked=block.checked if block.kind == "check" else None,
    )


def _manifest_flat_inlines(block: _Block) -> tuple[MarkdownImportInline, ...]:
    if block.kind == "table":
        return tuple(inline for row in block.cell_inlines for cell in row for inline in cell)
    return block.inlines


def _manifest_links(inlines: tuple[MarkdownImportInline, ...]) -> tuple[tuple[str, str], ...]:
    return tuple(
        (inline.text, inline.url)
        for inline in inlines
        if inline.url is not None and _is_safe_link(inline.url)
    )


def _validate_text(value: str, *, field: str) -> None:
    if _INVALID_XML_CHARACTER.search(value):
        raise UnsupportedSourceError from ValueError(
            f"{field} contains characters that XML cannot represent"
        )


def _has_frontmatter(lines: list[str]) -> bool:
    if not lines or lines[0].strip() != "---":
        return False
    return any(line.strip() in {"---", "..."} for line in lines[1:])


def _has_unsupported_footnote(
    markdown: str,
) -> bool:
    return _contains_footnote_reference(_outside_code_blocks(markdown))


def _outside_code_blocks(markdown: str) -> str:
    # Footnote-looking text inside a code block is literal source. Exclude
    # fenced and indented code before scanning for unsupported references.
    visible_lines: list[str] = []
    fence_character = ""
    fence_length = 0
    for line in markdown.splitlines(keepends=True):
        marker = re.match(r" {0,3}(`{3,}|~{3,})", line)
        if fence_character:
            if (
                marker
                and marker.group(1)[0] == fence_character
                and len(marker.group(1)) >= fence_length
            ):
                tail = line[marker.end() :].strip()
                if not tail:
                    fence_character = ""
                    fence_length = 0
            continue
        if marker and (marker.group(1)[0] != "`" or "`" not in line[marker.end() :]):
            fence_character = marker.group(1)[0]
            fence_length = len(marker.group(1))
            continue
        if line.startswith(("    ", "\t")):
            continue
        visible_lines.append(line)
    return "".join(visible_lines)


def _contains_footnote_reference(markdown: str) -> bool:
    index = 0
    while index < len(markdown):
        if markdown[index] == "\\":
            index += 2
            continue
        if markdown[index] == "`":
            delimiter_end = index + 1
            while delimiter_end < len(markdown) and markdown[delimiter_end] == "`":
                delimiter_end += 1
            delimiter = markdown[index:delimiter_end]
            closing = markdown.find(delimiter, delimiter_end)
            index = delimiter_end if closing < 0 else closing + len(delimiter)
            continue
        if markdown.startswith("[^", index):
            closing = markdown.find("]", index + 2)
            if closing > index + 2:
                return True
        index += 1
    return False


def _set_metadata(document: Document, title: str) -> None:
    properties = document.core_properties
    properties.title = title
    properties.author = "Sarj Wiki"
    properties.last_modified_by = "Sarj Wiki"
    properties.created = _FIXED_TIMESTAMP
    properties.modified = _FIXED_TIMESTAMP
    properties.revision = 1


def _configure_page(document: Document, *, branded: bool) -> None:
    section = document.sections[0]
    section.page_width = Pt(_PAGE_WIDTH_PT)
    section.page_height = Pt(_PAGE_HEIGHT_PT)
    vertical_margin = _DOCUMENT.vertical_margin_pt if branded else 72
    horizontal_margin = _DOCUMENT.horizontal_margin_pt if branded else 72
    section.top_margin = Pt(vertical_margin)
    section.right_margin = Pt(horizontal_margin)
    section.bottom_margin = Pt(vertical_margin)
    section.left_margin = Pt(horizontal_margin)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)


def _configure_styles(document: Document, *, branded: bool) -> None:
    styles = document.styles
    if not branded:
        for style_name in ("Normal", "Title", *(f"Heading {level}" for level in range(1, 7))):
            _set_style_font_family(
                _paragraph_style(styles[style_name]),
                font="Arial",
                complex_script_font="Arial",
            )
        title = _paragraph_style(styles["Title"])
        _set_style(
            title,
            font="Arial",
            size=24,
            color="000000",
            bold=True,
            after=12,
            keep_with_next=True,
        )
        _remove_unbranded_title_artifacts(title)
        return
    _set_style(
        _paragraph_style(styles["Normal"]),
        size=_DOCUMENT.body_size_pt,
        color=_INK,
        before=_DOCUMENT.paragraph_space_above_pt,
        after=_DOCUMENT.paragraph_space_below_pt,
        line_spacing=_DOCUMENT.line_spacing_percent / 100,
    )
    _set_style(
        _paragraph_style(styles["Heading 1"]),
        size=_DOCUMENT.h1_size_pt,
        color=_SARJ_PURPLE,
        bold=True,
        before=_DOCUMENT.paragraph_space_above_pt,
        after=_DOCUMENT.paragraph_space_below_pt,
        line_spacing=_DOCUMENT.line_spacing_percent / 100,
        keep_with_next=True,
    )
    _set_style(
        _paragraph_style(styles["Heading 2"]),
        size=_DOCUMENT.h2_size_pt,
        color=_SARJ_PURPLE,
        bold=True,
        before=_DOCUMENT.paragraph_space_above_pt,
        after=_DOCUMENT.paragraph_space_below_pt,
        line_spacing=_DOCUMENT.line_spacing_percent / 100,
        keep_with_next=True,
    )
    _set_style(
        _paragraph_style(styles["Heading 3"]),
        size=_DOCUMENT.h3_size_pt,
        color=_SARJ_PURPLE,
        bold=True,
        before=_DOCUMENT.paragraph_space_above_pt,
        after=_DOCUMENT.paragraph_space_below_pt,
        line_spacing=_DOCUMENT.line_spacing_percent / 100,
        keep_with_next=True,
    )
    for heading_level in range(4, 7):
        _set_style(
            _paragraph_style(styles[f"Heading {heading_level}"]),
            size=_DOCUMENT.h3_size_pt,
            color=_SARJ_PURPLE,
            bold=True,
            before=_DOCUMENT.paragraph_space_above_pt,
            after=_DOCUMENT.paragraph_space_below_pt,
            line_spacing=_DOCUMENT.line_spacing_percent / 100,
            keep_with_next=True,
        )
    _set_style(
        _paragraph_style(styles.add_style("Sarj Title", WD_STYLE_TYPE.PARAGRAPH)),
        size=_DOCUMENT.h1_size_pt,
        color=_SARJ_PURPLE,
        bold=True,
        before=_DOCUMENT.paragraph_space_above_pt,
        after=_DOCUMENT.paragraph_space_below_pt,
        line_spacing=_DOCUMENT.line_spacing_percent / 100,
        keep_with_next=True,
    )
    _set_style(
        _paragraph_style(styles.add_style("Sarj List", WD_STYLE_TYPE.PARAGRAPH)),
        size=_DOCUMENT.body_size_pt,
        color=_INK,
        before=_DOCUMENT.paragraph_space_above_pt,
        after=_DOCUMENT.paragraph_space_below_pt,
        line_spacing=_DOCUMENT.line_spacing_percent / 100,
    )
    quote = _paragraph_style(styles.add_style("Sarj Quote", WD_STYLE_TYPE.PARAGRAPH))
    _set_style(
        quote,
        size=_DOCUMENT.body_size_pt,
        color=_MUTED,
        italic=True,
        before=_DOCUMENT.paragraph_space_above_pt,
        after=_DOCUMENT.paragraph_space_below_pt,
        line_spacing=_DOCUMENT.line_spacing_percent / 100,
    )
    quote.paragraph_format.left_indent = Inches(0.25)
    quote.paragraph_format.right_indent = Inches(0.125)
    code = _paragraph_style(styles.add_style("Sarj Code", WD_STYLE_TYPE.PARAGRAPH))
    _set_style(
        code,
        font="Courier New",
        size=_DOCUMENT.body_size_pt,
        color=_INK,
        bold=False,
        before=_DOCUMENT.paragraph_space_above_pt,
        after=_DOCUMENT.paragraph_space_below_pt,
        line_spacing=_DOCUMENT.line_spacing_percent / 100,
    )
    code.paragraph_format.left_indent = Inches(0.125)
    code.paragraph_format.right_indent = Inches(0.125)
    _set_style(
        _paragraph_style(styles.add_style("Sarj Table Text", WD_STYLE_TYPE.PARAGRAPH)),
        size=_DOCUMENT.body_size_pt,
        color=_INK,
        before=_DOCUMENT.paragraph_space_above_pt,
        after=_DOCUMENT.paragraph_space_below_pt,
        line_spacing=_DOCUMENT.line_spacing_percent / 100,
    )


def _paragraph_style(style: BaseStyle) -> ParagraphStyle:
    if not isinstance(style, ParagraphStyle):
        msg = f"Expected a paragraph style, got {style.type!r}"
        raise TypeError(msg)
    return style


def _remove_unbranded_title_artifacts(style: ParagraphStyle) -> None:
    style_element = style.element
    if style_element is None:
        msg = f"Paragraph style {style.name!r} is detached"
        raise ValueError(msg)
    style_element.get_or_add_pPr().remove_all("w:pBdr")
    run_properties = style_element.get_or_add_rPr()
    fonts = run_properties.get_or_add_rFonts()
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attribute), None)
    color = run_properties.find(qn("w:color"))
    if color is not None:
        for attribute in ("w:themeColor", "w:themeTint", "w:themeShade"):
            color.attrib.pop(qn(attribute), None)


def _set_style(
    style: ParagraphStyle,
    *,
    size: float,
    color: str,
    after: float,
    font: str = _DOCUMENT.font_family,
    bold: bool | None = None,
    italic: bool | None = None,
    before: float = 0,
    line_spacing: float = 1.0,
    keep_with_next: bool = False,
) -> None:
    _set_style_font_family(
        style,
        font=font,
        complex_script_font=_COMPLEX_SCRIPT_FONT,
    )
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = italic
    style_element = style.element
    if style_element is None:
        msg = f"Paragraph style {style.name!r} is detached"
        raise ValueError(msg)
    paragraph_format = style.paragraph_format
    paragraph_format.space_before = Pt(before)
    paragraph_format.space_after = Pt(after)
    paragraph_format.line_spacing = line_spacing
    paragraph_format.keep_with_next = keep_with_next
    paragraph_format.widow_control = True


def _set_style_font_family(
    style: ParagraphStyle,
    *,
    font: str,
    complex_script_font: str,
) -> None:
    style.font.name = font
    style_element = style.element
    if style_element is None:
        msg = f"Paragraph style {style.name!r} is detached"
        raise ValueError(msg)
    fonts = style_element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), font)
    fonts.set(qn("w:hAnsi"), font)
    fonts.set(qn("w:eastAsia"), font)
    fonts.set(qn("w:cs"), complex_script_font)
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attribute), None)


def _configure_numbering(document: Document, *, branded: bool) -> _NumberingIds:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(element.get(qn("w:numId"))) for element in numbering.findall(qn("w:num"))]
    next_abstract = max(abstract_ids, default=-1) + 1
    next_num = max(num_ids, default=0) + 1
    values: list[int] = []
    schemes = (
        ("bullet", ("●", "○", "▪"), _DOCUMENT.font_family if branded else "Arial"),
        ("decimal", ("%1.", "%2.", "%3."), _DOCUMENT.font_family if branded else "Arial"),
        ("bullet", ("☐", "☐", "☐"), "DejaVu Sans"),
        ("bullet", ("☒", "☒", "☒"), "DejaVu Sans"),
    )
    for offset, (number_format, markers, marker_font) in enumerate(schemes):
        abstract_id = next_abstract + offset
        num_id = next_num + offset
        numbering.append(
            _abstract_numbering(
                abstract_id,
                number_format=number_format,
                markers=markers,
                marker_font=marker_font,
            )
        )
        concrete = OxmlElement("w:num")
        concrete.set(qn("w:numId"), str(num_id))
        reference = OxmlElement("w:abstractNumId")
        reference.set(qn("w:val"), str(abstract_id))
        concrete.append(reference)
        numbering.append(concrete)
        values.append(num_id)
    return _NumberingIds(
        bullet=values[0],
        ordered=values[1],
        unchecked=values[2],
        checked=values[3],
        ordered_abstract=next_abstract + 1,
        next_num=next_num + len(schemes),
    )


def _abstract_numbering(
    abstract_id: int,
    *,
    number_format: str,
    markers: tuple[str, str, str],
    marker_font: str,
) -> BaseOxmlElement:
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multilevel = OxmlElement("w:multiLevelType")
    multilevel.set(qn("w:val"), "hybridMultilevel")
    abstract.append(multilevel)
    for level, marker in enumerate(markers):
        level_element = OxmlElement("w:lvl")
        level_element.set(qn("w:ilvl"), str(level))
        _append_value_element(level_element, "w:start", "1")
        _append_value_element(level_element, "w:numFmt", number_format)
        _append_value_element(level_element, "w:lvlText", marker)
        _append_value_element(level_element, "w:lvlJc", "left")
        paragraph_properties = OxmlElement("w:pPr")
        indentation = OxmlElement("w:ind")
        indentation.set(qn("w:left"), str(720 + level * 360))
        indentation.set(qn("w:hanging"), "360")
        paragraph_properties.append(indentation)
        level_element.append(paragraph_properties)
        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), marker_font)
        fonts.set(qn("w:hAnsi"), marker_font)
        run_properties.append(fonts)
        level_element.append(run_properties)
        abstract.append(level_element)
    return abstract


def _append_value_element(parent: BaseOxmlElement, tag: str, value: str) -> None:
    element = OxmlElement(tag)
    element.set(qn("w:val"), value)
    parent.append(element)


def _xml_member(owner: object, member_name: str) -> BaseOxmlElement:
    element = getattr(owner, member_name, None)
    if not isinstance(element, BaseOxmlElement):
        msg = f"python-docx {type(owner).__name__} has no usable {member_name} XML member"
        raise TypeError(msg)
    return element


def _properties_element(owner: object, member_name: str, tag: str) -> BaseOxmlElement:
    container = _xml_member(owner, member_name)
    properties = container.find(qn(tag))
    if isinstance(properties, BaseOxmlElement):
        return properties
    properties = OxmlElement(tag)
    container.insert(0, properties)
    return properties


def _configure_header(document: Document, logo_bytes: bytes, *, title: str) -> None:
    paragraph = document.sections[0].header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(_HEADER.space_above_pt)
    paragraph.paragraph_format.space_after = Pt(_HEADER.space_below_pt)
    paragraph.paragraph_format.line_spacing = _HEADER.line_spacing_percent / 100
    paragraph.paragraph_format.keep_with_next = False
    title_run = paragraph.add_run(
        normalize_header_title(title, max_chars=_HEADER.max_title_chars) + HEADER_TITLE_SEPARATOR
    )
    _format_run(
        title_run,
        font=_HEADER.font_family,
        size=_HEADER.size_pt,
        color=_HEADER.color.removeprefix("#"),
    )
    title_run.bold = _HEADER.bold
    dimensions = _fitted_logo_dimensions()
    paragraph.add_run().add_picture(
        BytesIO(logo_bytes),
        width=Pt(dimensions.width),
        height=Pt(dimensions.height),
    )
    drawing_properties = _xml_member(paragraph, "_p").xpath(".//wp:docPr")
    if drawing_properties:
        drawing_properties[0].set("name", "Sarj logo")
        drawing_properties[0].set("title", "Sarj logo")
        drawing_properties[0].set("descr", trusted_document_logo_description())


def _fitted_logo_dimensions() -> _LogoDimensions:
    asset_width = CURRENT_BRAND_MANIFEST.assets.logo_pixel_width
    asset_height = CURRENT_BRAND_MANIFEST.assets.logo_pixel_height
    if asset_width is None or asset_height is None:  # pragma: no cover - manifest invariant
        msg = "Current Sarj logo asset must define pixel dimensions"
        raise RuntimeError(msg)
    scale = min(_HEADER.logo_width_pt / asset_width, _HEADER.logo_height_pt / asset_height)
    return _LogoDimensions(asset_width * scale, asset_height * scale)


def _configure_footer(document: Document) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(4)
    label = paragraph.add_run("Sarj  ·  ")
    _format_run(label, size=_DOCUMENT.caption_size_pt, color=_MUTED)
    _append_page_field(paragraph)


def _append_page_field(paragraph: Paragraph) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate):
        run = paragraph.add_run()
        _format_run(run, size=_DOCUMENT.caption_size_pt, color=_MUTED)
        _xml_member(run, "_r").append(element)
    value = paragraph.add_run("1")
    _format_run(value, size=_DOCUMENT.caption_size_pt, color=_MUTED)
    final = paragraph.add_run()
    _format_run(final, size=_DOCUMENT.caption_size_pt, color=_MUTED)
    _xml_member(final, "_r").append(end)


@dataclass(slots=True)
class _AppendListState:
    ordered_by_level: dict[int, int]
    next_num: int
    default_ordered_available: bool = True


def _append_blocks(
    document: Document,
    blocks: Iterable[_Block],
    numbering: _NumberingIds,
    *,
    branded: bool,
) -> None:
    lists = _AppendListState({}, numbering.next_num)
    for block in blocks:
        if block.kind == "list":
            _append_list_block(document, block, numbering, lists, branded=branded)
        elif block.kind == "check":
            _drop_ordered_at_or_above(lists.ordered_by_level, block.level)
            paragraph = document.add_paragraph(style="Sarj List" if branded else "List Paragraph")
            _set_numbering(
                paragraph,
                numbering.checked if block.checked else numbering.unchecked,
                block.level,
            )
            _append_inlines(paragraph, block.inlines, branded=branded)
        else:
            _append_nonlist_block(document, block, branded=branded)
            lists.ordered_by_level.clear()


def _append_nonlist_block(document: Document, block: _Block, *, branded: bool) -> None:
    if block.kind == "heading":
        paragraph = document.add_paragraph(style=f"Heading {min(max(block.level, 1), 6)}")
        _append_inlines(paragraph, block.inlines, branded=branded)
    elif block.kind == "paragraph":
        paragraph = document.add_paragraph(style="Normal")
        _append_inlines(paragraph, block.inlines, branded=branded)
    elif block.kind == "quote":
        paragraph = document.add_paragraph(style="Sarj Quote" if branded else "Quote")
        if branded:
            _set_paragraph_box(paragraph, fill=_SARJ_PURPLE_LIGHT, left_border=_SARJ_PURPLE)
        _append_inlines(paragraph, block.inlines, branded=branded)
    elif block.kind == "code":
        _append_code_block(document, block.text, branded=branded)
    elif block.kind == "table":
        _append_table(document, block, branded=branded)
    elif block.kind == "rule":
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        _set_paragraph_bottom_border(paragraph)


def _append_list_block(
    document: Document,
    block: _Block,
    numbering: _NumberingIds,
    state: _AppendListState,
    *,
    branded: bool,
) -> None:
    paragraph = document.add_paragraph(style="Sarj List" if branded else "List Paragraph")
    for level in tuple(state.ordered_by_level):
        if level > block.level:
            del state.ordered_by_level[level]
    if block.ordered and block.level not in state.ordered_by_level:
        _start_ordered_list(document, block, numbering, state)
    if not block.ordered:
        _drop_ordered_at_or_above(state.ordered_by_level, block.level)
    _set_numbering(
        paragraph,
        state.ordered_by_level[block.level] if block.ordered else numbering.bullet,
        block.level,
    )
    _append_inlines(paragraph, block.inlines, branded=branded)


def _start_ordered_list(
    document: Document, block: _Block, numbering: _NumberingIds, state: _AppendListState
) -> None:
    if state.default_ordered_available:
        state.ordered_by_level[block.level] = numbering.ordered
        state.default_ordered_available = False
        if block.start != 1:
            _append_numbering_override(
                document, num_id=numbering.ordered, level=block.level, start=block.start
            )
        return
    state.ordered_by_level[block.level] = state.next_num
    _append_concrete_numbering(
        document,
        abstract_id=numbering.ordered_abstract,
        num_id=state.next_num,
        level=block.level,
        start=block.start,
    )
    state.next_num += 1


def _drop_ordered_at_or_above(ordered_by_level: dict[int, int], level: int) -> None:
    for active_level in tuple(ordered_by_level):
        if active_level >= level:
            del ordered_by_level[active_level]


def _append_concrete_numbering(
    document: Document,
    *,
    abstract_id: int,
    num_id: int,
    level: int,
    start: int,
) -> None:
    concrete = OxmlElement("w:num")
    concrete.set(qn("w:numId"), str(num_id))
    _append_value_element(concrete, "w:abstractNumId", str(abstract_id))
    _append_level_start_override(concrete, level=level, start=start)
    document.part.numbering_part.element.append(concrete)


def _append_numbering_override(
    document: Document,
    *,
    num_id: int,
    level: int,
    start: int,
) -> None:
    numbering = document.part.numbering_part.element
    concrete = numbering.find(f"{qn('w:num')}[@{qn('w:numId')}='{num_id}']")
    if not isinstance(concrete, BaseOxmlElement):
        msg = f"Concrete numbering definition is missing: {num_id}"
        raise TypeError(msg)
    _append_level_start_override(concrete, level=level, start=start)


def _append_level_start_override(
    concrete: BaseOxmlElement,
    *,
    level: int,
    start: int,
) -> None:
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), str(level))
    _append_value_element(override, "w:startOverride", str(start))
    concrete.append(override)


def _set_numbering(paragraph: Paragraph, num_id: int, level: int) -> None:
    properties = _properties_element(paragraph, "_p", "w:pPr")
    existing = properties.find(qn("w:numPr"))
    if existing is not None:
        properties.remove(existing)
    num_properties = OxmlElement("w:numPr")
    _append_value_element(num_properties, "w:ilvl", str(min(max(level, 0), 2)))
    _append_value_element(num_properties, "w:numId", str(num_id))
    properties.append(num_properties)


def _append_code_block(document: Document, text: str, *, branded: bool) -> None:
    paragraph = document.add_paragraph(style="Sarj Code" if branded else "Normal")
    _set_paragraph_direction(paragraph, text)
    if branded:
        _set_paragraph_box(paragraph, fill=_CODE_FILL)
    for index, line in enumerate(text.split("\n")):
        if index:
            paragraph.add_run().add_break()
        # An empty source line is represented by the adjacent break. Adding a
        # space here changes the code block's content after DOCX conversion.
        run = paragraph.add_run(line)
        _format_run(
            run,
            font="Courier New",
            size=_DOCUMENT.body_size_pt if branded else 10,
            color=_INK if branded else "000000",
        )


def _append_table(document: Document, block: _Block, *, branded: bool) -> None:
    if not block.rows:
        return
    column_count = len(block.rows[0])
    widths = _table_widths(block.rows, column_count)
    table = document.add_table(rows=len(block.rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _configure_table_geometry(table, widths, branded=branded)
    for row_index, (source_row, target_row) in enumerate(zip(block.rows, table.rows, strict=True)):
        row_properties = _properties_element(target_row, "_tr", "w:trPr")
        keep_together = OxmlElement("w:cantSplit")
        keep_together.set(qn("w:val"), "true")
        row_properties.append(keep_together)
        if block.has_header and row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            row_properties.append(repeat)
        source_inlines = block.cell_inlines[row_index]
        for _value, inlines, cell, width in zip(
            source_row, source_inlines, target_row.cells, widths, strict=True
        ):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.width = Inches(width / 1_440)
            _set_cell_width(cell, width)
            if branded and block.has_header and row_index == 0:
                _set_cell_fill(cell, _SARJ_PURPLE_LIGHT)
            paragraph = cell.paragraphs[0]
            paragraph.style = _paragraph_style(
                document.styles["Sarj Table Text" if branded else "Normal"]
            )
            _append_inlines(
                paragraph,
                inlines,
                force_bold=block.has_header and row_index == 0,
                branded=branded,
            )
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def _configure_table_geometry(
    table: Table,
    widths: tuple[int, ...],
    *,
    branded: bool,
) -> None:
    table_xml = _xml_member(table, "_tbl")
    properties = _required_child(table_xml, "w:tblPr")
    table_width = properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.insert(0, table_width)
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    indentation = OxmlElement("w:tblInd")
    indentation.set(qn("w:w"), "0")
    indentation.set(qn("w:type"), "dxa")
    properties.insert_element_before(
        indentation,
        "w:tblBorders",
        "w:shd",
        "w:tblLayout",
        "w:tblCellMar",
        "w:tblLook",
        "w:tblCaption",
        "w:tblDescription",
        "w:tblPrChange",
    )
    existing_layout = properties.first_child_found_in("w:tblLayout")
    if existing_layout is not None:
        properties.remove(existing_layout)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    properties.insert_element_before(
        layout,
        "w:tblCellMar",
        "w:tblLook",
        "w:tblCaption",
        "w:tblDescription",
        "w:tblPrChange",
    )
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), _BORDER if branded else "A6A6A6")
        borders.append(border)
    properties.insert_element_before(
        borders,
        "w:shd",
        "w:tblLayout",
        "w:tblCellMar",
        "w:tblLook",
        "w:tblCaption",
        "w:tblDescription",
        "w:tblPrChange",
    )
    margins = OxmlElement("w:tblCellMar")
    for edge, width in (("top", 100), ("left", 120), ("bottom", 100), ("right", 120)):
        margin = OxmlElement(f"w:{edge}")
        margin.set(qn("w:w"), str(width))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    properties.insert_element_before(
        margins,
        "w:tblLook",
        "w:tblCaption",
        "w:tblDescription",
        "w:tblPrChange",
    )
    grid_columns = _required_child(table_xml, "w:tblGrid").findall(qn("w:gridCol"))
    for grid_column, width in zip(grid_columns, widths, strict=True):
        # ``CT_TblGridCol.w`` treats a bare int as EMUs. Write the OOXML dxa
        # value directly so Google Docs does not collapse every column to a
        # handful of twips during import.
        grid_column.set(qn("w:w"), str(width))


def _required_child(parent: BaseOxmlElement, tag: str) -> BaseOxmlElement:
    child = parent.find(qn(tag))
    if not isinstance(child, BaseOxmlElement):
        msg = f"Required WordprocessingML child is missing: {tag}"
        raise TypeError(msg)
    return child


def _set_cell_width(cell: _Cell, width: int) -> None:
    properties = _properties_element(cell, "_tc", "w:tcPr")
    cell_width = properties.find(qn("w:tcW"))
    if cell_width is None:
        cell_width = OxmlElement("w:tcW")
        properties.insert(0, cell_width)
    cell_width.set(qn("w:w"), str(width))
    cell_width.set(qn("w:type"), "dxa")


def _set_cell_fill(cell: _Cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    _properties_element(cell, "_tc", "w:tcPr").insert_element_before(
        shading,
        "w:noWrap",
        "w:tcMar",
        "w:textDirection",
        "w:tcFitText",
        "w:vAlign",
        "w:hideMark",
        "w:headers",
        "w:cellIns",
        "w:cellDel",
        "w:cellMerge",
        "w:tcPrChange",
    )


def _append_inlines(
    paragraph: Paragraph,
    inlines: tuple[MarkdownImportInline, ...],
    *,
    force_bold: bool = False,
    branded: bool,
) -> None:
    text = "".join(inline.text for inline in inlines)
    _set_paragraph_direction(paragraph, text)
    for inline in inlines:
        run = paragraph.add_run()
        for index, part in enumerate(inline.text.split("\n")):
            if index:
                run.add_break()
            if part:
                run.add_text(part)
        if inline.url is not None:
            _format_inline_run(
                run,
                inline,
                force_bold=force_bold,
                hyperlink=True,
                branded=branded,
            )
            _wrap_run_in_hyperlink(paragraph, run, inline.url)
        else:
            _format_inline_run(run, inline, force_bold=force_bold, branded=branded)


def _format_inline_run(
    run: Run,
    inline: MarkdownImportInline,
    *,
    force_bold: bool,
    hyperlink: bool = False,
    branded: bool,
) -> None:
    if inline.bold or force_bold:
        run.bold = True
    if inline.italic:
        run.italic = True
    if inline.code:
        _format_run(
            run,
            font="Courier New",
            size=_DOCUMENT.body_size_pt if branded else 10,
            color=_INK if branded else "000000",
        )
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), _CODE_FILL if branded else "E7E6E6")
        _properties_element(run, "_r", "w:rPr").append(shading)
    elif hyperlink:
        run.font.color.rgb = RGBColor.from_string(_SARJ_PURPLE if branded else "0563C1")
        run.underline = True
    _set_run_direction(run, inline.text)


def _format_run(
    run: Run,
    *,
    font: str = _DOCUMENT.font_family,
    size: float,
    color: str,
) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    properties = _properties_element(run, "_r", "w:rPr")
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), font)
    _set_run_direction(run, run.text)


def _set_paragraph_direction(paragraph: Paragraph, text: str) -> None:
    if not first_strong_is_rtl(text):
        return
    properties = _properties_element(paragraph, "_p", "w:pPr")
    bidi = properties.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        properties.append(bidi)
    bidi.set(qn("w:val"), "true")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _set_run_direction(run: Run, text: str) -> None:
    if not first_strong_is_rtl(text):
        return
    properties = _properties_element(run, "_r", "w:rPr")
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    fonts.set(qn("w:cs"), _COMPLEX_SCRIPT_FONT)
    fonts.attrib.pop(qn("w:cstheme"), None)
    rtl = properties.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        properties.append(rtl)
    rtl.set(qn("w:val"), "true")


def _wrap_run_in_hyperlink(paragraph: Paragraph, run: Run, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    hyperlink.set(qn("w:history"), "1")
    run_xml = _xml_member(run, "_r")
    parent = run_xml.getparent()
    if parent is None:
        msg = "Run is detached from its paragraph"
        raise ValueError(msg)
    parent.remove(run_xml)
    hyperlink.append(run_xml)
    _xml_member(paragraph, "_p").append(hyperlink)


def _is_safe_link(url: str) -> bool:
    return urlparse(url).scheme.casefold() in {"http", "https", "mailto"}


def _set_paragraph_box(
    paragraph: Paragraph,
    *,
    fill: str,
    left_border: str | None = None,
) -> None:
    properties = _properties_element(paragraph, "_p", "w:pPr")
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    if left_border is not None:
        borders = OxmlElement("w:pBdr")
        border = OxmlElement("w:left")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "18")
        border.set(qn("w:space"), "8")
        border.set(qn("w:color"), left_border)
        borders.append(border)
        properties.insert_element_before(
            borders,
            "w:shd",
            "w:tabs",
            "w:bidi",
            "w:spacing",
            "w:ind",
            "w:jc",
            "w:rPr",
            "w:sectPr",
            "w:pPrChange",
        )
    properties.insert_element_before(
        shading,
        "w:tabs",
        "w:bidi",
        "w:spacing",
        "w:ind",
        "w:jc",
        "w:rPr",
        "w:sectPr",
        "w:pPrChange",
    )


def _set_paragraph_bottom_border(paragraph: Paragraph) -> None:
    borders = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "8")
    border.set(qn("w:space"), "1")
    border.set(qn("w:color"), _BORDER)
    borders.append(border)
    _properties_element(paragraph, "_p", "w:pPr").insert_element_before(
        borders,
        "w:shd",
        "w:tabs",
        "w:bidi",
        "w:spacing",
        "w:ind",
        "w:jc",
        "w:rPr",
        "w:sectPr",
        "w:pPrChange",
    )


def _table_widths(rows: tuple[tuple[str, ...], ...], columns: int) -> tuple[int, ...]:
    weights = [max(4.0, max(len(row[index]) for row in rows) ** 0.5) for index in range(columns)]
    minimum = _CONTENT_WIDTH_DXA * 0.14
    raw = [max(minimum, _CONTENT_WIDTH_DXA * weight / sum(weights)) for weight in weights]
    scale = _CONTENT_WIDTH_DXA / sum(raw)
    widths = [int(value * scale) for value in raw]
    widths[-1] += _CONTENT_WIDTH_DXA - sum(widths)
    return tuple(widths)


def _plain_inlines(text: str) -> tuple[MarkdownImportInline, ...]:
    return (MarkdownImportInline(text),) if text else ()


def _inline_text(inlines: tuple[MarkdownImportInline, ...]) -> str:
    return "".join(inline.text for inline in inlines)


def _include_display_title(blocks: tuple[_Block, ...], title: str) -> bool:
    # An authored leading H1 owns the visible title hierarchy. File metadata
    # remains independent; never replace authored text or inline formatting.
    return bool(title.strip()) and not (
        blocks and blocks[0].kind == "heading" and blocks[0].level == 1
    )


@dataclass(slots=True)
class _ListContext:
    ordered: bool
    start: int
    item_index: int = -1


@dataclass(slots=True)
class _BlockParseState:
    blocks: list[_Block] = field(default_factory=list)
    list_stack: list[_ListContext] = field(default_factory=list)
    item_paragraphs: list[int] = field(default_factory=list)
    quote_depth: int = 0
    pending_kind: Literal["heading", "paragraph"] | None = None
    pending_heading_level: int = 0
    table_rows: list[list[tuple[MarkdownImportInline, ...]]] | None = None
    current_table_row: list[tuple[MarkdownImportInline, ...]] | None = None


def _parse_blocks(
    markdown: str,
) -> tuple[_Block, ...]:
    state = _BlockParseState()
    for token in _MARKDOWN.parse(markdown):
        _consume_block_token(state, token)
    if (
        state.quote_depth
        or state.list_stack
        or state.item_paragraphs
        or state.table_rows is not None
        or state.pending_kind
    ):
        raise UnsupportedSourceError
    return tuple(state.blocks)


def _consume_block_token(state: _BlockParseState, token: Token) -> None:
    kind = token.type
    if kind in {
        "blockquote_open",
        "blockquote_close",
        "bullet_list_open",
        "ordered_list_open",
        "bullet_list_close",
        "ordered_list_close",
        "list_item_open",
        "list_item_close",
        "heading_open",
        "paragraph_open",
        "heading_close",
        "paragraph_close",
    }:
        _consume_container_token(state, token)
    elif kind in {
        "table_open",
        "tr_open",
        "tr_close",
        "table_close",
        "thead_open",
        "thead_close",
        "tbody_open",
        "tbody_close",
        "th_open",
        "th_close",
        "td_open",
        "td_close",
    }:
        _consume_table_token(state, kind)
    elif kind == "inline":
        _consume_block_inline(state, token)
    elif kind == "hr":
        if state.quote_depth or state.list_stack:
            raise UnsupportedSourceError
        state.blocks.append(_Block("rule"))
    elif kind in {"fence", "code_block"}:
        if state.quote_depth or state.list_stack or state.table_rows is not None:
            raise UnsupportedSourceError
        state.blocks.append(_Block("code", text=token.content.removesuffix("\n")))
    else:
        raise UnsupportedSourceError


def _consume_container_token(state: _BlockParseState, token: Token) -> None:
    kind = token.type
    if kind in {
        "bullet_list_open",
        "ordered_list_open",
        "bullet_list_close",
        "ordered_list_close",
        "list_item_open",
        "list_item_close",
    }:
        _consume_list_container_token(state, token)
    elif kind == "blockquote_open":
        state.quote_depth += 1
        if state.quote_depth > 1 or state.list_stack:
            raise UnsupportedSourceError
    elif kind == "blockquote_close":
        state.quote_depth -= 1
    elif kind == "heading_open":
        if state.quote_depth or state.list_stack or state.table_rows is not None:
            raise UnsupportedSourceError
        state.pending_kind = "heading"
        state.pending_heading_level = int(token.tag.removeprefix("h"))
    elif kind == "paragraph_open" and state.table_rows is None:
        state.pending_kind = "paragraph"
        if state.list_stack:
            if not state.item_paragraphs or state.item_paragraphs[-1]:
                raise UnsupportedSourceError
            state.item_paragraphs[-1] += 1


def _consume_list_container_token(state: _BlockParseState, token: Token) -> None:
    kind = token.type
    if kind in {"bullet_list_open", "ordered_list_open"}:
        if state.quote_depth or len(state.list_stack) > _MAX_LIST_LEVEL:
            raise UnsupportedSourceError
        start = token.attrGet("start") if kind == "ordered_list_open" else None
        state.list_stack.append(
            _ListContext(ordered=kind == "ordered_list_open", start=int(start or 1))
        )
    elif kind in {"bullet_list_close", "ordered_list_close"}:
        if not state.list_stack:
            raise UnsupportedSourceError
        state.list_stack.pop()
    elif kind == "list_item_open":
        if not state.list_stack:
            raise UnsupportedSourceError
        state.list_stack[-1].item_index += 1
        state.item_paragraphs.append(0)
    elif kind == "list_item_close":
        if not state.item_paragraphs:
            raise UnsupportedSourceError
        state.item_paragraphs.pop()


def _consume_block_inline(state: _BlockParseState, token: Token) -> None:
    inlines = _parse_inline_tokens(token.children)
    if state.table_rows is not None:
        if state.current_table_row is None:
            raise UnsupportedSourceError
        state.current_table_row.append(inlines)
        return
    if state.pending_kind is None:
        raise UnsupportedSourceError
    text = _inline_text(inlines)
    if state.pending_kind == "heading":
        block = _Block("heading", text=text, level=state.pending_heading_level, inlines=inlines)
    elif state.list_stack:
        context = state.list_stack[-1]
        marker = _strip_check_marker(inlines)
        block = _Block(
            "check" if marker.checked is not None else "list",
            text=_inline_text(marker.inlines),
            level=len(state.list_stack) - 1,
            ordered=context.ordered,
            start=context.start + context.item_index,
            checked=bool(marker.checked),
            inlines=marker.inlines,
        )
    else:
        block = _Block("quote" if state.quote_depth else "paragraph", text=text, inlines=inlines)
    state.blocks.append(block)
    state.pending_kind = None


def _consume_table_token(state: _BlockParseState, kind: str) -> None:
    if kind == "table_open":
        if state.quote_depth or state.list_stack or state.table_rows is not None:
            raise UnsupportedSourceError
        state.table_rows = []
    elif kind == "tr_open":
        if state.table_rows is None or state.current_table_row is not None:
            raise UnsupportedSourceError
        state.current_table_row = []
    elif kind == "tr_close":
        if state.table_rows is None or state.current_table_row is None:
            raise UnsupportedSourceError
        state.table_rows.append(state.current_table_row)
        state.current_table_row = None
    elif kind == "table_close":
        if state.table_rows is None or state.current_table_row is not None or not state.table_rows:
            raise UnsupportedSourceError
        state.blocks.append(_table_block_from_rows(state.table_rows))
        state.table_rows = None
    elif state.table_rows is None:
        raise UnsupportedSourceError


def _table_block_from_rows(table_rows: list[list[tuple[MarkdownImportInline, ...]]]) -> _Block:
    width = max(len(row) for row in table_rows)
    normalized_inlines = tuple(tuple(row + [()] * (width - len(row))) for row in table_rows)
    rows = tuple(tuple(_inline_text(cell) for cell in row) for row in normalized_inlines)
    return _Block("table", rows=rows, cell_inlines=normalized_inlines, has_header=True)


@dataclass(slots=True)
class _InlineParseState:
    result: list[MarkdownImportInline] = field(default_factory=list)
    bold_depth: int = 0
    italic_depth: int = 0
    link_url: str | None = None


def _parse_inline_tokens(tokens: list[Token] | None) -> tuple[MarkdownImportInline, ...]:
    if tokens is None:
        raise UnsupportedSourceError
    state = _InlineParseState()
    for token in tokens:
        _consume_inline_token(state, token)
        if state.bold_depth < 0 or state.italic_depth < 0:
            raise UnsupportedSourceError
    if state.bold_depth or state.italic_depth or state.link_url is not None:
        raise UnsupportedSourceError
    return tuple(state.result)


def _consume_inline_token(state: _InlineParseState, token: Token) -> None:
    kind = token.type
    match kind:
        case "strong_open":
            state.bold_depth += 1
        case "strong_close":
            state.bold_depth -= 1
        case "em_open":
            state.italic_depth += 1
        case "em_close":
            state.italic_depth -= 1
        case "link_open":
            if state.link_url is not None:
                raise UnsupportedSourceError
            destination = token.attrGet("href")
            if not isinstance(destination, str) or not _is_safe_link(destination):
                raise UnsupportedSourceError
            state.link_url = destination
        case "link_close":
            if state.link_url is None:
                raise UnsupportedSourceError
            state.link_url = None
        case "text" | "code_inline" | "softbreak" | "hardbreak":
            content = " " if kind == "softbreak" else "\n" if kind == "hardbreak" else token.content
            _append_inline(
                state.result,
                content,
                bold=state.bold_depth > 0,
                italic=state.italic_depth > 0,
                code=kind == "code_inline",
                url=state.link_url,
            )
        case _:
            raise UnsupportedSourceError


def _append_inline(
    result: list[MarkdownImportInline],
    text: str,
    *,
    bold: bool,
    italic: bool,
    code: bool,
    url: str | None,
) -> None:
    if not text:
        return
    value = MarkdownImportInline(text, bold=bold, italic=italic, code=code, url=url)
    previous = result[-1] if result else None
    if previous is not None and (
        previous.bold,
        previous.italic,
        previous.code,
        previous.url,
    ) == (bold, italic, code, url):
        result[-1] = MarkdownImportInline(
            previous.text + text,
            bold=bold,
            italic=italic,
            code=code,
            url=url,
        )
    else:
        result.append(value)


def _strip_check_marker(inlines: tuple[MarkdownImportInline, ...]) -> _CheckMarker:
    marker_length = 4
    visible = _inline_text(inlines)
    if len(visible) < marker_length or visible[:marker_length].casefold() not in {"[ ] ", "[x] "}:
        return _CheckMarker(None, inlines)
    marker_characters = 0
    for inline in inlines:
        if marker_characters >= marker_length:
            break
        if inline.bold or inline.italic or inline.code or inline.url is not None:
            return _CheckMarker(None, inlines)
        marker_characters += len(inline.text)
    remaining = marker_length
    stripped: list[MarkdownImportInline] = []
    for inline in inlines:
        if remaining >= len(inline.text):
            remaining -= len(inline.text)
            continue
        stripped.append(
            MarkdownImportInline(
                inline.text[remaining:],
                bold=inline.bold,
                italic=inline.italic,
                code=inline.code,
                url=inline.url,
            )
        )
        remaining = 0
    return _CheckMarker(visible[1].casefold() == "x", tuple(stripped))


def _validate_heading_hierarchy(blocks: tuple[_Block, ...]) -> None:
    previous_level = 0
    for block in blocks:
        if block.kind != "heading":
            continue
        if previous_level == 0 and block.level != 1:
            raise UnsupportedSourceError
        if previous_level and block.level > previous_level + 1:
            raise UnsupportedSourceError
        previous_level = block.level


def _canonicalize_archive(raw: bytes) -> bytes:
    output = BytesIO()
    with (
        ZipFile(BytesIO(raw)) as source,
        ZipFile(output, "w", compression=ZIP_DEFLATED, compresslevel=9) as target,
    ):
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    return output.getvalue()
